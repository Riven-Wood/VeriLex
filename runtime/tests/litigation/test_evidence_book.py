import io
import hashlib
import pytest
from pypdf import PdfReader
from tests.litigation.test_workflows import env
from tests.litigation.test_workbench import desk
from verilex_runtime.litigation.evidence_book import directory_pdf,save,export
from verilex_runtime.tools.registry import ToolError


def test_booklet_matches_directory_and_keeps_original(desk):
    service,ctx,_,_=desk
    root=service._root(ctx);path=root/'样本.pdf';data=directory_pdf(['第'+str(i)+'行' for i in range(80)]);path.write_bytes(data)
    original=hashlib.sha256(data).hexdigest()
    record=save(service,ctx,{'title':'本案证据目录','entries':[{'name':'证据二','group':'合同','purpose':'证明履行','content':'双方签署的合同及履行凭证','path':'样本.pdf','first':2,'last':3},{'name':'证据一','path':'样本.pdf','first':1,'last':1}]},0)
    result=export(service,ctx,record['revision'],'pdf');reader=PdfReader(root/result['path'])
    assert len(reader.pages)==result['pages']==3
    assert result['directory_pages']==0
    assert result['entries'][0]['first']==1 and result['entries'][0]['last']==2
    assert '本案证据目录' not in reader.pages[0].extract_text()
    assert '第35行' in reader.pages[0].extract_text() and '第0行' in reader.pages[2].extract_text()
    from docx import Document
    doc=export(service,ctx,result['record']['revision'],'docx')
    table=Document(root/doc['path']).tables[0]
    assert [c.text for c in table.rows[0].cells]==['编号','证据名称','证据内容','证明目的','证据来源','页码']
    assert table.rows[1].cells[-1].text=='1—2'
    assert table.rows[1].cells[2].text=='双方签署的合同及履行凭证'
    assert hashlib.sha256(path.read_bytes()).hexdigest()==original
    path.write_bytes(directory_pdf(['已更新']))
    with pytest.raises(ToolError):export(service,ctx,record['revision'],'pdf')


def test_range_failure_does_not_publish_partial_book(desk):
    service,ctx,_,_=desk;root=service._root(ctx)
    (root/'样本.pdf').write_bytes(directory_pdf(['单页']))
    record=save(service,ctx,{'entries':[{'name':'超范围','path':'样本.pdf','first':1,'last':2}]},0)
    with pytest.raises(ToolError):export(service,ctx,record['revision'],'pdf')
    assert not list((root/'文书起草').glob('证据册-*.pdf'))


def test_whole_file_selection_reorder_and_manual_text_persist(desk):
    from verilex_runtime.litigation.evidence_book import prepare,read
    service,ctx,_,_=desk;root=ctx.root_path
    (root/'二.pdf').write_bytes(directory_pdf(['第二份']))
    (root/'一.pdf').write_bytes(directory_pdf(['第一份']*50))
    record=save(service,ctx,{'entries':[{'name':'二','path':'二.pdf'},{'name':'一','path':'一.pdf','content':'人工核对内容','manual':True}]},0)
    record=prepare(service,ctx,record['revision'])
    assert [e['page_count'] for e in record['data']['entries']]==[1,2]
    record['data']['entries'].reverse()
    record=save(service,ctx,record['data'],record['revision'])
    result=export(service,ctx,record['revision'],'pdf')
    reader=PdfReader(root/result['path'])
    assert '第一份' in reader.pages[0].extract_text() and '第二份' in reader.pages[2].extract_text()
    assert read(service,ctx)['data']['entries'][0]['content']=='人工核对内容'


def test_native_word_and_image_conversion_keeps_originals(desk):
    import sys
    from pathlib import Path
    if sys.platform!='darwin' or not Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').is_file():pytest.skip('local LibreOffice required')
    from docx import Document
    from PIL import Image
    service,ctx,_,_=desk;root=ctx.root_path
    doc=Document();doc.add_paragraph('证据内容：合同约定价款为一百元。');doc.add_table(rows=1,cols=2).cell(0,0).text='表格证据';doc.save(root/'合同.docx')
    Image.new('RGB',(400,200),'blue').save(root/'照片.png')
    before={name:hashlib.sha256((root/name).read_bytes()).hexdigest() for name in ['合同.docx','照片.png']}
    record=save(service,ctx,{'entries':[{'name':name,'path':name} for name in before]},0)
    result=export(service,ctx,record['revision'],'pdf');reader=PdfReader(root/result['path'])
    assert len(reader.pages)==2
    assert '一百元' in reader.pages[0].extract_text() and '表格证据' in reader.pages[0].extract_text()
    assert len(reader.pages[1].images)==1
    assert before=={name:hashlib.sha256((root/name).read_bytes()).hexdigest() for name in before}


def test_conversion_failure_and_permission_do_not_publish(desk,monkeypatch):
    from verilex_runtime.litigation import evidence_book as book
    service,ctx,workspace,_=desk;root=ctx.root_path
    (root/'损坏.docx').write_bytes(b'not a document')
    record=save(service,ctx,{'entries':[{'name':'损坏','path':'损坏.docx'}]},0)
    def failed(*_):raise ToolError('litigation_conversion_failed','转换失败')
    monkeypatch.setattr(book,'convert',failed)
    with pytest.raises(ToolError,match='损坏'):export(service,ctx,record['revision'],'pdf')
    assert not list((root/'文书起草').glob('证据册-*.pdf'))
    workspace.allow_write=False
    with pytest.raises(ToolError):book.prepare(service,ctx,record['revision'])


def test_ai_partial_read_tracks_gaps_and_protects_edits(desk):
    from verilex_runtime.litigation.evidence_ai import next_material,save_summary
    from verilex_runtime.litigation.workbench import WorkBinding
    service,ctx,_,_=desk;root=ctx.root_path
    (root/'长.pdf').write_bytes(directory_pdf(['第一份证据。'*12]*210))
    (root/'后.pdf').write_bytes(directory_pdf(['第二份证据。']))
    record=save(service,ctx,{'entries':[{'name':n,'path':n} for n in ['长.pdf','后.pdf']]},0)
    service.workbench.bind(ctx,ctx.thread_id,WorkBinding(kind='evidence',materials=['长.pdf','后.pdf'],expected_revision=0,evidence_revision=record['revision']))
    part=next_material(service,ctx)
    assert part['number']==1 and part['next_start'] is not None
    summary={'content':'概述','purpose':'证明合同履行','source':'当事人提供'}
    with pytest.raises(ToolError,match='超出'):next_material(service,ctx,999999)
    save_summary(service,ctx,part['token'],part['expected_revision'],summary)
    from verilex_runtime.litigation.evidence_book import read
    assert any('其余内容未由本工具读取' in gap for gap in read(service,ctx)['data']['entries'][0]['gaps'])
    part=next_material(service,ctx);assert part['number']==2
    from verilex_runtime.litigation.evidence_book import read
    edited=read(service,ctx);edited['data']['entries'][1].update(content='律师确认',manual=True)
    save(service,ctx,edited['data'],edited['revision'])
    with pytest.raises(ToolError,match='变化'):save_summary(service,ctx,part['token'],part['expected_revision'],summary)
    assert next_material(service,ctx)['done'] is True
    assert read(service,ctx)['data']['entries'][1]['content']=='律师确认'


def test_text_extraction_gap_does_not_overwrite_submitted_summary(desk):
    from verilex_runtime.litigation.evidence_ai import next_material,save_summary
    from verilex_runtime.litigation.workbench import WorkBinding
    from verilex_runtime.litigation.evidence_book import read
    from pypdf import PdfWriter
    service,ctx,_,_=desk;root=ctx.root_path
    writer=PdfWriter();writer.add_blank_page(595,842)
    with (root/'扫描.pdf').open('wb') as f:writer.write(f)
    record=save(service,ctx,{'entries':[{'name':'扫描','path':'扫描.pdf'}]},0)
    service.workbench.bind(ctx,ctx.thread_id,WorkBinding(kind='evidence',materials=['扫描.pdf'],expected_revision=0,evidence_revision=record['revision']))
    part=next_material(service,ctx);assert not part['text'] and part['gaps']
    save_summary(service,ctx,part['token'],part['expected_revision'],{'content':'由其他读取渠道取得的说明','purpose':'核对履行','source':'当事人提供'})
    saved=read(service,ctx)['data']['entries'][0]
    assert saved['status']=='unread' and saved['content']=='由其他读取渠道取得的说明' and saved['purpose']=='核对履行'
    assert any('没有可提取文字' in gap for gap in saved['gaps'])


def test_api_roundtrip_separate_exports(desk):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from verilex_runtime.api.litigation import build_litigation_router
    service,ctx,_,_=desk;(ctx.root_path/'材料.pdf').write_bytes(directory_pdf(['测试材料']))
    app=FastAPI();app.include_router(build_litigation_router(service));client=TestClient(app)
    base=f'/v1/litigation/workspaces/{ctx.workspace_id}/evidence-directory'
    saved=client.put(base,json={'expected_revision':0,'data':{'entries':[{'name':'材料','path':'材料.pdf'}]}}).json()
    prepared=client.post(base+'/prepare',json={'expected_revision':saved['revision']}).json()
    pdf=client.post(base+'/export/pdf',json={'expected_revision':prepared['revision']}).json()
    docx=client.post(base+'/export/docx',json={'expected_revision':pdf['record']['revision']}).json()
    assert pdf['pages']==1 and pdf['directory_pages']==0 and pdf['path']!=docx['path']


def test_ai_fills_other_columns_without_overwriting_manual_purpose(desk):
    from verilex_runtime.litigation.evidence_ai import next_material,save_summary
    from verilex_runtime.litigation.evidence_book import read
    from verilex_runtime.litigation.workbench import WorkBinding
    service,ctx,_,_=desk
    (ctx.root_path/'材料.pdf').write_bytes(directory_pdf(['合同价款100元']))
    record=save(service,ctx,{'entries':[{'name':'合同','path':'材料.pdf','purpose':'人工确认目的','manual':True,'manual_fields':['purpose']}]},0)
    service.workbench.bind(ctx,ctx.thread_id,WorkBinding(kind='evidence',materials=['材料.pdf'],expected_revision=0,evidence_revision=record['revision']))
    part=next_material(service,ctx)
    save_summary(service,ctx,part['token'],part['expected_revision'],{'content':'合同价款100元','purpose':'模型改写目的','source':'待核对'})
    item=read(service,ctx)['data']['entries'][0]
    assert item['purpose']=='人工确认目的' and item['content']=='合同价款100元'


def test_single_group_exports_without_group_heading_or_common_purpose(desk):
    from docx import Document
    service,ctx,_,_=desk
    (ctx.root_path/'材料.pdf').write_bytes(directory_pdf(['样本']))
    record=save(service,ctx,{'groups':[{'id':'only','name':'第一组证据','purpose':'整组目的不应输出','numbering':'restart','start_number':40}],
        'entries':[{'name':'证据甲','path':'材料.pdf','group':'only','purpose':'本条目的保留'},{'name':'证据乙','path':'材料.pdf','group':'only'}]},0)
    result=export(service,ctx,record['revision'],'docx');table=Document(ctx.root_path/result['path']).tables[0]
    assert len(table.rows)==3
    assert [row.cells[0].text for row in table.rows[1:]]==['1','2']
    text=' '.join(cell.text for row in table.rows for cell in row.cells)
    assert '本条目的保留' in text and '第一组证据' not in text and '整组目的不应输出' not in text
    assert result['record']['data']['groups'][0]['purpose']=='整组目的不应输出'
