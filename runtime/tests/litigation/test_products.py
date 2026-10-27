import io
import os
from unittest.mock import patch
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest
from tests.litigation.test_workflows import env
from verilex_runtime.litigation.products import Products
from verilex_runtime.litigation.editor import document_text
from verilex_runtime.api.litigation import build_litigation_router
from verilex_runtime.tools.registry import ToolError


def body(s):return {'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':s}]}]}

def create(service,ctx,s='第一份'):return service.products.create(ctx,'答辩状','documents',body(s))


def test_two_documents_reopen_draft_and_save_history(env):
    service,ctx,_=env;a=create(service,ctx);b=create(service,ctx,'第二份')
    assert a['id']!=b['id'] and a['path']!=b['path']
    assert a['title']!=b['title']
    service.products.draft(ctx,a['id'],a['file_hash'],body('未正式保存'))
    restarted=Products(service)
    assert restarted.read(ctx,a['id'])['draft_current']
    saved=restarted.save(ctx,a['id'],a['file_hash'],body('正式正文'))
    assert saved['revision']==2
    assert restarted.list(ctx)['errors']==[]
    assert len(restarted.list(ctx)['products'])==2
    assert document_text(restarted.read(ctx,b['id'])['document'])=='第二份'
    assert Document(ctx.root_path/saved['path']).paragraphs[0].text=='正式正文'
    assert restarted.history(ctx,a['id'])['versions'][0]['document']==body('第一份')
    assert not restarted.read(ctx,a['id'])['draft_current']
    assert service.files(ctx,'文书起草/编辑记录')['entries']==[]


def test_external_word_update_wins_and_old_draft_is_retained(env):
    service,ctx,_=env;a=create(service,ctx)
    service.products.draft(ctx,a['id'],a['file_hash'],body('我的草稿'))
    path=ctx.root_path/a['path'];d=Document(path);d.paragraphs[0].text='外部修改';d.save(path)
    before=path.read_bytes();r=service.products.read(ctx,a['id'])
    assert r['external_update'] and r['document'] is None and not r['draft_current']
    assert r['external_content']['paragraphs'][0]['text']=='外部修改'
    assert r['draft']['document']==body('我的草稿')
    for method in (service.products.save,service.products.draft):
        with pytest.raises(ToolError):method(ctx,a['id'],a['file_hash'],body('过时回写'))
    assert path.read_bytes()==before


def test_concurrent_save_rejects_stale_version_and_permission_changes(env):
    service,ctx,workspace=env;a=create(service,ctx)
    b=service.products.save(ctx,a['id'],a['file_hash'],body('新版'))
    with pytest.raises(ToolError):service.products.save(ctx,a['id'],a['file_hash'],body('晚到的保存'))
    workspace.allow_write=False
    with pytest.raises(ToolError):service.products.draft(ctx,a['id'],b['file_hash'],body('拒绝'))
    assert service.products.read(ctx,a['id'])['document']==body('新版')


def test_failed_publication_restores_old_word_and_keeps_draft(env):
    service,ctx,_=env;a=create(service,ctx);path=ctx.root_path/a['path'];before=path.read_bytes();link=os.link
    def fail_new(src,dst):
        if str(dst)==str(path) and '.verilex-edit-history-' not in str(src):raise OSError('simulated disk failure')
        return link(src,dst)
    with patch('verilex_runtime.litigation.products.os.link',side_effect=fail_new):
        with pytest.raises(ToolError):service.products.save(ctx,a['id'],a['file_hash'],body('待恢复'))
    assert path.read_bytes()==before
    assert service.products.read(ctx,a['id'])['draft']['document']==body('待恢复')


def test_draft_autosaves_do_not_create_history_per_keystroke(env):
    service,ctx,_=env;a=create(service,ctx)
    for s in ['甲','甲乙','甲乙丙']:service.products.draft(ctx,a['id'],a['file_hash'],body(s))
    assert list((ctx.root_path/'文书起草/编辑记录').glob('.verilex-edit-history-*'))==[]


def test_ids_and_symlinks_cannot_escape_workspace(env,tmp_path):
    service,ctx,_=env;a=create(service,ctx);path=ctx.root_path/a['path']
    for product_id in ['../secret','/etc/passwd','x'*32]:
        with pytest.raises(ToolError):service.products.read(ctx,product_id)
    external=tmp_path/'external.docx';path.rename(external);path.symlink_to(external)
    with pytest.raises(ToolError):service.products.read(ctx,a['id'])


def test_api_actual_create_draft_save_and_reopen(env):
    service,ctx,_=env;app=FastAPI();app.include_router(build_litigation_router(service));client=TestClient(app)
    base=f'/v1/litigation/workspaces/{ctx.workspace_id}/editor-products'
    r=client.post(base,json={'title':'新文书','kind':'documents','document':body('中文')});assert r.status_code==200
    a=r.json();target=base+'/'+a['id'];payload={'base_hash':a['file_hash'],'document':body('新正文')}
    assert client.put(target+'/draft',json=payload).status_code==200
    assert client.put(target,json=payload).status_code==200
    assert client.get(target).json()['document']==body('新正文')
    assert len(client.get(base).json()['products'])==1
    assert len(client.get(target+'/history').json()['versions'])==1


def test_personal_template_import_keeps_source_and_produces_independent_document(env):
    service,ctx,_=env;a=create(service,ctx,'个人模板内容')
    template=service.products.save_template(ctx,a['id'],a['file_hash'],body('个人模板内容'),'答辩模板')
    path=ctx.root_path/template['path'];before=path.read_bytes()
    copied=service.products.import_copy(ctx,template['path'],'模板副本','documents')
    assert document_text(copied['document'])=='个人模板内容'
    assert path.read_bytes()==before and copied['id']!=a['id']
    assert service.products.templates(ctx)['templates'][0]['path']==template['path']


def test_import_complex_word_requires_explicit_text_copy(env):
    from tests.litigation.test_editor import sample
    service,ctx,_=env;path=ctx.root_path/'复杂模板.docx';path.write_bytes(sample());before=path.read_bytes()
    with pytest.raises(ToolError):service.products.import_copy(ctx,path.name,'复杂模板','facts')
    r=service.products.import_copy(ctx,path.name,'正文副本','documents',True)
    assert '嵌套表格' in document_text(r['document'])
    assert path.read_bytes()==before


def test_late_draft_after_external_update_is_recoverable_without_overwriting_word(env):
    service,ctx,_=env;a=create(service,ctx);path=ctx.root_path/a['path']
    d=Document(path);d.paragraphs[0].text='外部新正文';d.save(path);before=path.read_bytes()
    with pytest.raises(ToolError):service.products.draft(ctx,a['id'],a['file_hash'],body('尚未保存的编辑'))
    current=service.products.read(ctx,a['id'])
    assert current['recovery']['document']==body('尚未保存的编辑')
    assert path.read_bytes()==before


def test_office_removing_custom_metadata_still_loads_latest_external_content(env):
    from verilex_runtime.documents.docx_package import parts,pack
    service,ctx,_=env;a=create(service,ctx)
    a=service.products.save(ctx,a['id'],a['file_hash'],body('保存过的正文'),'更新标题')
    path=ctx.root_path/a['path']
    d=Document(path);d.paragraphs[0].text='办公软件修改的最新正文';d.save(path)
    items=parts(path.read_bytes());items.pop('verilex/editor.json',None)
    path.write_bytes(pack(items));before=path.read_bytes()
    current=service.products.read(ctx,a['id'])
    assert current['revision']==2 and current['title']=='更新标题'
    assert current['external_update'] and current['document'] is None
    assert current['external_content']['paragraphs'][0]['text']=='办公软件修改的最新正文'
    assert service.products.list(ctx)['products'][0]['id']==a['id']
    assert service.products.list(ctx)['products'][0]['path']==a['path']
    with pytest.raises(ToolError):service.products.save(ctx,a['id'],a['file_hash'],body('过时保存'))
    assert path.read_bytes()==before


def test_index_cannot_redirect_edit_to_original_or_another_product(env):
    import json
    service,ctx,_=env;a=create(service,ctx);b=create(service,ctx,'另一份')
    index=ctx.root_path/'文书起草/编辑记录'/f"{a['id']}-index.json"
    stored=json.loads(index.read_text())
    for destination in ['合同.txt',b['path']]:
        index.write_text(json.dumps({**stored,'path':destination}),encoding='utf8')
        with pytest.raises(ToolError):service.products.save(ctx,a['id'],a['file_hash'],body('错误目标'))
    assert Document(ctx.root_path/b['path']).paragraphs[0].text=='另一份'


@pytest.mark.skipif(not hasattr(os, 'fork'), reason='POSIX abrupt process exit')
def test_process_exit_during_publication_keeps_formal_word_reopenable(env):
    service, ctx, _ = env
    item = create(service, ctx, '崩溃前正文')
    path = ctx.root_path / item['path']
    child = os.fork()
    if child == 0:
        link, replace = os.link, os.replace
        def stop_before_publish(operation):
            def wrapped(src, dst, *args, **kwargs):
                if str(dst) == str(path) and str(src) != str(path):
                    os._exit(73)
                return operation(src, dst, *args, **kwargs)
            return wrapped
        os.link = stop_before_publish(link)
        os.replace = stop_before_publish(replace)
        try:
            service.products.save(ctx, item['id'], item['file_hash'], body('未完成的新正文'))
        finally:
            os._exit(74)
    _, status = os.waitpid(child, 0)
    assert os.waitstatus_to_exitcode(status) == 73
    restarted = Products(service)
    assert len(restarted.list(ctx)['products']) == 1
    assert path.exists()
    assert restarted.read(ctx, item['id'])['document'] == body('崩溃前正文')
    assert len(restarted.list(ctx)['products']) == 1
    assert restarted.read(ctx, item['id'])['draft']['document'] == body('未完成的新正文')


def test_pending_recovery_never_overwrites_an_external_save(env):
    import json
    service, ctx, _ = env; item = create(service, ctx)
    path = ctx.root_path / item['path']
    backup = path.parent / ('.verilex-edit-history-'+'b'*32+'.docx')
    path.rename(backup)
    journal = path.parent / ('.verilex-edit-pending-'+'c'*32+'.json')
    journal.write_text(json.dumps({'target':path.name,'backup':backup.name}))
    external = Document(); external.add_paragraph('外部保存的新正文'); external.save(path)
    expected = path.read_bytes()
    loaded = Products(service).read(ctx, item['id'])
    assert loaded['external_update'] and path.read_bytes() == expected
    assert backup.exists() and not journal.exists()


def test_manual_removal_is_not_undone_using_old_history(env):
    service, ctx, _ = env; item = create(service, ctx)
    service.products.save(ctx, item['id'], item['file_hash'], body('新正文'))
    (ctx.root_path/item['path']).unlink()
    assert Products(service).list(ctx)['products'] == []
    with pytest.raises(ToolError): Products(service).read(ctx, item['id'])


def test_recovery_rejects_a_journal_target_outside_the_product_folder(env):
    import json
    service, ctx, _ = env; item = create(service, ctx)
    folder = (ctx.root_path/item['path']).parent
    (folder/('.verilex-edit-pending-'+'d'*32+'.json')).write_text(
        json.dumps({'target':'../outside.docx','backup':'.verilex-edit-history-'+'b'*32+'.docx'}))
    with pytest.raises(ToolError): Products(service).list(ctx)
    assert not (ctx.root_path/'outside.docx').exists()


@pytest.mark.skipif(not hasattr(os, 'fork'), reason='POSIX abrupt process exit')
def test_exit_while_writing_recovery_journal_does_not_block_intact_documents(env):
    import json
    service, ctx, _ = env; item = create(service, ctx, '原文仍完整')
    child = os.fork()
    if child == 0:
        def incomplete_dump(value, stream, *args, **kwargs):
            stream.write('{'); stream.flush(); os._exit(75)
        json.dump = incomplete_dump
        try:
            service.products.save(ctx, item['id'], item['file_hash'], body('新草稿'))
        finally:
            os._exit(76)
    _, status = os.waitpid(child, 0)
    assert os.waitstatus_to_exitcode(status) == 75
    restarted = Products(service)
    assert len(restarted.list(ctx)['products']) == 1
    assert restarted.read(ctx, item['id'])['document'] == body('原文仍完整')


def test_native_templates_preserve_word_parts_and_independent_identity(env):
    from docx.shared import Pt
    from tests.litigation.test_editor import sample
    from verilex_runtime.documents.docx_package import parts
    service,ctx,_=env
    doc=Document(io.BytesIO(sample()))
    doc.sections[0].header.paragraphs[0].text='模板页眉'
    doc.sections[0].footer.paragraphs[0].text='模板页脚'
    doc.paragraphs[0].runs[0].font.size=Pt(18)
    source=ctx.root_path/'自备模板.docx';doc.save(source);original=source.read_bytes()
    a=service.products.import_copy(ctx,source.name,'本案答辩稿','documents',document_type='答辩状')
    assert a['native_word'] and a['document_type']=='答辩状'
    original_parts=parts(original)
    def unchanged(data):
        result=parts(data)
        assert all(result[k]==v for k,v in original_parts.items() if k not in ('[Content_Types].xml','_rels/.rels'))
    opened=service.products.open_word(ctx,a['id'],a['file_hash'])['product']
    unchanged((ctx.root_path/opened['path']).read_bytes())
    saved=service.products.save_template(ctx,opened['id'],opened['file_hash'],opened['document'],'常用答辩格式')
    template=ctx.root_path/saved['path'];before=template.read_bytes();unchanged(before)
    b=service.products.import_copy(ctx,saved['path'],'本案答辩稿','documents',document_type='答辩状')
    assert b['id']!=a['id'] and b['path']!=a['path'] and b['title']=='本案答辩稿（2）'
    unchanged((ctx.root_path/b['path']).read_bytes())
    service.products.rename(ctx,b['id'],b['file_hash'],'第二案件答辩')
    assert source.read_bytes()==original and template.read_bytes()==before


def test_native_template_rejects_stale_body_external_edits_and_readonly(env):
    service,ctx,workspace=env;a=create(service,ctx)
    a=service.products.open_word(ctx,a['id'],a['file_hash'])['product']
    for hash,document in [('0'*64,a['document']),(a['file_hash'],body('未保存输入'))]:
        with pytest.raises(ToolError):service.products.save_template(ctx,a['id'],hash,document,'模板')
    workspace.allow_write=False
    with pytest.raises(ToolError):service.products.save_template(ctx,a['id'],a['file_hash'],a['document'],'模板')
    workspace.allow_write=True
    path=ctx.root_path/a['path'];doc=Document(path);doc.paragraphs[0].text='外部修改';doc.save(path)
    changed=service.products.read(ctx,a['id'])
    with pytest.raises(ToolError):service.products.save_template(ctx,a['id'],changed['file_hash'],a['document'],'模板')
    assert service.products.templates(ctx)=={'templates':[]}


def test_template_api_preserves_native_copy_type_and_rejects_outside_path(env,tmp_path):
    service,ctx,_=env;a=create(service,ctx)
    app=FastAPI();app.include_router(build_litigation_router(service));client=TestClient(app)
    url=f'/v1/litigation/workspaces/{ctx.workspace_id}'
    r=client.post(url+'/editor-import',json={'path':a['path'],'title':'自定义名称','kind':'documents','document_type':'上诉状'})
    assert r.status_code==200 and r.json()['document_type']=='上诉状' and r.json()['native_word']
    with pytest.raises(ToolError):service.products.import_copy(ctx,'../越界.docx','越界','documents')
    outside=tmp_path/'outside.docx';outside.write_bytes((ctx.root_path/a['path']).read_bytes())
    linked=ctx.root_path/'link.docx';linked.symlink_to(outside)
    with pytest.raises(ToolError):service.products.import_copy(ctx,linked.name,'链接','documents')


def test_upload_template_exact_bytes_names_permissions_and_reuse(env):
    from tests.litigation.test_editor import sample
    service,ctx,workspace=env
    original=sample()
    first=service.products.import_template(ctx,original,'external.docx','上传格式')
    second=service.products.import_template(ctx,original,'external.docx','上传格式')
    assert first['name']=='上传格式' and second['name']=='上传格式（2）'
    assert (ctx.root_path/first['path']).read_bytes()==original
    assert (ctx.root_path/second['path']).read_bytes()==original
    copy=service.products.import_copy(ctx,first['path'],'本案文书','documents',document_type='上传格式')
    assert copy['native_word'] and copy['document_type']=='上传格式'
    assert (ctx.root_path/first['path']).read_bytes()==original
    for data,filename,title in [(b'bad','x.docx','格式'),(b'','x.docx','格式'),(original,'x.doc','格式'),(original,'x.docx','../越界'),(original,'x.docx','...')]:
        with pytest.raises(ToolError):service.products.import_template(ctx,data,filename,title)
    workspace.allow_write=False
    with pytest.raises(ToolError):service.products.import_template(ctx,original,'x.docx','只读')


def test_template_upload_http_limit_allows_word_and_rejects_oversize(env):
    import zipfile
    from verilex_runtime.api.limits import RequestBodyLimitMiddleware
    from tests.litigation.test_editor import sample
    service,ctx,_=env
    stream=io.BytesIO(sample())
    with zipfile.ZipFile(stream,'a',compression=zipfile.ZIP_STORED) as archive:archive.writestr('padding.bin',b'x'*(1100*1024))
    data=stream.getvalue()
    app=FastAPI();app.include_router(build_litigation_router(service));app.add_middleware(RequestBodyLimitMiddleware)
    client=TestClient(app);url=f'/v1/litigation/workspaces/{ctx.workspace_id}/editor-templates'
    result=client.post(url,params={'filename':'external.docx','title':'上传'},content=data,headers={'content-type':'application/octet-stream'})
    assert result.status_code==201
    assert (ctx.root_path/result.json()['path']).read_bytes()==data
    assert client.post(url,params={'filename':'x.docx','title':'过大'},content=b'x'*(20*1024*1024+1)).status_code==413
    assert client.post(url+'/other',content=data).status_code==413
