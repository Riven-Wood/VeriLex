import io

import pytest
from docx import Document
from docx.shared import Pt

from verilex_runtime.litigation.native_word import project, update_document
from verilex_runtime.tools.registry import ToolError


def test_splitting_paragraph_preserves_surviving_run_format_and_other_parts():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('原有').bold = True
    paragraph.add_run('文字').font.size = Pt(18)
    stream = io.BytesIO(); document.save(stream)
    source = stream.getvalue()
    before = project(source)
    after = {'type': 'doc', 'content': [
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': '原有'}]},
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': '文字'}]},
    ]}
    result=update_document(source,before,after)
    reopened=Document(io.BytesIO(result))
    assert [p.text for p in reopened.paragraphs]==['原有','文字']
    assert next(r for r in reopened.paragraphs[0].runs if r.text=='原有').bold
    assert next(r for r in reopened.paragraphs[1].runs if r.text=='文字').font.size==Pt(18)
    from verilex_runtime.documents.docx_package import parts
    original,updated=parts(source),parts(result)
    assert all(original[name]==updated[name] for name in original if name!='word/document.xml')
    assert Document(io.BytesIO(source)).paragraphs[0].text=='原有文字'


def test_merge_preserves_run_formats_when_paragraph_properties_match():
    document=Document();document.add_paragraph().add_run('第一段').bold=True
    document.add_paragraph().add_run('第二段').font.size=Pt(18)
    stream=io.BytesIO();document.save(stream);source=stream.getvalue()
    after={'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':'第一段第二段'}]}]}
    result=Document(io.BytesIO(update_document(source,project(source),after)))
    assert result.paragraphs[0].text=='第一段第二段'
    assert result.paragraphs[0].runs[0].bold
    assert result.paragraphs[0].runs[1].font.size==Pt(18)


def test_merge_refuses_to_flatten_different_paragraph_formatting():
    document=Document();document.add_paragraph('第一段');document.add_paragraph('第二段').paragraph_format.left_indent=Pt(24)
    stream=io.BytesIO();document.save(stream);source=stream.getvalue()
    after={'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':'第一段第二段'}]}]}
    with pytest.raises(ToolError,match='不能完整保留原有文字格式'):
        update_document(source,project(source),after)


def test_multi_paragraph_rewrite_keeps_uniform_run_formatting():
    document=Document()
    document.add_paragraph().add_run('原文甲').bold=True
    document.add_paragraph().add_run('原文乙').bold=True
    stream=io.BytesIO();document.save(stream);source=stream.getvalue()
    after={'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':text}]} for text in ('新段一','新段二','新段三')]}
    reopened=Document(io.BytesIO(update_document(source,project(source),after)))
    assert [p.text for p in reopened.paragraphs]==['新段一','新段二','新段三']
    assert all(next(r for r in p.runs if r.text).bold for p in reopened.paragraphs)


def test_projection_does_not_claim_complete_when_header_has_content():
    from verilex_runtime.litigation.native_word import projection_is_complete
    document=Document();document.add_paragraph('正文')
    buffer=io.BytesIO();document.save(buffer)
    assert projection_is_complete(buffer.getvalue())
    document.sections[0].header.paragraphs[0].text='页眉中的额外内容'
    buffer=io.BytesIO();document.save(buffer)
    assert not projection_is_complete(buffer.getvalue())
