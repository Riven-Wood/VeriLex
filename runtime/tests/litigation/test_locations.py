from docx import Document
from tests.litigation.test_workflows import env
from tests.litigation.test_workbench import desk
from verilex_runtime.litigation.models import Citation,SourceLocator


def test_duplicate_word_quote_never_guesses_and_structural_locator_is_exact(desk):
    service,ctx,_,_=desk;path=service._root(ctx)/'来源.docx'
    doc=Document();doc.add_paragraph('相同引文');doc.add_table(rows=1,cols=1).cell(0,0).text='相同引文';doc.save(path)
    source=Citation(path='来源.docx',quote='相同引文')
    assert service.locate_source(ctx,source)['status']=='ambiguous'
    from verilex_runtime.litigation.editor import inspect_docx
    inspection=inspect_docx(path.read_bytes())
    source.locator=SourceLocator(file_hash=inspection['file_hash'],kind='docx',key=inspection['paragraphs'][1]['id'],index=1,part='word/document.xml')
    result=service.locate_source(ctx,source)
    assert result['status']=='located' and result['locator']['index']==1
    doc.paragraphs[0].text='已经修改';doc.save(path)
    result=service.locate_source(ctx,source)
    assert result['status']=='located' and result['notice']
    path.unlink();assert service.locate_source(ctx,source)['status']=='missing'
