import base64
import io
from dataclasses import replace
import pytest
from docx import Document
from pypdf import PdfReader
from tests.litigation.test_workflows import env
from tests.litigation.test_workbench import desk
from tests.litigation.test_products import body
from verilex_runtime.tools.registry import ToolError
from verilex_runtime.litigation.editor import document_text

EMPTY={'type':'doc','content':[{'type':'paragraph'}]}
def initial(desk,title='答辩状'):
    service,ctx,_,_=desk
    p=service.products.create(ctx,title,'documents',EMPTY)
    snapshot=service.editor_discussions.prepare(ctx,{'product_id':p['id'],'base_hash':p['file_hash'],'document':EMPTY,'operation':'generate'})
    service.editor_discussions.bind(ctx,snapshot)
    return service,ctx,p

def test_initial_generation_fills_bound_document_once_then_proposals_preserve_body(desk):
    s,c,p=initial(desk)
    s.editor_discussions.generate(c,'答辩状\n请依法驳回原告请求。')
    current=s.products.read(c,p['id'])
    assert current['revision']==2
    assert document_text(current['document'])=='答辩状\n请依法驳回原告请求。'
    with pytest.raises(ToolError):s.editor_discussions.generate(c,'覆盖')
    c=replace(c,turn_id='rewrite-turn')
    s.editor_discussions.bind(c,s.editor_discussions.prepare(c,{'product_id':p['id'],'base_hash':current['file_hash'],'document':current['document']}))
    suggestion=s.editor_discussions.propose(c,'答辩状\n新的正文。')
    assert s.products.read(c,p['id'])['document']==current['document']
    with pytest.raises(ToolError):s.editor_discussions.generate(c,'绕过批注')
    accepted=s.editor_discussions.decide(c,suggestion['id'],'accept',current['document'])
    assert document_text(accepted['document'])=='答辩状\n新的正文。'

def test_generation_cannot_overwrite_new_typing_or_other_thread(desk):
    s,c,p=initial(desk)
    s.products.draft(c,p['id'],p['file_hash'],body('用户正在输入'))
    with pytest.raises(ToolError):s.editor_discussions.generate(c,'AI内容')
    with pytest.raises(ToolError):s.editor_discussions.generate(replace(c,thread_id='thread-2'),'AI内容')
    assert s.products.read(c,p['id'])['draft']['document']==body('用户正在输入')
    with pytest.raises(ToolError):s.editor_discussions.prepare(c,{'product_id':p['id'],'base_hash':p['file_hash'],'document':EMPTY,'operation':'generate'})

def test_initial_authorization_rejects_saved_body_or_selection(desk):
    s,c,_,_=desk
    p=s.products.create(c,'答辩状','documents',body('已有正文'))
    with pytest.raises(ToolError):s.editor_discussions.prepare(c,{'product_id':p['id'],'base_hash':p['file_hash'],'document':EMPTY,'operation':'generate'})
    with pytest.raises(ToolError):s.editor_discussions.prepare(c,{'product_id':p['id'],'base_hash':p['file_hash'],'document':body('已有正文'),'operation':'generate','selection':{'from':1,'to':3,'text':'已有'}})

def test_export_current_body_to_real_docx_and_pdf_and_reject_stale(desk):
    s,c,p=initial(desk)
    document=body('答辩状\nCurrent body 2026 手动修改后的正文')
    word=s.products.export_document(c,p['id'],p['file_hash'],document,'答辩状','docx')
    exported=Document(io.BytesIO(base64.b64decode(word['content'])))
    assert '手动修改后的正文' in '\n'.join(p.text for p in exported.paragraphs)
    with pytest.raises(ToolError):s.products.export_document(c,p['id'],p['file_hash'],EMPTY,'答辩状','docx')
    pdf=s.products.export_document(c,p['id'],word['product']['file_hash'],document,'答辩状','pdf')
    pages=PdfReader(io.BytesIO(base64.b64decode(pdf['content']))).pages
    assert 'Current body 2026' in '\n'.join(p.extract_text() for p in pages)
    assert s.products.read(c,p['id'])['document']==document
    assert (c.root_path/'合同.txt').read_text()=='2026年1月2日双方签订合同。应支付货款100元。'

def test_pdf_conversion_concurrent_edit_is_reported_not_exported(desk,monkeypatch):
    s,c,p=initial(desk)
    def conversion(data,ext):
        current=s.products.read(c,p['id'])
        s.products.draft(c,p['id'],current['file_hash'],body('转换期间的新输入'))
        return b'%PDF-1.7'
    monkeypatch.setattr('verilex_runtime.litigation.evidence_conversion.convert',conversion)
    with pytest.raises(ToolError,match='转换期间'):s.products.export_document(c,p['id'],p['file_hash'],body('旧正文'),'文书','pdf')
    assert s.products.read(c,p['id'])['draft']['document']==body('转换期间的新输入')


def cross_table():
    paragraph=lambda text:{'type':'paragraph','content':[{'type':'text','text':text,'marks':[{'type':'textStyle','attrs':{'fontFamily':'仿宋','fontSize':'12pt'}}]}]}
    rows=[]
    for values in [('序号','证据名称','页码','质证意见'),('5','对方合同','1—2','原件待核对，合同不能单独证明付款。'),('8','收据','3','请结合付款流水核对。')]:
        rows.append({'type':'tableRow','content':[{'type':'tableCell','attrs':{'colwidth':[width]},'content':[paragraph(text)]} for width,text in zip([45,110,65,340],values)]})
    return {'type':'doc','content':[{'type':'heading','attrs':{'level':1},'content':[{'type':'text','text':'质证意见'}]},{'type':'table','content':rows},paragraph('质证人：待确认')]}


def test_structured_generation_creates_native_editable_table_and_preserves_original_materials(desk):
    s,c,p=initial(desk,'质证意见');document=cross_table()
    s.editor_discussions.generate(c,document=document)
    current=s.products.read(c,p['id'])
    assert any(n['type']=='table' for n in current['document']['content'])
    output=s.products.export_document(c,p['id'],current['file_hash'],current['document'],'质证意见','docx')
    word=Document(io.BytesIO(base64.b64decode(output['content'])))
    assert word.paragraphs[0].text=='质证意见'
    assert word.paragraphs[0].alignment==1
    table=word.tables[0]
    assert len(table.columns)==4 and len(table.rows)==3
    assert [cell.text for cell in table.rows[1].cells][:3]==['5','对方合同','1—2']
    assert table.columns[3].width>table.columns[1].width>table.columns[0].width
    assert table.cell(1,3).paragraphs[0].runs[0].font.size.pt==12
    with pytest.raises(ToolError):s.editor_discussions.generate(c,document=cross_table())
    assert (c.root_path/'合同.txt').read_text()=='2026年1月2日双方签订合同。应支付货款100元。'


def test_structured_generation_rejects_invalid_documents_and_does_not_overwrite_drafts(desk):
    s,c,p=initial(desk)
    for document in [cross_table()|{'other':True},{'type':'doc','content':[{'type':'image'}]}]:
        with pytest.raises(ToolError):s.editor_discussions.generate(c,document=document)
    with pytest.raises(ToolError):s.editor_discussions.generate(c,'text',cross_table())
    invalid=cross_table();invalid['content'][1]['content'][0]['content'][0]['attrs']['colwidth']=[-1]
    with pytest.raises(ToolError):s.editor_discussions.generate(c,document=invalid)
    assert s.products.read(c,p['id'])['revision']==p['revision']
    s.products.draft(c,p['id'],p['file_hash'],body('尚未保存的人工正文'))
    with pytest.raises(ToolError):s.editor_discussions.generate(c,document=cross_table())
    assert document_text(s.products.read(c,p['id'])['draft']['document'])=='尚未保存的人工正文'


def test_cross_template_allows_user_requested_prose(desk):
    s,c,p=initial(desk,'质证意见')
    s.editor_discussions.generate(c,'按用户要求，以段落说明质证意见。')
    result=s.products.read(c,p['id'])
    assert '以段落说明质证意见' in document_text(result['document'])
