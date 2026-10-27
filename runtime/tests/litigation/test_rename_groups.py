import io
from unittest.mock import patch
from zipfile import ZipFile
import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from tests.litigation.test_workflows import env
from tests.litigation.test_workbench import desk
from tests.litigation.test_products import body,create
from verilex_runtime.litigation.products import Products,META
from verilex_runtime.litigation.evidence_book import save,export,directory_pdf,Inventory,numbered_entries
from verilex_runtime.tools.registry import ToolError


def test_rename_preserves_docx_parts_id_type_draft_and_exports(env):
    service,ctx,_=env;a=create(service,ctx);old=ctx.root_path/a['path'];raw=old.read_bytes()
    service.products.draft(ctx,a['id'],a['file_hash'],body('待保存的正文'))
    b=service.products.rename(ctx,a['id'],a['file_hash'],'案件甲-第二稿.docx')
    assert b['id']==a['id'] and b['document_type']==a['document_type'] and b['title']=='案件甲-第二稿'
    assert not old.exists() and b['draft_current'] and b['draft']['document']==body('待保存的正文')
    with ZipFile(io.BytesIO(raw)) as before,ZipFile(ctx.root_path/b['path']) as after:
        assert {name:before.read(name) for name in before.namelist() if name!=META}=={name:after.read(name) for name in after.namelist() if name!=META}
    assert len(service.products.list(ctx)['products'])==1
    saved=service.products.save(ctx,b['id'],b['file_hash'],body('待保存的正文'))
    output=service.products.export_document(ctx,b['id'],saved['file_hash'],saved['document'],saved['title'],'docx')
    assert output['filename']=='案件甲-第二稿.docx'
    assert '待保存的正文' in Document(ctx.root_path/output['path']).paragraphs[0].text


def test_rename_collision_stale_hash_bad_name_and_permission(env):
    service,ctx,workspace=env;a=create(service,ctx);other=create(service,ctx);raw=(ctx.root_path/other['path']).read_bytes()
    b=service.products.rename(ctx,a['id'],a['file_hash'],other['title'])
    assert b['title']!=other['title'] and (ctx.root_path/other['path']).read_bytes()==raw
    with pytest.raises(ToolError):service.products.rename(ctx,a['id'],a['file_hash'],'旧版本')
    for name in ('../外部','/绝对','..','错误:名称'):
        with pytest.raises(ToolError):service.products.rename(ctx,a['id'],b['file_hash'],name)
    workspace.allow_write=False
    with pytest.raises(ToolError):service.products.rename(ctx,a['id'],b['file_hash'],'只读')


def test_rename_interruption_recovers_index_without_duplicates(env):
    service,ctx,_=env;a=create(service,ctx);publish=service.products._publish
    def fail_index(context,path,*args,**kwargs):
        if str(path).endswith('-index.json'):raise OSError('interrupted index')
        return publish(context,path,*args,**kwargs)
    with patch.object(service.products,'_publish',side_effect=fail_index):
        with pytest.raises(ToolError):service.products.rename(ctx,a['id'],a['file_hash'],'恢复名称')
    recovered=Products(service).read(ctx,a['id'])
    assert recovered['title']=='恢复名称' and recovered['document']==a['document']
    assert len(service.products.list(ctx)['products'])==1
    assert not list((ctx.root_path/'文书起草/编辑记录').glob('rename-pending-*'))


def test_grouped_directory_numbering_merge_style_and_book_order(desk):
    service,ctx,_,_=desk
    for name in ('甲','乙','丙'):(ctx.root_path/f'{name}.pdf').write_bytes(directory_pdf([name+'证据正文']))
    data={'title':'合成分组目录','groups':[{'id':'g1','name':'第六组证据','purpose':'合同成立','numbering':'restart','start_number':40},{'id':'g2','name':'第七组证据','purpose':'合同履行','numbering':'continue'}],
          'entries':[{'name':'乙','path':'乙.pdf','group':'g2'},{'name':'甲','path':'甲.pdf','group':'g1','purpose':'人工确认','manual_fields':['purpose']},{'name':'丙','path':'丙.pdf','group':'g2'}]}
    record=save(service,ctx,data,0)
    assert [row['name'] for row in record['data']['entries']]==['甲','乙','丙']
    pdf=export(service,ctx,record['revision'],'pdf');word=export(service,ctx,pdf['record']['revision'],'docx')
    assert [page.extract_text().strip() for page in PdfReader(ctx.root_path/pdf['path']).pages]==['甲证据正文','乙证据正文','丙证据正文']
    doc=Document(ctx.root_path/word['path']);table=doc.tables[0]
    assert table.rows[1].cells[0].text=='第六组证据，证明：合同成立'
    assert table.rows[1]._tr.tc_lst[0].tcPr.gridSpan.val==6
    assert [table.rows[i].cells[0].text for i in (2,4,5)]==['40','41','42']
    assert [table.rows[i].cells[-1].text for i in (2,4,5)]==['1','2','3']
    assert doc.styles['Normal'].element.rPr.rFonts.get(qn('w:eastAsia'))=='楷体'
    assert table.rows[0].cells[0]._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))=='EAEAEA'
    edited=word['record'];edited['data']['groups'][1]['numbering']='restart'
    edited=save(service,ctx,edited['data'],edited['revision'])
    assert [number for number,*_ in numbered_entries(Inventory.model_validate(edited['data']))]==[40,1,2]
    assert edited['data']['entries'][0]['purpose']=='人工确认'


def test_legacy_group_labels_survive_and_empty_groups_do_not_reset_numbers():
    value=Inventory.model_validate({'groups':[{'id':'empty','name':'空组','numbering':'restart','start_number':8}], 'entries':[{'name':'一','path':'一.pdf','group':'合同'},{'name':'二','path':'二.pdf'}]})
    assert value.groups[-1].name=='合同'
    assert [number for number,*_ in numbered_entries(value)]==[1,2]
    with pytest.raises(ValueError):Inventory.model_validate({'groups':[{'id':'a','name':'甲'},{'id':'a','name':'乙'}]})
