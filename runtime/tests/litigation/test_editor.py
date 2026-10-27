import io
from docx import Document
from lxml import etree
import pytest
from verilex_runtime.documents.docx_package import W, parts, pack, xml, serialize
from verilex_runtime.litigation.editor import inspect_docx, patch_docx, ParagraphEdit, render_docx, validate_document
from verilex_runtime.tools.registry import ToolError


def sample():
    doc=Document(); p=doc.add_paragraph(); p.add_run('重复').bold=True; p.add_run('中文正文').italic=True
    doc.add_paragraph('重复中文正文')
    table=doc.add_table(rows=1,cols=1); table.cell(0,0).text='表格正文'
    table.cell(0,0).add_table(rows=1,cols=1).cell(0,0).text='嵌套表格'
    out=io.BytesIO();doc.save(out);return out.getvalue()


def test_structural_patch_preserves_other_paragraphs_parts_and_format():
    data=sample(); before=parts(data); inspected=inspect_docx(data)
    p=inspected['paragraphs'][0]
    output=patch_docx(data,inspected['file_hash'],[ParagraphEdit(p['id'],p['text'],'重复新的正文')])
    after=parts(output)
    assert all(v==after[k] for k,v in before.items() if k!='word/document.xml')
    old=list(xml(before['word/document.xml']).iter(f'{{{W}}}p'))
    new=list(xml(after['word/document.xml']).iter(f'{{{W}}}p'))
    assert all(etree.tostring(a)==etree.tostring(b) for a,b in zip(old[1:],new[1:]))
    assert new[0].find(f'.//{{{W}}}b') is not None
    assert [p['text'] for p in inspect_docx(output)['paragraphs']][:2]==['重复新的正文','重复中文正文']
    nested=next(p for p in inspected['paragraphs'] if p['text']=='嵌套表格')
    assert '已修改' in [p['text'] for p in inspect_docx(patch_docx(data,inspected['file_hash'],[ParagraphEdit(nested['id'],nested['text'],'已修改')]))['paragraphs']]


@pytest.mark.parametrize('old,new',[('','中文'),('中文',''),('中文','中文增加'),('中文','前中文'),('中文','中追加文'),('甲乙丙','甲丁丙')])
def test_empty_insert_delete_and_middle(old,new):
    d=Document();d.add_paragraph(old);b=io.BytesIO();d.save(b);data=b.getvalue();info=inspect_docx(data);p=info['paragraphs'][0]
    result=patch_docx(data,info['file_hash'],[ParagraphEdit(p['id'],old,new)])
    assert inspect_docx(result)['paragraphs'][0]['text']==new


@pytest.mark.parametrize('tag',['drawing','footnoteReference','commentReference','fldChar'])
def test_complex_run_is_readonly_and_package_retained(tag):
    items=parts(sample());root=xml(items['word/document.xml']);p=next(root.iter(f'{{{W}}}p'))
    etree.SubElement(p.find(f'{{{W}}}r'),f'{{{W}}}{tag}')
    items['word/document.xml']=serialize(root);data=pack(items);info=inspect_docx(data)
    assert not info['paragraphs'][0]['editable']
    with pytest.raises(ToolError):patch_docx(data,info['file_hash'],[ParagraphEdit(info['paragraphs'][0]['id'],info['paragraphs'][0]['text'],'替换')])
    second=info['paragraphs'][1]
    out=patch_docx(data,info['file_hash'],[ParagraphEdit(second['id'],second['text'],'第二段')])
    assert etree.tostring(next(xml(parts(out)['word/document.xml']).iter(f'{{{W}}}p')))==etree.tostring(p)


def test_stale_and_ambiguous_edits_fail_closed():
    data=sample();info=inspect_docx(data);p=info['paragraphs'][0];e=ParagraphEdit(p['id'],p['text'],'新正文')
    with pytest.raises(ToolError):patch_docx(data,'0'*64,[e])
    with pytest.raises(ToolError):patch_docx(data,info['file_hash'],[e,e])
    with pytest.raises(ToolError):patch_docx(data,info['file_hash'],[ParagraphEdit('//w:p',p['text'],'新正文')])
    with pytest.raises(ToolError):patch_docx(data,info['file_hash'],[ParagraphEdit(p['id'],'错误原文','新正文')])


def text(s,marks=None):return {'type':'text','text':s,**({'marks':marks} if marks else {})}
def paragraph(s):return {'type':'paragraph','content':[text(s)]}


def test_structured_document_exports_chinese_formats_lists_and_tables():
    value={'type':'doc','content':[
        {'type':'heading','attrs':{'level':1,'textAlign':'center'},'content':[text('答辩状')]},
        {'type':'paragraph','attrs':{'indent':24,'lineHeight':'2'},'content':[text('中文正文',[{'type':'bold'},{'type':'underline'},{'type':'textStyle','attrs':{'fontFamily':'仿宋','fontSize':'14pt'}}])]},
        {'type':'orderedList','attrs':{'start':3},'content':[{'type':'listItem','content':[paragraph('第一项')]},{'type':'listItem','content':[paragraph('第二项')]}]},
        {'type':'table','content':[{'type':'tableRow','content':[{'type':'tableCell','content':[paragraph('证据名称')]},{'type':'tableCell','content':[paragraph('证明目的')]}]}]},
    ]}
    data=render_docx(value);doc=Document(io.BytesIO(data))
    assert [p.text for p in doc.paragraphs]==['答辩状','中文正文','第一项','第二项']
    assert doc.paragraphs[1].runs[0].bold and doc.paragraphs[1].runs[0].underline
    assert doc.paragraphs[1].runs[0].font.size.pt==14
    assert doc.tables[0].cell(0,1).text=='证明目的'
    assert doc.paragraphs[2]._p.find(f'.//{{{W}}}numId') is not None


def test_unknown_or_merged_content_cannot_silently_disappear():
    for value in [{'type':'doc','content':[{'type':'image','attrs':{'src':'x'}}]},
                  {'type':'doc','content':[{'type':'paragraph','content':[text('x',[{'type':'unknown'}])]}]},
                  {'type':'doc','content':[{'type':'table','content':[{'type':'tableRow','content':[{'type':'tableCell','attrs':{'colspan':2},'content':[paragraph('x')]}]}]}]}]:
        with pytest.raises(ToolError):validate_document(value)


def test_list_continuation_paragraph_does_not_advance_number():
    value = {'type': 'doc', 'content': [{'type': 'orderedList', 'content': [
        {'type': 'listItem', 'content': [paragraph('第一项'), paragraph('第一项说明')]},
        {'type': 'listItem', 'content': [paragraph('第二项')]}]}]}
    doc = Document(io.BytesIO(render_docx(value)))
    assert [p.text for p in doc.paragraphs] == ['第一项', '第一项说明', '第二项']
    assert [p._p.find(f'.//{{{W}}}numPr') is not None for p in doc.paragraphs] == [True, False, True]
