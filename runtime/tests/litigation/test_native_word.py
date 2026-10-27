import base64
import io
from copy import deepcopy
from docx import Document
from docx.shared import Pt
import pytest
from tests.litigation.test_workflows import env
from tests.litigation.test_editor_discussion import desk
from verilex_runtime.litigation.native_word import project
from verilex_runtime.litigation.editor import document_text
from verilex_runtime.documents.docx_package import parts
from verilex_runtime.tools.registry import ToolError

def body(*lines):return {'type':'doc','content':[{'type':'paragraph',**({'content':[{'type':'text','text':text}]} if text else {})} for text in lines]}
def open_native(service,ctx):
    p=service.products.create(ctx,'答辩状','documents',body('甲方付款','保持此段'))
    result=service.products.open_word(ctx,p['id'],p['file_hash']);return result['product'],base64.b64decode(result['content'])
def styled(data):
    doc=Document(io.BytesIO(data));doc.paragraphs[0].runs[0].font.size=Pt(18)
    doc.sections[0].header.paragraphs[0].text='合成页眉'
    doc.add_table(rows=1,cols=2).cell(0,0).text='表格原文'
    out=io.BytesIO();doc.save(out);return out.getvalue()
def test_word_bytes_save_reopen_export_preserves_package_and_conflict(env):
    s,c,_=env;p,data=open_native(s,c);raw=styled(data)
    saved=s.products.save_word(c,p['id'],p['file_hash'],base64.b64encode(raw).decode())
    assert saved['native_word'] and not saved['external_update']
    result=s.products.export_document(c,p['id'],saved['file_hash'],saved['document'],saved['title'],'docx')
    actual=parts(base64.b64decode(result['content']));original=parts(raw)
    for name in original:
        if name not in ('verilex/editor.json','[Content_Types].xml','_rels/.rels'):assert actual[name]==original[name],name
    assert result['product']['file_hash']==saved['file_hash']
    with pytest.raises(ToolError):s.products.save_word(c,p['id'],p['file_hash'],base64.b64encode(data).decode())
    assert s.products.open_word(c,p['id'],saved['file_hash'])['product']['document']==saved['document']
def test_ai_patch_preserves_unrelated_word_parts_and_table(env):
    s,c,_=env;p,data=open_native(s,c);raw=styled(data)
    saved=s.products.save_word(c,p['id'],p['file_hash'],base64.b64encode(raw).decode())
    changed=deepcopy(saved['document']);changed['content'][0]['content'][0]['text']='甲方及时付款'
    result=s.products.save(c,p['id'],saved['file_hash'],changed)
    after=(c.root_path/result['path']).read_bytes();doc=Document(io.BytesIO(after))
    assert doc.paragraphs[0].text=='甲方及时付款' and doc.paragraphs[0].runs[0].font.size.pt==18
    assert doc.tables[0].cell(0,0).text=='表格原文'
    for name,content in parts(raw).items():
        if name not in ('word/document.xml','verilex/editor.json','[Content_Types].xml','_rels/.rels'):assert parts(after)[name]==content,name
    assert result['document']==changed

def test_rewrite_paragraph_count_and_table_cell_edit(env):
    s,c,_=env;p,data=open_native(s,c)
    changed=body('第一段','第二段','第三段')
    saved=s.products.save(c,p['id'],p['file_hash'],changed)
    assert saved['document']==changed
    raw=styled((c.root_path/saved['path']).read_bytes())
    saved=s.products.save_word(c,p['id'],saved['file_hash'],base64.b64encode(raw).decode())
    changed=deepcopy(saved['document']);changed['content'][-1]['content'][0]['content'][0]['content'][0]['content'][0]['text']='新表格文字'
    result=s.products.save(c,p['id'],saved['file_hash'],changed)
    assert Document(c.root_path/result['path']).tables[0].cell(0,0).text=='新表格文字'

def test_native_proposal_accept_and_undo(desk):
    s,c,_,_=desk;p,data=open_native(s,c)
    payload={'product_id':p['id'],'base_hash':p['file_hash'],'document':p['document'],'selection':{'from':3,'to':5,'text':'付款'}}
    s.editor_discussions.bind(c,s.editor_discussions.prepare(c,payload))
    proposal=s.editor_discussions.propose(c,'及时付款')
    saved=s.editor_discussions.decide(c,proposal['id'],'accept',p['document'])
    assert document_text(saved['document'])=='甲方及时付款\n保持此段'
    undone=s.editor_discussions.decide(c,proposal['id'],'undo',saved['document'])
    assert undone['document']==p['document']
    assert parts((c.root_path/undone['path']).read_bytes())['word/document.xml']==parts(data)['word/document.xml']

def test_legacy_current_draft_is_migrated_without_losing_body(env):
    s,c,_=env;p=s.products.create(c,'答辩状','documents',body('旧文'))
    s.products.draft(c,p['id'],p['file_hash'],body('未正式保存的草稿'))
    result=s.products.open_word(c,p['id'],p['file_hash'])
    assert document_text(result['product']['document'])=='未正式保存的草稿'
    assert project(base64.b64decode(result['content']))==result['product']['document']


def test_workspace_exports_do_not_overwrite_and_keep_current_bytes(env):
    s,c,_=env;p,data=open_native(s,c)
    (c.root_path/'答辩状.docx').write_bytes(b'original file')
    first=s.products.export_document(c,p['id'],p['file_hash'],p['document'],p['title'],'docx')
    second=s.products.export_document(c,p['id'],p['file_hash'],p['document'],p['title'],'docx')
    assert first['path']=='答辩状（2）.docx' and second['path']=='答辩状（3）.docx'
    assert (c.root_path/first['path']).read_bytes()==data
    assert (c.root_path/'答辩状.docx').read_bytes()==b'original file'
    assert s.products.read(c,p['id'])['file_hash']==p['file_hash']

def test_compact_default_only_changes_old_generated_spacing(env):
    from verilex_runtime.litigation.editor import render_docx
    from verilex_runtime.litigation.native_word import compact_generated_spacing
    raw=render_docx(body('第一段','第二段'))
    adjusted=compact_generated_spacing(raw)
    assert Document(io.BytesIO(adjusted)).styles['Normal'].paragraph_format.space_after.pt==0
    for name,content in parts(raw).items():
        if name!='word/styles.xml':assert parts(adjusted)[name]==content
    doc=Document(io.BytesIO(raw));doc.styles['Normal'].paragraph_format.space_after=Pt(20)
    out=io.BytesIO();doc.save(out);assert compact_generated_spacing(out.getvalue())==out.getvalue()
    s,c,_=env;p,data=open_native(s,c)
    assert Document(io.BytesIO(data)).styles['Normal'].paragraph_format.space_after.pt==0
    assert s.products.open_word(c,p['id'],p['file_hash'])['product']['file_hash']==p['file_hash']


def test_legacy_spacing_migration_is_once_and_retains_explicit_paragraph_settings(env):
    from verilex_runtime.litigation.editor import render_docx
    from verilex_runtime.documents.docx_package import pack
    from verilex_runtime.litigation.products import META
    import json
    s,c,_=env;p=s.products.create(c,'旧答辩状','documents',body('原文','第二段'))
    path=c.root_path/p['path'];record=s.products.read(c,p['id']);record.pop('draft_spacing_version')
    old=Document(io.BytesIO(render_docx(record['document'])));old.paragraphs[0].paragraph_format.space_after=Pt(15)
    out=io.BytesIO();old.save(out);path.write_bytes(s.products._encode(record,out.getvalue()))
    before=s.products.read(c,p['id']);opened=s.products.open_word(c,p['id'],before['file_hash'])
    doc=Document(io.BytesIO(base64.b64decode(opened['content'])))
    assert doc.styles['Normal'].paragraph_format.space_after.pt==0
    assert doc.paragraphs[0].paragraph_format.space_after.pt==15
    assert doc.paragraphs[0].text=='原文'
    after=opened['product'];assert after['draft_spacing_version']==1
    assert s.products.open_word(c,p['id'],after['file_hash'])['product']['file_hash']==after['file_hash']
