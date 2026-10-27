import hashlib
import io
from dataclasses import replace
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from tests.litigation.test_workflows import env
from tests.litigation.test_workbench import desk
from verilex_runtime.litigation import cross_examination as cross
from verilex_runtime.litigation.workbench import WorkBinding
from verilex_runtime.tools.registry import ToolError
from verilex_runtime.api.litigation import build_litigation_router


def setup(service,ctx):
    (ctx.root_path/'参照.txt').write_text('对方主张已付款，我方要求核对银行流水。')
    return cross.save(service,ctx,{'title':'被告质证意见','references':['参照.txt'],'entries':[
        {'id':'a','number':'12','name':'合同','path':'合同.txt','pages':'1—2','group':'第一组','opinion':'人工意见'},
        {'id':'b','name':'收据','group':'第二组','opinion':'待核对原件'}]},0)


def binding(service,ctx,record,entry=None):
    service.workbench.bind(ctx,ctx.thread_id,WorkBinding(kind='evidence',materials=cross.material_paths(cross.Record.model_validate(record['data'])),expected_revision=0,cross_revision=record['revision'],cross_entry_id=entry))


def test_order_export_restoration_and_originals(desk):
    service,ctx,_,_=desk
    before=(ctx.root_path/'合同.txt').read_bytes()
    record=setup(service,ctx)
    record['data']['entries'].reverse()
    record=cross.save(service,ctx,record['data'],1)
    result=cross.export(service,ctx,2,'docx');doc=Document(ctx.root_path/result['path'])
    assert doc.paragraphs[0].text=='被告质证意见'
    table=doc.tables[0]
    assert [c.text for c in table.rows[0].cells]==['序号','证据名称','页码','质证意见']
    assert table.rows[1].cells[0].text=='第二组'
    assert table.rows[2].cells[1].text=='收据'
    assert table.rows[4].cells[0].text=='12'
    assert '人工意见' in table.rows[4].cells[3].text
    assert cross.export(service,ctx,2,'docx')['path']!=result['path']
    assert (ctx.root_path/'合同.txt').read_bytes()==before
    assert cross.read(service,ctx)['data']['entries'][0]['id']=='b'
    with pytest.raises(ToolError):cross.save(service,ctx,record['data'],1)
    with pytest.raises(ToolError):cross.export(service,ctx,1,'docx')


def test_suggestions_preserve_manual_text_and_are_not_exported_until_accepted(desk):
    service,ctx,_,_=desk;record=setup(service,ctx);binding(service,ctx,record)
    receipt=cross.next_entry(service,ctx)
    assert '对方主张已付款' in receipt['text'] and '双方签订合同' in receipt['text']
    assert receipt['entry']['id']=='a' and receipt['next_start'] is None
    cross.save_suggestion(service,ctx,receipt['token'],record['revision'],{'opinion':'AI建议：核对原件'})
    row=cross.read(service,ctx)['data']['entries'][0]
    assert row['opinion']=='人工意见' and row['suggestion']['opinion']=='AI建议：核对原件'
    result=cross.export(service,ctx,2,'docx')
    text=''.join(c.text for row in Document(ctx.root_path/result['path']).tables[0].rows for c in row.cells)
    assert '人工意见' in text and 'AI建议' not in text
    assert cross.next_entry(service,ctx)['entry']['id']=='b'
    result=cross.decide(service,ctx,'a',2,'accept')
    assert result['data']['entries'][0]['opinion']=='AI建议：核对原件'
    assert result['data']['entries'][0]['number']=='12'
    assert result['data']['entries'][0]['suggestion'] is None


def test_read_coverage_stale_files_stale_records_and_single_entry_scope(desk):
    service,ctx,_,_=desk;record=setup(service,ctx)
    (ctx.root_path/'合同.txt').write_text('材料文字。'*5000)
    binding(service,ctx,record,'a');receipt=cross.next_entry(service,ctx)
    receipt=cross.next_entry(service,ctx,receipt['next_start']+1)
    cross.save_suggestion(service,ctx,receipt['token'],1,{'opinion':'仅针对已读部分'})
    assert any('其余内容未由本工具读取' in gap for gap in cross.read(service,ctx)['data']['entries'][0]['gaps'])
    assert cross.next_entry(service,ctx)['done']
    with pytest.raises(ToolError):cross.save_suggestion(service,ctx,receipt['token'],1,{'opinion':'重复'})
    (ctx.root_path/'合同.txt').write_text('外部修改')
    with pytest.raises(ToolError,match='材料已变化'):cross.decide(service,ctx,'a',2,'accept')
    (ctx.root_path/'合同.txt').unlink()
    result=cross.decide(service,ctx,'a',2,'discard')
    assert result['data']['entries'][0]['opinion']=='人工意见'
    assert result['data']['entries'][1]['suggestion'] is None


def test_changed_during_read_and_unauthorized_targets(desk):
    service,ctx,workspace,threads=desk;record=setup(service,ctx);binding(service,ctx,record)
    receipt=cross.next_entry(service,ctx)
    (ctx.root_path/'参照.txt').write_text('新材料')
    with pytest.raises(ToolError,match='材料已有变化'):cross.save_suggestion(service,ctx,receipt['token'],1,{'opinion':'旧建议'})
    with pytest.raises(ToolError):cross.next_entry(service,replace(ctx,thread_id='thread-2'))
    with pytest.raises(ToolError):service.workbench.save(ctx,{'title':'非法替代','sections':[{'id':'x','title':'x','body':'x'}]},0)
    threads[ctx.thread_id].workspace_id='other'
    with pytest.raises(ToolError):cross.next_entry(service,ctx)
    workspace.allow_write=False
    with pytest.raises(ToolError):cross.save(service,ctx,record['data'],1)
    with pytest.raises(ToolError):cross.export(service,ctx,1,'docx')


def test_path_and_validation_fail_closed(desk):
    service,ctx,_,_=desk
    for data in [{'entries':[{'id':'a','name':'坏路径','path':'../秘密.txt'}]}, {'entries':[{'id':'a','name':'a'},{'id':'a','name':'b'}]}, {'title':' '}, {'entries':[{'id':'a','name':' '}]}]:
        with pytest.raises(ToolError):cross.save(service,ctx,data,0)
    (ctx.root_path/'链接.txt').symlink_to(ctx.root_path/'合同.txt')
    with pytest.raises(ToolError):cross.save(service,ctx,{'entries':[{'id':'a','name':'链接','path':'链接.txt'}]},0)
    assert cross.read(service,ctx)['revision']==0


def test_api_roundtrip_decision_and_export(desk):
    service,ctx,_,_=desk;app=FastAPI();app.include_router(build_litigation_router(service));client=TestClient(app)
    base=f'/v1/litigation/workspaces/{ctx.workspace_id}/cross-examination'
    assert client.get(base).json()['data']['entries']==[]
    response=client.put(base,json={'expected_revision':0,'data':{'entries':[{'id':'manual','name':'庭审出示证据','opinion':'核对后补充'}]}})
    assert response.status_code==200 and response.json()['revision']==1
    assert client.post(base+'/export/docx',json={'expected_revision':1}).status_code==200
    assert client.post(base+'/export/exe',json={'expected_revision':1}).status_code==422


def test_native_pdf_preserves_chinese_table(desk):
    from pathlib import Path
    if not Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').is_file():pytest.skip('LibreOffice required')
    service,ctx,_,_=desk;record=setup(service,ctx)
    result=cross.export(service,ctx,record['revision'],'pdf')
    pdf=PdfReader(ctx.root_path/result['path']);text=''.join(p.extract_text() for p in pdf.pages)
    assert '质证意见' in text and '人工意见' in text and '合同' in text
    assert all(p.mediabox.width<p.mediabox.height for p in pdf.pages)
