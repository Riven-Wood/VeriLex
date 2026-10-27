import base64
import io
from dataclasses import replace
from docx import Document
from docx.oxml.ns import qn
from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest
from tests.litigation.test_workflows import env
from tests.litigation.test_workbench import desk
from verilex_runtime.api.litigation import build_litigation_router
from verilex_runtime.documents.docx_package import parts
from verilex_runtime.tools.registry import ToolError

EMPTY={'type':'doc','content':[{'type':'paragraph'}]}

def test_custom_name_generation_native_export_uses_sample_typography(desk):
    s,c,_,_=desk
    p=s.products.create(c,'张某案-二稿.docx','documents',EMPTY,'起诉状')
    assert p['title']=='张某案-二稿' and p['document_type']=='起诉状'
    p=s.products.open_word(c,p['id'],p['file_hash'])['product']
    s.editor_discussions.bind(c,s.editor_discussions.prepare(c,{'product_id':p['id'],'base_hash':p['file_hash'],'document':p['document'],'operation':'generate'}))
    s.editor_discussions.generate(c,'民事起诉状\n原告：测试甲。\n请求支付款项。\n此致\n测试法院\n具状人：测试甲\n2026年10月26日')
    saved=s.products.read(c,p['id'])
    exported=s.products.export_document(c,p['id'],saved['file_hash'],saved['document'],saved['title'],'docx')
    doc=Document(io.BytesIO(base64.b64decode(exported['content'])))
    normal=doc.styles['Normal'];title=doc.styles['Heading 1']
    assert normal.font.color.rgb.__str__()=='000000'
    assert normal.font.size.pt==14 and normal.element.rPr.rFonts.get(qn('w:eastAsia'))=='仿宋'
    assert normal.paragraph_format.first_line_indent.pt==28
    assert normal.paragraph_format.line_spacing==1.5 and normal.paragraph_format.space_after.pt==0
    assert title.font.size.pt==18 and title.element.rPr.rFonts.get(qn('w:eastAsia'))=='黑体'
    assert title.element.rPr.color.get(qn('w:themeColor')) is None
    assert doc.paragraphs[0].style.name=='Heading 1' and doc.paragraphs[0].alignment==1
    assert doc.paragraphs[-2].alignment==2 and doc.paragraphs[-1].alignment==2
    assert doc.sections[0].page_width.twips==11906 and doc.sections[0].left_margin.twips==1474
    assert 'PAGE' in doc.sections[0].footer._element.xml
    assert (c.root_path/exported['path']).name=='张某案-二稿.docx'
    reopened=s.products.open_word(c,p['id'],saved['file_hash'])
    assert parts(base64.b64decode(reopened['content']))['word/styles.xml']==parts(base64.b64decode(exported['content']))['word/styles.xml']

def test_delete_moves_only_owned_document_and_does_not_resurrect(env):
    s,c,w=env
    p=s.products.create(c,'自定义文件','documents',EMPTY,'答辩状')
    other=s.products.create(c,'另一份','documents',EMPTY,'上诉状')
    exported=s.products.export_document(c,p['id'],p['file_hash'],EMPTY,p['title'],'docx')
    p=exported['product'];data=(c.root_path/p['path']).read_bytes()
    with pytest.raises(ToolError):s.products.delete(c,p['id'],'0'*64)
    assert (c.root_path/p['path']).read_bytes()==data
    w.allow_write=False
    with pytest.raises(ToolError):s.products.delete(c,p['id'],p['file_hash'])
    w.allow_write=True
    moved=s.products.delete(c,p['id'],p['file_hash'])
    assert not (c.root_path/p['path']).exists()
    assert (c.root_path/moved['path']).read_bytes()==data
    assert (c.root_path/exported['path']).exists()
    assert [row['id'] for row in s.products.list(c,'documents')['products']]==[other['id']]
    with pytest.raises(ToolError):s.products.read(c,p['id'])
    with pytest.raises(ToolError):s.products.save(c,p['id'],p['file_hash'],EMPTY)
    recreated=s.products.create(c,'自定义文件','documents',EMPTY,'答辩状')
    assert recreated['id']!=p['id']
    with pytest.raises(ToolError):s.products.delete(c,p['id'],p['file_hash'])
    assert (c.root_path/recreated['path']).exists()

def test_api_custom_name_and_delete(env):
    s,c,_=env;app=FastAPI();app.include_router(build_litigation_router(s));client=TestClient(app)
    base=f'/v1/litigation/workspaces/{c.workspace_id}/editor-products'
    created=client.post(base,json={'title':'自定义名称','kind':'documents','document':EMPTY,'document_type':'上诉状'})
    assert created.status_code==200
    p=created.json();assert p['document_type']=='上诉状'
    result=client.request('DELETE',base+'/'+p['id'],json={'base_hash':p['file_hash']})
    assert result.status_code==200 and '已删除文书' in result.json()['path']
    assert client.get(base).json()['products']==[]
