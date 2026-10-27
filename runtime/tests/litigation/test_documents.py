"""Synthetic ordinary-form pipeline, publication and hard-stop regressions."""
from copy import deepcopy
from dataclasses import replace
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
import json
import pytest
from docx import Document
from lxml import etree
from tests.litigation.test_workflows import env, NOW, make_database, repository, law_env
from verilex_runtime.litigation.documents import DocumentProposal
from verilex_runtime.litigation.document_templates import catalogue
from verilex_runtime.litigation.document_render import render_document, validate_package
from verilex_runtime.litigation.exporter import LitigationExporter
from verilex_runtime.litigation.service import LitigationService
from verilex_runtime.artifacts.publisher import ArtifactPublisher
from verilex_runtime.artifacts.blob_store import ArtifactBlobStore
from verilex_runtime.tools.registry import ToolError, ToolRegistry
from verilex_runtime.litigation.tools import register_litigation_tools


def values(template_id='civil-person',long=False):
    definition=next(t for t in catalogue() if t['id']==template_id)
    special={'plaintiff_name':'虚构原告甲','defendant':'虚构被告乙，住虚构市测试路2号', 'sex':'男','birth_date':'1990-01-02','date':'2026-09-09','copies':'1','measure':'冻结','court':'虚构市人民法院','legal_agent':'无','agent':'无','reasons':('本案合成事实，金额100元。\n' * 300 if long else '根据合成合同请求付款100元。')}
    return {'template_id':template_id,'plaintiff_count':1,'defendant_count':1,'plaintiff_type':definition.get('party_scope',{}).get('plaintiff_type','natural'),'defendant_type':definition.get('party_scope',{}).get('defendant_type','natural'),**{k:v for k,v in definition.get('party_scope',{}).items() if k.startswith('third_party_')},'fields':[{'key':f['key'],'value':special.get(f['key'], f.get('choices', ['1990-01-02' if f['kind']=='date' else '男' if f['kind']=='sex' else '合成测试内容'])[0]),'sources':[],'pending':[]} for f in definition['fields']],'laws':[],'summary':'合成测试文书'}


def confirmed(service,ctx,template_id='civil-person'):
    payload=values(template_id)
    saved=service.save(ctx,'documents',payload,0)
    assert not saved['data']['confirmed']
    assert all(f['pending']==(['请当事人选择'] if f['key'].startswith('mediation_') else ['来源待核对']) for f in saved['data']['fields'])
    return service.confirm_document(ctx,payload,1)


def exporter_for(tmp_path,ctx):
    db=make_database(tmp_path,workspace_root=ctx.root_path)
    with db.connect() as connection:
        connection.execute("INSERT INTO tool_calls (id,turn_id,provider_call_id,name,arguments_json,risk_level,status,started_at) VALUES ('tool-2','turn-1','provider-2','litigation_export','{}','high','running',?)",(NOW,))
    return LitigationExporter(ArtifactPublisher(repository(db),ArtifactBlobStore(tmp_path/'artifacts'),now=lambda:NOW),temporary_root=tmp_path/'stage',now=lambda:NOW)


@pytest.mark.parametrize('template_id',[t['id'] for t in catalogue()])
def test_complete_human_pipeline_publishes_docx_and_restarts(env,tmp_path,template_id):
    service,ctx,_=env
    saved=confirmed(service,ctx,template_id)
    assert LitigationService(service.workspaces).read(ctx,'documents')==saved
    exporter=exporter_for(tmp_path,ctx)
    records=exporter.export(service,ctx,'documents')
    assert len(records)==2
    doc=Document(ctx.root_path/records[0].published_relative_path)
    text='\n'.join(doc.element.xpath('.//w:t/text()'))
    assert '虚构原告甲' in text and '2026年9月9日' in text
    assert '{{' not in text and 'AI 草稿，须律师审查' in text
    before=(ctx.root_path/records[0].published_relative_path).read_bytes()
    again=exporter.export(service,replace(ctx,tool_call_id='tool-2'),'documents')
    assert records[0].published_relative_path!=again[0].published_relative_path
    assert (ctx.root_path/records[0].published_relative_path).read_bytes()==before
    report=(ctx.root_path/records[1].published_relative_path).read_text()
    assert 'court.gov.cn' in report
    assert '已检查必填字段和文件结构' in report


@pytest.mark.parametrize('change',[{'confirmed':True},{'reviewed_fields':['plaintiff_name']},{'dependencies':{}},{'plaintiff_count':2},{'defendant_count':3}])
def test_model_cannot_confer_review_or_expand_parties(env,change):
    service,ctx,_=env
    with pytest.raises(ToolError):service.save(ctx,'documents',{**values(),**change},0)
    assert service.read(ctx,'documents')['revision']==0


def test_export_requires_review_and_rechecks_permission(env,tmp_path):
    service,ctx,workspace=env
    service.save(ctx,'documents',values(),0)
    exporter=exporter_for(tmp_path,ctx)
    with pytest.raises(ToolError,match='逐项核对'):exporter.export(service,ctx,'documents')
    service.confirm_document(ctx,values(),1)
    workspace.allow_write=False
    with pytest.raises(ToolError):exporter.export(service,ctx,'documents')


@pytest.mark.parametrize('bad',['','{{old_case}}','×××','待确认','2026-02-30'])
def test_missing_placeholder_and_bad_date_block_confirmation(env,bad):
    service,ctx,_=env
    service.save(ctx,'documents',values(),0)
    data=values()
    next(f for f in data['fields'] if f['key']=='date')['value']=bad
    with pytest.raises(ToolError):service.confirm_document(ctx,data,1)
    assert service.read(ctx,'documents')['revision']==1


def test_pending_cannot_be_erased_by_export(env,tmp_path):
    service,ctx,_=env
    saved=service.save(ctx,'documents',values(),0)
    proposal={k:saved['data'][k] for k in DocumentProposal.model_fields}
    with pytest.raises(ToolError): service.confirm_document(ctx,proposal,1)


def test_upstream_and_stale_confirmation_fail_without_overwrite(env,tmp_path):
    service,ctx,_=env
    confirmed(service,ctx)
    with pytest.raises(ToolError):service.confirm_document(ctx,values(),1)
    service.save_case(ctx,{'title':'案件修改'},0)
    with pytest.raises(ToolError,match='案件资料已更新'):exporter_for(tmp_path,ctx).export(service,ctx,'documents')


def test_changed_source_is_blocked_at_export(env,tmp_path):
    service,ctx,_=env
    payload=values();payload['fields'][0]['sources']=[{'path':'合同.txt','quote':'应支付货款100元。'}]
    service.save(ctx,'documents',payload,0)
    service.confirm_document(ctx,payload,1)
    (ctx.root_path/'合同.txt').write_text('材料已改变')
    with pytest.raises(ToolError,match='原文不一致'):exporter_for(tmp_path,ctx).export(service,ctx,'documents')


@pytest.mark.parametrize('template_id',[t['id'] for t in catalogue()])
def test_long_text_and_multiline_survive_without_fixed_row_clipping(tmp_path,template_id):
    payload=values(template_id,long=True);path=tmp_path/'long.docx'
    assert render_document(path,DocumentProposal.model_validate(payload))=={'residue':'passed','required_fields':'passed','format':'passed'}
    text='\n'.join(Document(path).element.xpath('.//w:t/text()'))
    assert text.count('本案合成事实，金额100元。')==300
    if template_id.startswith('private-lending-'):
        assert len(Document(path).tables)==6
    else:
        assert not Document(path).tables


@pytest.mark.parametrize('failure',['residue','format','extra_part','missing_part'])
def test_final_package_corruption_stops_before_publication(env,tmp_path,failure):
    service,ctx,_=env;confirmed(service,ctx)
    exporter=exporter_for(tmp_path,ctx)
    def corrupt(path,parts,tree):
        with ZipFile(path) as archive:changed={n:archive.read(n) for n in archive.namelist()}
        if failure=='residue':changed['word/document.xml']=changed['word/document.xml'].replace('虚构原告甲'.encode(),'上次案件某公司'.encode())
        elif failure=='format':changed['word/document.xml']=changed['word/document.xml'].replace(b'w:sz w:val="40"',b'w:sz w:val="90"')
        elif failure=='extra_part':changed['word/old-case.xml']=b'old case'
        else:del changed['word/styles.xml']
        with ZipFile(path,'w',ZIP_DEFLATED) as archive:
            for n,b in changed.items():archive.writestr(n,b)
        return validate_package(path,parts,tree)
    with patch('verilex_runtime.litigation.document_render.validate_package',corrupt):
        with pytest.raises(ToolError):exporter.export(service,ctx,'documents')
    assert not list((ctx.root_path/'文书').glob('*.docx'))
    assert not list((ctx.root_path/'文书').glob('*.md'))


def test_template_tamper_and_missing_keys_are_not_accepted(env,tmp_path):
    service,ctx,_=env
    data=values();data['fields'].pop()
    with pytest.raises(ToolError):service.save(ctx,'documents',data,0)
    with patch('verilex_runtime.litigation.document_templates.catalogue',return_value=[{**catalogue()[0],'sha256':'bad'}]):
        with pytest.raises(ToolError,match='样式校验失败'):service.save(ctx,'documents',values(),0)


def test_numbered_laws_require_source_read_and_export_current_thread(env,tmp_path):
    service,ctx,_=env
    data=values();next(f for f in data['fields'] if f['key']=='reasons')['value']='依据第一条提出请求'
    saved=service.save(ctx,'documents',data,0)
    assert saved['revision']==1
    law=law_env(service)['issues'][0]['laws'][0];data['laws']=[law]
    service.save(ctx,'documents',data,1)
    # HTTP review context has no thread and may not invent law verification.
    service.confirm_document(replace(ctx,thread_id=None),data,2)
    with pytest.raises(ToolError):exporter_for(tmp_path,ctx).export(service,replace(ctx,thread_id='other'),'documents')


@pytest.mark.asyncio
async def test_agent_read_exposes_only_real_template_schema(env):
    service,ctx,_=env;registry=ToolRegistry();register_litigation_tools(registry,service)
    result=await registry.get('litigation_read').execute(ctx,{'kind':'documents'})
    read=json.loads(result.output)
    assert len(read['templates'])==15
    assert 'confirmed' not in read['schema']['properties']


def test_authenticated_api_reviews_and_reloads_real_record(env,tmp_path):
    from fastapi.testclient import TestClient
    from verilex_runtime.app import create_app
    from verilex_runtime.container import RuntimeServices
    from verilex_runtime.settings import RuntimeSettings
    from verilex_runtime.providers.service import ProviderService
    from verilex_runtime.storage.provider_repository import ProviderRepository
    service,ctx,workspace=env;db=make_database(tmp_path,workspace_root=ctx.root_path)
    app=create_app(RuntimeSettings(host='127.0.0.1',port=0,state_dir=tmp_path,api_token='synthetic-docs-test'),RuntimeServices(provider_service=ProviderService(ProviderRepository(db),object(),object()),litigation_service=service))
    client=TestClient(app,base_url='http://127.0.0.1',raise_server_exceptions=False)
    base='/v1/litigation/workspaces/workspace-1';headers={'X-VeriLex-Token':'synthetic-docs-test'}
    service.save(ctx,'documents',values(),0)
    assert client.get(base+'/document-templates').status_code==401
    assert len(client.get(base+'/document-templates',headers=headers).json()['templates'])==15
    assert client.put(base+'/documents/confirmation',json={'data':values(),'expected_revision':1}).status_code==401
    saved=client.put(base+'/documents/confirmation',headers=headers,json={'data':values(),'expected_revision':1})
    assert saved.status_code==200 and saved.json()['data']['confirmed']
    assert client.get(base,headers=headers).json()['documents']['revision']==2
    assert client.put(base+'/documents/confirmation',headers=headers,json={'data':values(),'expected_revision':1}).status_code==409
    assert client.put(base+'/documents/confirmation',headers=headers,json={'data':{**values(),'confirmed':True},'expected_revision':2}).status_code==422
    workspace.allow_write=False
    assert client.put(base+'/documents/confirmation',headers=headers,json={'data':values(),'expected_revision':2}).status_code>=400


def lending_sample():
    from pathlib import Path
    sample=json.loads((Path(__file__).parent/'fixtures/private-lending-person-2025.json').read_text())
    return {'template_id':'private-lending-person-2025','plaintiff_count':1,'defendant_count':1,
            'plaintiff_type':'natural','defendant_type':'natural','third_party_count':0,
            'fields':[{'key':key,'value':value,'sources':[],'pending':[]} for key,value in sample.items()],
            'laws':[],'summary':'自然人无息借款合成金样'}


@pytest.mark.parametrize('change',[{'plaintiff_type':'organization'},{'defendant_type':'organization'},
                                  {'plaintiff_type':None},{'third_party_count':1}])
def test_lending_rejects_unsupported_party_scope(env,change):
    service,ctx,_=env
    with pytest.raises(ToolError):service.save(ctx,'documents',{**lending_sample(),**change},0)
    assert service.read(ctx,'documents')['revision']==0


def test_lending_model_cannot_choose_mediation_but_human_can(env,tmp_path):
    service,ctx,_=env;data=lending_sample()
    saved=service.save(ctx,'documents',data,0)
    wishes=[f for f in saved['data']['fields'] if f['key'].startswith('mediation_')]
    assert len(wishes)==7 and all(f['value']=='' and f['pending']==['请当事人选择'] for f in wishes)
    invalid=deepcopy(data)
    next(f for f in invalid['fields'] if f['key']=='mediation_choice')['value']='模型决定调解'
    with pytest.raises(ToolError):service.confirm_document(ctx,invalid,1)
    service.confirm_document(ctx,data,1)
    records=exporter_for(tmp_path,ctx).export(service,ctx,'documents')
    doc=Document(ctx.root_path/records[0].published_relative_path)
    text='\n'.join(doc.element.xpath('.//w:t/text()'))
    assert '暂不确定，想要了解更多内容' in text
    assert '虚构被告乙' in text and '100元人民币' in text
    assert '民事答辩状' not in text and '了解□' not in text
    assert '{{' not in text


@pytest.mark.parametrize('template_id',[t['id'] for t in catalogue() if t['id'].startswith('private-lending-')])
def test_lending_preserves_official_column_labels_and_flowing_rows(tmp_path,template_id):
    from pathlib import Path
    from verilex_runtime.litigation.document_templates import template
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    source=etree.parse(str(Path(__file__).resolve().parents[3]/'scripts/litigation/sources/private-lending-2025.xml'))
    _,content=template(template_id)
    built=etree.fromstring(content)
    original_tables=source.findall('w:tbl',ns)
    result_tables=built.findall('.//w:tbl',ns)
    assert len(original_tables)==len(result_tables)==6
    for before,after in zip(original_tables,result_tables):
        assert etree.tostring(before.find('w:tblGrid',ns),method='c14n',exclusive=True)==etree.tostring(after.find('w:tblGrid',ns),method='c14n',exclusive=True)
        for old,new in zip(before.findall('w:tr',ns),after.findall('w:tr',ns)):
            if len(old.findall('w:tc',ns))>1:
                assert old.find('w:tc',ns).xpath('.//w:t/text()',namespaces=ns)==new.find('w:tc',ns).xpath('.//w:t/text()',namespaces=ns)
    assert all(h.get('{'+ns['w']+'}hRule')!='exact' for h in built.findall('.//w:trHeight',ns))
    assert not built.findall('.//w:cantSplit',ns)
    assert not built.findall('.//w:pStyle',ns)
    assert not built.xpath('.//@r:id',namespaces={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})


def test_legacy_document_without_new_party_fields_still_exports(env,tmp_path):
    service,ctx,_=env
    saved=confirmed(service,ctx)
    for name in ('plaintiff_type','defendant_type','third_party_count'):saved['data'].pop(name)
    from verilex_runtime.litigation.documents import ready
    assert ready(service,ctx,saved['data']).template_id=='civil-person'


ORGANIZATION_TEMPLATES = [t['id'] for t in catalogue() if t.get('party_scope') and not t.get('third_parties')]


@pytest.mark.parametrize('template_id', ORGANIZATION_TEMPLATES)
@pytest.mark.parametrize('role', ['plaintiff_type', 'defendant_type'])
def test_organization_scope_mismatch_rejected_at_every_stage(env,tmp_path,template_id,role):
    from verilex_runtime.litigation.documents import ready
    service,ctx,_=env
    data=values(template_id)
    bad=deepcopy(data)
    bad[role]='natural' if data[role]=='organization' else 'organization'
    with pytest.raises(ToolError,match='类型'):
        service.save(ctx,'documents',bad,0)
    assert service.read(ctx,'documents')['revision']==0
    service.save(ctx,'documents',data,0)
    with pytest.raises(ToolError,match='类型'):
        service.confirm_document(ctx,bad,1)
    saved=service.confirm_document(ctx,data,1)
    saved['data'][role]=bad[role]
    with pytest.raises(ToolError,match='类型'):
        ready(service,ctx,saved['data'])


@pytest.mark.parametrize('template_id', ORGANIZATION_TEMPLATES)
def test_organization_gold_sample_fills_only_matching_official_party_rows(env,tmp_path,template_id):
    from pathlib import Path
    service,ctx,_=env
    data=values(template_id)
    sample=json.loads((Path(__file__).parent/'fixtures'/f'{template_id}.json').read_text())
    for field in data['fields']:field['value']=sample[field['key']]
    saved=service.save(ctx,'documents',data,0)
    assert all(f['value']=='' for f in saved['data']['fields'] if f['key'].startswith('mediation_'))
    service.confirm_document(ctx,data,1)
    records=exporter_for(tmp_path,ctx).export(service,ctx,'documents')
    doc=Document(ctx.root_path/records[0].published_relative_path)
    for role,t,natural_row,org_row in [('plaintiff',0,2,3),('defendant',1,1,2)]:
        rows=doc.tables[t].rows
        if data[role+'_type']=='organization':
            assert rows[natural_row].cells[1].text=='无（不适用）'
            assert sample[role+'_name'] in rows[org_row].cells[1].text
            assert sample[role+'_representative'] in rows[org_row+1].cells[1].text
            assert sample[role+'_phone'] in rows[org_row+1].cells[-1].text
            assert sample[role+'_ownership'] in rows[org_row+2].cells[1].text
        else:
            assert sample[role+'_name'] in rows[natural_row].cells[1].text
            assert rows[org_row].cells[1].text=='无（不适用）'
    text='\n'.join(doc.element.xpath('.//w:t/text()'))
    assert '民事答辩状' not in text and '{{' not in text


@pytest.mark.parametrize('template_id', ORGANIZATION_TEMPLATES)
@pytest.mark.parametrize('change', [{'plaintiff_type':None},{'defendant_type':None},{'third_party_count':1},{'plaintiff_count':2},{'defendant_count':2}])
def test_organization_scope_remains_single_and_explicit(env,template_id,change):
    service,ctx,_=env
    with pytest.raises(ToolError):service.save(ctx,'documents',{**values(template_id),**change},0)
    assert service.read(ctx,'documents')['revision']==0


@pytest.mark.asyncio
async def test_large_template_catalogue_reads_only_selected_fields(env):
    from verilex_runtime.litigation.document_templates import template
    service,ctx,_=env;registry=ToolRegistry();register_litigation_tools(registry,service)
    read=registry.get('litigation_read').execute
    result=json.loads((await read(ctx,{'kind':'documents'})).output)
    assert all('fields' not in t for t in result['templates'])
    assert 'template' not in result
    selected='private-lending-person-third-natural-2025'
    result=json.loads((await read(ctx,{'kind':'documents','template_id':selected})).output)
    assert result['template']==template(selected)[0]
    assert result['template']['party_scope']['third_party_count']==1
    with pytest.raises(ToolError):await read(ctx,{'kind':'documents','template_id':'not-a-template'})
    with pytest.raises(ToolError):await read(ctx,{'kind':'facts','template_id':selected})
    # A much larger catalogue must not repeat all of its field arrays in the tool output.
    many=[{**catalogue()[0],'id':f'form-{i}'} for i in range(200)]
    with patch('verilex_runtime.litigation.document_templates.catalogue',return_value=many):
        result=await read(ctx,{'kind':'documents'})
        assert len(result.output)<120000


THIRD_PARTY_TEMPLATES=[t['id'] for t in catalogue() if t.get('third_parties')]


@pytest.mark.parametrize('template_id',THIRD_PARTY_TEMPLATES)
def test_third_party_requires_exact_count_type_and_fills_official_row(env,tmp_path,template_id):
    service,ctx,_=env;data=values(template_id)
    for change in [{'third_party_count':0},{'third_party_count':2},{'third_party_type':None},
                   {'third_party_type':'natural' if data['third_party_type']=='organization' else 'organization'}]:
        with pytest.raises(ToolError):service.save(ctx,'documents',{**data,**change},0)
    for f in data['fields']:
        if f['key']=='third_party_name':f['value']='虚构第三人丙'
    service.save(ctx,'documents',data,0)
    service.confirm_document(ctx,data,1)
    records=exporter_for(tmp_path,ctx).export(service,ctx,'documents')
    doc=Document(ctx.root_path/records[0].published_relative_path)
    natural=doc.tables[1].rows[5].cells[1].text
    organization=doc.tables[2].rows[0].cells[1].text
    assert '虚构第三人丙' in (natural if data['third_party_type']=='natural' else organization)
    assert (organization if data['third_party_type']=='natural' else natural)=='无（不适用）'


def test_model_field_order_is_normalized_to_the_official_review_order(env):
    service,ctx,_=env
    data=values('private-lending-person-third-natural-2025')
    data['fields'].reverse()
    saved=service.save(ctx,'documents',data,0)
    expected=next(t for t in catalogue() if t['id']==data['template_id'])['fields']
    assert [f['key'] for f in saved['data']['fields']]==[f['key'] for f in expected]
