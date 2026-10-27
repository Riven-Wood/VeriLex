from __future__ import annotations
import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import pytest
from docx import Document
from verilex_runtime.litigation.models import Fact, merge_facts
from verilex_runtime.litigation.service import LitigationService, OWNER
from verilex_runtime.litigation.exporter import LitigationExporter
from verilex_runtime.litigation.tools import register_litigation_tools
from verilex_runtime.tools.registry import ToolContext, ToolError, ToolRegistry
from verilex_runtime.storage.models import ToolCallStatus
from verilex_runtime.artifacts.blob_store import ArtifactBlobStore
from verilex_runtime.artifacts.publisher import ArtifactPublisher
from tests.artifacts._support import NOW, make_database, repository

@pytest.fixture
def env(tmp_path):
    root=tmp_path/'workspace'; root.mkdir()
    (root/'合同.txt').write_text('2026年1月2日双方签订合同。应支付货款100元。',encoding='utf8')
    workspace=SimpleNamespace(id='workspace-1',root_path=str(root),allow_read=True,allow_write=True)
    internal={}
    workspaces=SimpleNamespace(get=lambda _:workspace,
        register_internal_file=lambda wid,path,digest:internal.__setitem__((wid,path),digest),
        internal_file_digest=lambda wid,path:internal.get((wid,path)))
    service=LitigationService(workspaces)
    context=ToolContext(root,True,True,thread_id='thread-1',turn_id='turn-1',tool_call_id='tool-1',workspace_id='workspace-1')
    return service,context,workspace


def fact_proposal(**changes):
    fact={'id':'f1','date':'2026-01-02','text':'双方签订合同。','sources':[{'path':'合同.txt','quote':'2026年1月2日双方签订合同。'}]}
    fact.update(changes)
    return {'facts':[fact],'gaps':['付款材料缺失'],'relationships':'合同关系，须律师复核','summary':'新增合同事实'}


def evidence_proposal():
    return {'evidence':[{'id':'e1','group':'合同成立','name':'合同','kind':'书证','purpose':'证明合同成立','sources':[{'path':'合同.txt','quote':'双方签订合同。'}],'assessment':'内部专用意见，不应进入递交目录'}], 'gaps':['付款记录'],'contradictions':[],'suggestions':['补充付款凭证'],'summary':'整理证据'}


def law_env(service):
    text='第一条\n法律原文。'
    service.law_cache=SimpleNamespace(lookup=lambda _:SimpleNamespace(body_text=text))
    record=SimpleNamespace(name='law_read',status=ToolCallStatus.SUCCEEDED, arguments_json=json.dumps({'bbbs':'law1','clause':'第一条'}), result_json=json.dumps({'ok':True,'output':json.dumps({'bbbs':'law1','clause':'第一条','clause_text':['第一条','法律原文。'],'body_text':text})}))
    service.law_reads=SimpleNamespace(list_by_thread=lambda thread:[record] if thread=='thread-1' else [])
    return {'issues':[{'title':'合同是否成立','elements':[{'name':'合意','supporting':[{'path':'合同.txt','quote':'双方签订合同。'}]}],'laws':[{'bbbs':'law1','clause':'第一条','quote':'法律原文。'}]}],'summary':'核对要件'}


def test_incremental_conflict_keeps_confirmed_fact(env):
    service,ctx,_=env
    saved=service.save(ctx,'facts',fact_proposal(),0)
    item=saved['data']['facts'][0]
    reviewed=service.review_fact(ctx,'f1',{**item,'user_note':'我已核对原件'},1)
    result=service.save(ctx,'facts',fact_proposal(text='另一个事件。'),2)
    assert result['data']['facts'][0]==reviewed['data']['facts'][0]
    assert len(result['data']['conflicts'])==1
    assert result['revision']==3
    assert len(list(ctx.root_path.glob('事实/.verilex-revision-*.json')))==2

@pytest.mark.parametrize('source',[{'path':'合同.txt','quote':'不存在的事实'},{'path':'../外部.txt','quote':'x'},{'path':'合同.txt','page':1,'quote':'双方签订合同。'}])
def test_invalid_source_is_not_saved(env,source):
    service,ctx,_=env
    with pytest.raises(ToolError): service.save(ctx,'facts',fact_proposal(sources=[source]),0)
    assert not (ctx.root_path/'事实').exists()

@pytest.mark.parametrize('changes',[{'status':'confirmed'},{'user_note':'伪造用户确认'}])
def test_model_cannot_confirm_facts(env,changes):
    service,ctx,_=env
    with pytest.raises(ToolError,match='只有用户核对'): service.save(ctx,'facts',fact_proposal(**changes),0)


def test_missing_date_and_source_remain_pending(env):
    service,ctx,_=env
    item=service.save(ctx,'facts',fact_proposal(date='',sources=[]),0)['data']['facts'][0]
    assert item['status']=='pending'
    assert set(item['pending'])=={'日期待确认','事实来源待确认'}


def test_permissions_rechecked_after_revocation(env):
    service,ctx,workspace=env
    workspace.allow_write=False
    with pytest.raises(ToolError): service.save(ctx,'facts',fact_proposal(),0)


def test_stale_revision_does_not_replace_saved_facts(env):
    service,ctx,_=env
    service.save(ctx,'facts',fact_proposal(),0)
    with pytest.raises(ToolError,match='案件已有更新'): service.save(ctx,'facts',fact_proposal(),0)
    assert service.read(ctx,'facts')['revision']==1


def test_foreign_file_is_never_replaced(env):
    service,ctx,_=env
    (ctx.root_path/'案件信息.json').write_text('用户已有文件')
    with pytest.raises(ToolError): service.save_case(ctx,{},0)
    assert (ctx.root_path/'案件信息.json').read_text()=='用户已有文件'


def test_concurrent_foreign_file_wins_first_creation(env):
    service,ctx,_=env
    import tempfile
    original=tempfile.mkstemp
    def competing_writer(*args,**kwargs):
        result=original(*args,**kwargs)
        (ctx.root_path/'案件信息.json').write_text('用户原件')
        return result
    with patch('verilex_runtime.litigation.service.tempfile.mkstemp',competing_writer):
        with pytest.raises(ToolError): service.save_case(ctx,{},0)
    assert (ctx.root_path/'案件信息.json').read_text()=='用户原件'


def test_user_edit_during_update_is_retained(env):
    service,ctx,_=env
    service.save_case(ctx,{'title':'初始'},0)
    import tempfile
    original=tempfile.mkstemp
    def competing_writer(*args,**kwargs):
        result=original(*args,**kwargs)
        (ctx.root_path/'案件信息.json').write_text('用户改稿')
        return result
    with patch('verilex_runtime.litigation.service.tempfile.mkstemp',competing_writer):
        with pytest.raises(ToolError): service.save_case(ctx,{'title':'新稿'},1)
    assert (ctx.root_path/'案件信息.json').read_text()=='用户改稿'


def test_material_listing_skips_outputs_hidden_symlinks(env,tmp_path):
    service,ctx,_=env
    (ctx.root_path/'.private').mkdir(); (ctx.root_path/'.private'/'x.txt').write_text('no')
    (ctx.root_path/'link.txt').symlink_to(tmp_path/'outside')
    (ctx.root_path/'扫描.png').write_bytes(b'not-real-image')
    service.save(ctx,'facts',fact_proposal(),0)
    items=service.materials(ctx)['materials']
    assert {i['path'] for i in items}=={'合同.txt','扫描.png'}
    assert next(i for i in items if i['path']=='扫描.png')['supported'] is False


def test_evidence_order_retains_all_rows(env):
    service,ctx,_=env
    service.save(ctx,'evidence',evidence_proposal(),0)
    updated=service.arrange_evidence(ctx,[{'id':'e1','group':'更正分组'}],1)
    assert updated['data']['evidence'][0]['group']=='更正分组'
    with pytest.raises(ToolError):service.arrange_evidence(ctx,[],2)


def test_law_requires_successful_read_in_same_thread(env):
    service,ctx,_=env
    proposal=law_env(service)
    with pytest.raises(ToolError,match='当前案件对话'): save_issues(service,replace(ctx,thread_id='other'),proposal,0)
    result=save_issues(service,ctx,proposal,0)
    assert result['data']['law_verification']=='verified_against_thread_read'


def test_law_quote_must_match_cached_and_read_text(env):
    service,ctx,_=env
    proposal=law_env(service);proposal['issues'][0]['laws'][0]['quote']='虚构条文'
    with pytest.raises(ToolError):save_issues(service,ctx,proposal,0)


def test_export_publishes_real_evidence_docx_and_separate_assessment(env,tmp_path):
    service,ctx,_=env
    service.save(ctx,'facts',fact_proposal(),0)
    service.save(ctx,'evidence',evidence_proposal(),0)
    db=make_database(tmp_path,workspace_root=ctx.root_path)
    publisher=ArtifactPublisher(repository(db),ArtifactBlobStore(tmp_path/'artifacts'),now=lambda:NOW)
    exporter=LitigationExporter(publisher,temporary_root=tmp_path/'staging',now=lambda:NOW)
    records=exporter.export(service,ctx,'evidence')
    assert len(records)==2
    docx=next(r for r in records if r.display_name.endswith('.docx'))
    doc=Document(ctx.root_path/docx.published_relative_path)
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert '证明合同成立' in text and '内部专用意见' not in text
    assert 'AI 草稿，须律师审查' in text
    assert '内部专用意见' in (ctx.root_path/'证据/证据评估.md').read_text()
    assert len(repository(db).list_page('thread-1',limit=10,cursor=None).items)==2


def test_export_rejects_changed_source_before_publication(env,tmp_path):
    service,ctx,_=env
    service.save(ctx,'evidence',evidence_proposal(),0)
    (ctx.root_path/'合同.txt').write_text('材料已被替换')
    publisher=SimpleNamespace(publish_batch=lambda _:pytest.fail('must not publish'))
    exporter=LitigationExporter(publisher,temporary_root=tmp_path/'staging',now=lambda:NOW)
    with pytest.raises(ToolError):exporter.export(service,ctx,'evidence')


def test_export_rejects_tampered_law(env,tmp_path):
    service,ctx,_=env
    save_issues(service,ctx,law_env(service),0)
    path=ctx.root_path/'焦点与策略/issues.json'
    content=json.loads(path.read_text());content['data']['issues'][0]['laws'][0]['quote']='不存在的条文';path.write_text(json.dumps(content))
    exporter=LitigationExporter(SimpleNamespace(publish_batch=lambda _:pytest.fail('must not publish')),temporary_root=tmp_path/'staging',now=lambda:NOW)
    with pytest.raises(ToolError):exporter.export(service,ctx,'issues')


def test_case_theory_change_invalidates_strategy(env,tmp_path):
    service,ctx,_=env
    service.save_case(ctx,{'theory':'解除合同'},0)
    save_issues(service,ctx,law_env(service),0)
    strategy={'favorable':['合同'],'unfavorable':['欠缺付款'],'defenses':['对方否认'],'actions':['补证'],'opposing_counsel':['攻击付款'],'summary':'分析策略'}
    service.save(ctx,'strategy',strategy,0)
    service.save_case(ctx,{'theory':'继续履行'},1)
    exporter=LitigationExporter(SimpleNamespace(publish_batch=lambda _:pytest.fail('must not publish')),temporary_root=tmp_path/'staging',now=lambda:NOW)
    with pytest.raises(ToolError,match='已更新'): exporter.export(service,ctx,'strategy')

@pytest.mark.asyncio
async def test_tool_schema_and_approval_metadata(env):
    service,ctx,_=env
    registry=ToolRegistry();register_litigation_tools(registry,service)
    assert registry.get('litigation_save').risk_level.value=='high'
    assert registry.get('litigation_save').approval_behavior.value=='risk_default'
    result=await registry.get('litigation_read').execute(ctx,{'kind':'facts'})
    assert 'schema' in json.loads(result.output)
    result=await registry.get('litigation_save').execute(ctx,{'kind':'facts','payload_json':json.dumps(fact_proposal()),'expected_revision':0})
    assert json.loads(result.output)['revision']==1


def test_authenticated_api_save_reopen_review_and_conflict(env,tmp_path):
    from fastapi.testclient import TestClient
    from verilex_runtime.app import create_app
    from verilex_runtime.container import RuntimeServices
    from verilex_runtime.settings import RuntimeSettings
    from verilex_runtime.providers.service import ProviderService
    from verilex_runtime.storage.provider_repository import ProviderRepository
    service,ctx,_=env
    db=make_database(tmp_path,workspace_root=ctx.root_path)
    app=create_app(RuntimeSettings(host='127.0.0.1',port=0,state_dir=tmp_path,api_token='litigation-test'), RuntimeServices(provider_service=ProviderService(ProviderRepository(db),object(),object()),litigation_service=service))
    client=TestClient(app,base_url='http://127.0.0.1',raise_server_exceptions=False)
    endpoint='/v1/litigation/workspaces/workspace-1'
    headers={'X-VeriLex-Token':'litigation-test'}
    assert client.get(endpoint).status_code==401
    saved=client.put(endpoint,headers=headers,json={'expected_revision':0,'data':{'title':'合成案件'}})
    assert saved.status_code==200
    assert client.get(endpoint,headers=headers).json()['case']['data']['title']=='合成案件'
    conflict=client.put(endpoint,headers=headers,json={'expected_revision':0,'data':{'title':'过时修改'}})
    assert conflict.status_code in {400,409}
    assert conflict.json()['error']['code']=='litigation_revision_conflict'
    assert client.get(endpoint+'/materials',headers=headers).json()['materials'][0]['path']=='合同.txt'
    service.save(ctx,'facts',fact_proposal(),0)
    fact=service.read(ctx,'facts')['data']['facts'][0]
    response=client.patch(endpoint+'/facts/f1',headers=headers,json={'expected_revision':1,'data':{**fact,'user_note':'已经核对'}})
    assert response.status_code==200
    assert client.get(endpoint,headers=headers).json()['facts']['data']['facts'][0]['status']=='confirmed'
    restarted=LitigationService(service.workspaces)
    assert restarted.read(ctx,'case')['data']['title']=='合成案件'
    assert restarted.read(ctx,'facts')['data']['facts'][0]['user_note']=='已经核对'


def test_case_record_direct_symlink_is_rejected(env,tmp_path):
    service,ctx,_=env
    outside=tmp_path/'outside.json';outside.write_text('foreign')
    (ctx.root_path/'案件信息.json').symlink_to(outside)
    with pytest.raises(ToolError):service.save_case(ctx,{},0)
    assert outside.read_text()=='foreign'


def test_export_facts_preserves_pending_and_conflict_text(env,tmp_path):
    from verilex_runtime.litigation.render import render_markdown
    service,ctx,_=env
    service.save(ctx,'facts',fact_proposal(date=''),0)
    service.save(ctx,'facts',fact_proposal(text='不同提议',date=''),1)
    text=render_markdown('facts',service.snapshot(ctx))
    assert '日期待确认' in text and '待处理冲突' in text
    assert '不同提议' in text and '双方签订合同' in text
    assert '待确认事实 1 条' in text


def test_duplicate_evidence_ids_cannot_drop_materials_on_reorder(env):
    service,ctx,_=env
    proposal=evidence_proposal()
    proposal['evidence'].append({**proposal['evidence'][0], 'name':'另一份证据'})
    with pytest.raises(ToolError):service.save(ctx,'evidence',proposal,0)
    assert service.read(ctx,'evidence')['revision']==0


def save_issues(service, ctx, proposal, expected):
    from tests.litigation.test_updates import confirmed_outline
    confirmed_outline(service, ctx, proposal)
    return service.save(ctx, 'issues', proposal, expected)


@pytest.mark.parametrize('kind',['facts','evidence'])
def test_concurrent_review_during_render_stops_publication(env,tmp_path,kind):
    from verilex_runtime.litigation.render import render_markdown
    service,ctx,_=env
    service.save(ctx,'facts',fact_proposal(),0)
    service.save(ctx,'evidence',evidence_proposal(),0)
    db=make_database(tmp_path,workspace_root=ctx.root_path)
    exporter=LitigationExporter(ArtifactPublisher(repository(db),ArtifactBlobStore(tmp_path/'artifacts'),now=lambda:NOW),temporary_root=tmp_path/'stage',now=lambda:NOW)
    def change_while_rendering(*args):
        result=render_markdown(*args)
        service.save_case(ctx,{'theory':'用户在出稿期间更新案件理论'},0)
        return result
    with patch('verilex_runtime.litigation.exporter.render_markdown',change_while_rendering):
        with pytest.raises(ToolError,match='出稿过程中已更新'):exporter.export(service,ctx,kind)
    assert not list(ctx.root_path.rglob('*.md')) and not list(ctx.root_path.rglob('*.docx'))


def test_legacy_word_source_reuses_existing_reader_without_changing_material(env,tmp_path):
    import shutil, subprocess
    from verilex_runtime.litigation.models import Citation
    converter=shutil.which('textutil')
    if converter is None:pytest.skip('macOS textutil unavailable')
    service,ctx,_=env
    seed=tmp_path/'seed.docx';doc=Document();doc.add_paragraph('虚构合同双方约定支付货款100元。');doc.save(seed)
    target=ctx.root_path/'旧合同.doc'
    subprocess.run([converter,'-convert','doc','-output',str(target),str(seed)],check=True,capture_output=True)
    before=target.read_bytes()
    material=next(m for m in service.materials(ctx)['materials'] if m['path']=='旧合同.doc')
    assert material['supported']
    service.verify_sources(ctx,[Citation(path='旧合同.doc',quote='虚构合同双方约定支付货款100元。')])
    assert target.read_bytes()==before
    with pytest.raises(ToolError):service.verify_sources(ctx,[Citation(path='旧合同.doc',page=1,quote='虚构合同双方约定支付货款100元。')])


def test_legacy_word_without_converter_is_not_offered_as_readable(env):
    service,ctx,_=env
    (ctx.root_path/'旧合同.doc').write_bytes(b'not actual doc')
    with patch('verilex_runtime.tools.files._locate_textutil',return_value=None):
        material=next(m for m in service.materials(ctx)['materials'] if m['path']=='旧合同.doc')
        assert not material['supported'] and '.docx' in material['reason']


def test_pdf_quotes_tolerate_chinese_layout_but_not_changed_content(env, monkeypatch):
    from verilex_runtime.litigation.models import Citation
    service, ctx, _ = env
    # Actual reader remains responsible for file access and page validation;
    # simulate the extracted line wraps from a long Chinese case PDF here.
    monkeypatch.setattr(service, 'source_text', lambda *_: '成交价共计50\n万元，已付\n款。\n账号1 0\nnot liable')
    def source(quote, path='材料.pdf'):
        return Citation(path=path, page=1 if path.endswith('.pdf') else None, quote=quote)
    service.verify_sources(ctx, [source('成交价共计50万元，已付款。')])
    for quote in ('成交价共计500万元', '成交价共计50万元;已付款', '未付款', '账号10', 'notliable'):
        with pytest.raises(ToolError) as failed:
            service.verify_sources(ctx, [source(quote)])
        assert failed.value.code == 'litigation_quote_mismatch'
        assert 'PDF第1页' in str(failed.value)
    with pytest.raises(ToolError):
        service.verify_sources(ctx, [source('成交价共计50万元', '材料.txt')])
