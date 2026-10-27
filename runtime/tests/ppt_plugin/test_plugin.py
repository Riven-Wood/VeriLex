from dataclasses import replace
import io
import json
import zipfile
import pytest
from pptx import Presentation
from pptx.util import Inches
from verilex_runtime.artifacts.blob_store import ArtifactBlobStore
from verilex_runtime.artifacts.publisher import ArtifactPublisher
from verilex_runtime.tools.registry import ToolContext, ToolError, ToolRegistry, ApprovalBehavior
from verilex_runtime.ppt_plugin.service import PowerPointPlugin, create, open_checked, outline, replace_text
from verilex_runtime.ppt_plugin.tools import register_ppt_plugin
from tests.artifacts._support import NOW, make_database, repository

SLIDES = [{'title':'合成材料演示','bullets':['甲公司：虚构主体','项目金额：100万元'],'notes':'测试备注，不含真实案件资料'}, {'title':'下一步安排','bullets':['核对材料清单','确认演示时间']}]

@pytest.fixture
def env(tmp_path):
    root=tmp_path/'workspace';root.mkdir()
    db=make_database(tmp_path,workspace_root=root)
    with db.connect() as connection:
        connection.execute("INSERT INTO tool_calls (id,turn_id,provider_call_id,name,arguments_json,risk_level,status) VALUES ('tool-2','turn-1','provider-2','ppt_plugin_replace','{}','high','running')")
    repo=repository(db)
    publisher=ArtifactPublisher(repo,ArtifactBlobStore(tmp_path/'artifacts'),now=lambda:NOW)
    service=PowerPointPlugin(publisher,now=lambda:NOW)
    context=ToolContext(root,True,True,'thread-1','turn-1','tool-1','workspace-1')
    return service,context,repo,publisher


def test_create_publish_read_replace_retains_original(env):
    service,context,repo,_=env
    result=service.write(context,{'output_name':'汇报.pptx','slides':SLIDES})
    path=context.root_path/result['display_name']; original=path.read_bytes()
    read=service.read(context,{'path':path.name})
    assert read['slide_count']==2 and '测试备注' in read['pages'][0]['text']
    assert repo.get_owned(context.thread_id,result['artifact_id']).media_type.endswith('presentation')
    edited=service.write(replace(context,tool_call_id='tool-2'),{'path':path.name,'output_name':path.name,'old_text':'甲公司','new_text':'乙公司'},edit=True)
    assert edited['display_name']!=path.name and edited['replacements']==1
    assert path.read_bytes()==original
    assert '乙公司' in service.read(context,{'path':edited['display_name']})['pages'][0]['text']
    slide=Presentation(context.root_path/edited['display_name']).slides[0]
    assert slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.size.pt==20


@pytest.mark.parametrize('updates',[{'allow_write':False},{'allow_read':False},{'root_path':None},{'workspace_id':None},{'thread_id':None}])
def test_creation_needs_authority(env,updates):
    service,context,_,_=env
    with pytest.raises(ToolError):service.write(replace(context,**updates),{'output_name':'x.pptx','slides':SLIDES})
    assert not list(context.root_path.iterdir())

@pytest.mark.parametrize('name',['../x.pptx','/tmp/x.pptx','x.ppt','x/p.pptx','x\\p.pptx','.pptx','x:p.pptx'])
def test_invalid_output_names(env,name):
    service,context,_,_=env
    with pytest.raises(ToolError):service.write(context,{'output_name':name,'slides':SLIDES})
    assert not list(context.root_path.iterdir())


def test_source_symlink_and_traversal_rejected(env,tmp_path):
    service,context,_,_=env
    outside=tmp_path/'outside.pptx';create(SLIDES).save(outside)
    (context.root_path/'link.pptx').symlink_to(outside)
    for path in ['link.pptx','../outside.pptx',str(outside)]:
        with pytest.raises(ToolError):service.read(context,{'path':path})

@pytest.mark.parametrize('slides',[[],SLIDES*21,[{'title':'x','bullets':['x']*7}],[{'title':'x','bullets':['x'*121]}],[{'title':'x','bullets':['x'*120]*6}],[{'title':'x','bullets':['\x00']}],[{'title':'x','bullets':['a\n'*30]}]] )
def test_oversized_or_invalid_slides_rejected(slides):
    with pytest.raises(ToolError):create(slides)

@pytest.mark.parametrize('data',[b'no zip',b'PKbroken'])
def test_corrupt_input(data):
    with pytest.raises(ToolError):open_checked(data)


def test_macro_and_entity_rejected():
    for name,content in [('ppt/vbaProject.bin',b'macro'),('ppt/slides/slide1.xml',b'<!DOCTYPE x>')]:
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w') as archive:
            archive.writestr('ppt/presentation.xml','<xml/>');archive.writestr(name,content)
        with pytest.raises(ToolError):open_checked(stream.getvalue())


def test_reader_includes_textboxes_tables_and_notes():
    pres=create(SLIDES);slide=pres.slides[0]
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(1),Inches(1)).text='额外文字'
    slide.shapes.add_table(1,1,Inches(1),Inches(1),Inches(2),Inches(1)).table.cell(0,0).text='表格文字'
    text=outline(pres)['pages'][0]['text']
    assert '额外文字' in text and '表格文字' in text and '测试备注' in text


def test_split_runs_honest_failure_and_no_file(env):
    service,context,_,_=env
    pres=create(SLIDES);frame=pres.slides[0].placeholders[1].text_frame;frame.clear()
    frame.paragraphs[0].add_run().text='甲';frame.paragraphs[0].add_run().text='公司'
    pres.save(context.root_path/'original.pptx')
    with pytest.raises(ToolError,match='跨格式'):
        service.write(context,{'path':'original.pptx','output_name':'new.pptx','old_text':'甲公司','new_text':'乙公司'},edit=True)
    assert not (context.root_path/'new.pptx').exists()

@pytest.mark.asyncio
async def test_registry_schemas_and_approval(env):
    _,context,_,publisher=env
    registry=ToolRegistry();register_ppt_plugin(registry,publisher,now=lambda:NOW)
    definition=registry.get('ppt_plugin_create')
    assert definition.approval_behavior==ApprovalBehavior.RISK_DEFAULT
    assert definition.risk_level.value=='high'
    result=await definition.execute(context,{'output_name':'工具调用.pptx','slides':SLIDES})
    assert result.ok and json.loads(result.output)['artifact_id']


def test_read_output_is_bounded_and_marked():
    pres=create(SLIDES)
    pres.slides[0].placeholders[1].text='长'*50000
    result=outline(pres)
    assert sum(len(page['text']) for page in result['pages'])==48000
    assert result['pages'][0]['truncated'] and result['pages'][1]['truncated']


def test_changed_count_and_source_read_permission(env):
    service,context,_,_=env
    create(SLIDES).save(context.root_path/'read.pptx')
    with pytest.raises(ToolError):service.read(replace(context,allow_read=False),{'path':'read.pptx'})


def test_duplicate_zip_entries_rejected():
    stream=io.BytesIO()
    with pytest.warns(UserWarning,match='Duplicate name'):
        with zipfile.ZipFile(stream,'w') as archive:
            archive.writestr('ppt/presentation.xml','<xml/>')
            archive.writestr('ppt/presentation.xml','<xml/>')
    with pytest.raises(ToolError):open_checked(stream.getvalue())
