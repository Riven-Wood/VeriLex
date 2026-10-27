import hashlib
import json
from pathlib import Path
import sys
import zipfile

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
import pytest

from verilex_runtime.excel_plugin.service import ExcelPlugin, WriteBlock
from verilex_runtime.excel_plugin.server import create_server
from verilex_runtime.mcp.schema import validate_mcp_input_schema


@pytest.fixture
def sample(tmp_path):
    workbook = Workbook()
    s = workbook.active
    s.title = '费用'
    s.append(['项目', '金额', '备注'])
    s.append(['甲', 10, '样例'])
    s.append(['乙', 20, None])
    s['B4'] = '=SUM(B2:B3)'
    s['A1'].font = Font(bold=True)
    s['B2'].number_format = '0.00'
    s['B2'].fill = PatternFill('solid', fgColor='FFFF00')
    s.column_dimensions['A'].width = 24
    s.freeze_panes = 'A2'
    workbook.create_sheet('说明')['A1'] = '合成测试'
    workbook.save(tmp_path / '原件.xlsx')
    return tmp_path


def test_read_and_deliver_preserves_original_style_formula(sample):
    p = ExcelPlugin(sample)
    original = (sample / '原件.xlsx').read_bytes()
    result = p.inspect('原件.xlsx', '费用', 'A1:C4')
    assert len(result['sheets']) == 2
    assert result['data']['values'][3][1] == '=SUM(B2:B3)'
    assert result['data']['cached_values'][3][1] is None
    out = p.deliver('结果.xlsx', [WriteBlock(sheet='费用', start_cell='A5', rows=[['丙', 30, '新增']], copy_style_from_row=2)], '原件.xlsx')
    assert out['source_preserved'] is True
    assert original == (sample / '原件.xlsx').read_bytes()
    wb = load_workbook(sample / '结果.xlsx')
    assert wb['费用']['B5'].value == 30
    assert wb['费用']['B5'].number_format == '0.00'
    assert wb['费用']['B5'].fill.fgColor.rgb == '00FFFF00'
    assert wb['费用']['B4'].value == '=SUM(B2:B3)'
    assert wb['费用'].column_dimensions['A'].width == 24
    assert wb['说明']['A1'].value == '合成测试'


def test_create_summary_new_file(sample):
    p = ExcelPlugin(sample)
    values = p.inspect('原件.xlsx', '费用', 'A2:B3')['data']['values']
    p.deliver('汇总.xlsx', [WriteBlock(sheet='汇总', rows=[['合计'], [sum(row[1] for row in values)]])])
    assert load_workbook(sample / '汇总.xlsx')['汇总']['A2'].value == 30
    assert load_workbook(sample / '汇总.xlsx')['汇总']['A1'].font.bold


@pytest.mark.parametrize('output', ['原件.xlsx', '../越界.xlsx', 'old.xls', 'macro.xlsm', 'missing/out.xlsx'])
def test_reject_bad_output_and_keep_original(sample, output):
    original = (sample / '原件.xlsx').read_bytes()
    with pytest.raises((ValueError, FileNotFoundError)):
        ExcelPlugin(sample).deliver(output, [WriteBlock(sheet='费用', rows=[[1]])], '原件.xlsx')
    assert original == (sample / '原件.xlsx').read_bytes()


def test_symlink_read_write_rejected(sample, tmp_path):
    (sample / '链接.xlsx').symlink_to(sample / '原件.xlsx')
    with pytest.raises(ValueError):
        ExcelPlugin(sample).inspect('链接.xlsx')
    (sample / 'alias').symlink_to(sample, target_is_directory=True)
    with pytest.raises(ValueError):
        ExcelPlugin(sample).deliver('alias/new.xlsx', [WriteBlock(sheet='S', rows=[[1]])])


def test_complex_workbook_read_only(sample):
    wb = load_workbook(sample / '原件.xlsx')
    chart = BarChart(); chart.add_data(Reference(wb['费用'], min_col=2, min_row=1, max_row=3), titles_from_data=True)
    wb['费用'].add_chart(chart, 'E1'); wb.save(sample / '图表.xlsx')
    p = ExcelPlugin(sample)
    assert p.inspect('图表.xlsx')['complex_parts']
    with pytest.raises(ValueError, match='仅支持读取'):
        p.deliver('out.xlsx', [WriteBlock(sheet='费用', rows=[[1]])], '图表.xlsx')
    assert not (sample / 'out.xlsx').exists()


def test_failure_does_not_publish_partial_output(sample):
    p = ExcelPlugin(sample)
    with pytest.raises(Exception):
        p.deliver('失败.xlsx', [WriteBlock(sheet='费用', rows=[[1]]), WriteBlock(sheet='费用', start_cell='bad', rows=[[2]])], '原件.xlsx')
    assert not (sample / '失败.xlsx').exists()
    with pytest.raises(Exception):
        p.deliver('危险.xlsx', [WriteBlock(sheet='费用', rows=[[{'type':'formula','value':'=WEBSERVICE("https://example.com")'}]])])
    assert not (sample / '危险.xlsx').exists()


def test_pagination_and_long_text(sample):
    p = ExcelPlugin(sample)
    result = p.inspect('原件.xlsx', '费用', 'A1:C1000')
    assert result['data']['truncated'] is True
    assert result['data']['next_range'] == 'A667:C1000'
    with pytest.raises(ValueError, match='长度'):
        p.deliver('long.xlsx', [WriteBlock(sheet='S', rows=[['x' * 32768]])])


@pytest.mark.asyncio
async def test_real_stdio_protocol_schemas_and_output(sample):
    entry = Path(__file__).resolve().parents[2] / 'verilex_runtime/excel_plugin/entry.py'
    params = StdioServerParameters(command=sys.executable, args=[str(entry), str(sample)])
    async with stdio_client(params) as (read, write):
        async with ClientSession(read, write) as session:
            await session.initialize()
            tools = (await session.list_tools()).tools
            assert {t.name for t in tools} == {'excel_list_workbooks', 'excel_read_workbook', 'excel_deliver_workbook'}
            for tool in tools:
                validate_mcp_input_schema(tool.inputSchema)
            read_result = await session.call_tool('excel_read_workbook', {'path': '原件.xlsx', 'sheet': '费用', 'cell_range': 'A1:C4'})
            assert not read_result.isError
            result = await session.call_tool('excel_deliver_workbook', {'output': 'MCP成果.xlsx', 'source': '原件.xlsx', 'writes': [{'sheet': '费用', 'start_cell': 'B2', 'rows': [[99]]}]})
            assert not result.isError
            assert load_workbook(sample / 'MCP成果.xlsx')['费用']['B2'].value == 99
            assert load_workbook(sample / '原件.xlsx')['费用']['B2'].value == 10

@pytest.mark.asyncio
async def test_catalog_command_through_real_verilex_pool(sample):
    from verilex_runtime.mcp.client import MCPConnectionPool, MCPResolvedServer
    from verilex_runtime.mcp.importer import MCPServerDraft
    pool = MCPConnectionPool()
    server = MCPResolvedServer('excel-test', MCPServerDraft(display_name='Excel', transport='stdio', command='verilex-excel-plugin', cwd=str(sample)), 0)
    try:
        probe = await pool.probe(server)
        assert len(probe.tools) == 3
        lease = await pool.acquire(server, 'test', 'excel-test')
        try:
            result = await lease.call_tool('excel_deliver_workbook', {'output': '连接成果.xlsx', 'writes': [{'sheet': '汇总', 'rows': [['合计'], [30]]}]})
            assert not result.isError
            assert load_workbook(sample / '连接成果.xlsx')['汇总']['A2'].value == 30
        finally:
            await lease.aclose()
    finally:
        await pool.aclose()


def test_zip_bomb_and_macro_rejected(sample):
    import verilex_runtime.excel_plugin.service as service
    path = sample / 'bad.xlsx'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('xl/vbaProject.bin', b'fake')
    with pytest.raises(ValueError, match='宏'):
        ExcelPlugin(sample).inspect('bad.xlsx')
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('large.bin', b'x' * (service.MAX_EXPANDED + 1))
    with pytest.raises(ValueError, match='解压'):
        ExcelPlugin(sample).inspect('bad.xlsx')


def test_publish_does_not_overwrite_racing_file(sample, monkeypatch):
    import os
    original_link = os.link
    def race(src, dst, **kwargs):
        (sample / 'race.xlsx').write_bytes(b'concurrent-owner')
        return original_link(src, dst, **kwargs)
    monkeypatch.setattr(os, 'link', race)
    with pytest.raises(FileExistsError):
        ExcelPlugin(sample).deliver('race.xlsx', [WriteBlock(sheet='S', rows=[[1]])])
    assert (sample / 'race.xlsx').read_bytes() == b'concurrent-owner'
    assert not list(sample.glob('.excel-*.tmp'))


def test_explicit_cell_types_preserve_literal_strings_and_local_formulas(tmp_path):
    from datetime import datetime
    p=ExcelPlugin(tmp_path)
    p.deliver('显式类型.xlsx',[WriteBlock(sheet='Sheet1',rows=[[
        '2026-02-30','=WEBSERVICE("https://example.com")',
        {'type':'date','value':'2026-02-28'},
        {'type':'formula','value':'=INDIRECT("A1")'},
        {'type':'formula','value':'=HYPERLINK("#Sheet1!A1","跳转")'},
        {'type':'formula','value':'=CELL("row",A1)'},
        {'type':'formula','value':'=FILTERXML("<r><v>1</v></r>","//v")'},
    ]])])
    sheet=load_workbook(tmp_path/'显式类型.xlsx')['Sheet1']
    assert sheet['A1'].value=='2026-02-30' and sheet['A1'].data_type=='s'
    assert sheet['B1'].data_type=='s'
    assert sheet['C1'].value==datetime(2026,2,28)
    assert all(sheet.cell(1,i).data_type=='f' for i in range(4,8))


@pytest.mark.parametrize('formula',[
    '=HYPERLINK("https://example.com","跳转")','=HYPERLINK(A1,"跳转")',
    '=INDIRECT("[external.xlsx]Sheet1!A1")','=INDIRECT(A1)',
    '=WEBSERVICE("https://example.com")',"='external.xlsx'!A1",'=CALL("x")',
])
def test_external_or_unestablished_formula_effect_is_still_rejected(tmp_path,formula):
    p=ExcelPlugin(tmp_path)
    from verilex_runtime.excel_plugin.vendor.errors import UnsafeFormulaError
    with pytest.raises(UnsafeFormulaError):
        p.deliver('拒绝.xlsx',[WriteBlock(sheet='Sheet1',rows=[[{'type':'formula','value':formula}]])])
    assert not (tmp_path/'拒绝.xlsx').exists()


def test_unrelated_edit_does_not_rejudge_existing_formulas(tmp_path):
    wb=Workbook();wb.active.title='Sheet1';wb.active['A1']='=INDIRECT(B1)';wb.save(tmp_path/'原件.xlsx')
    before=(tmp_path/'原件.xlsx').read_bytes()
    ExcelPlugin(tmp_path).deliver('副本.xlsx',[WriteBlock(sheet='Sheet1',start_cell='C1',rows=[['修改']])],'原件.xlsx')
    assert (tmp_path/'原件.xlsx').read_bytes()==before
    assert load_workbook(tmp_path/'副本.xlsx')['Sheet1']['A1'].value=='=INDIRECT(B1)'
