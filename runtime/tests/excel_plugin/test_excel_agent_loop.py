"""Real selected MCP transport + Agent loop using a local scripted provider."""
import asyncio
import hashlib
import json
import pytest
from openpyxl import Workbook, load_workbook
from tests.agent.test_tool_loop import env, make_engine, emit_into
from tests.mcp.test_service_lifecycle import MemoryKeyring
from agent_fakes import scripted_factory, text_delta, tool_call_delta
from verilex_runtime.mcp.client import MCPConnectionPool
from verilex_runtime.mcp.importer import MCPServerDraft
from verilex_runtime.mcp.receipts import TestReceiptStore, ToolDecision
from verilex_runtime.mcp.secrets import MCPSecretJournal, MCPSecretManager
from verilex_runtime.mcp.service import MCPService, MCPSecretPatch
from verilex_runtime.mcp.tools import stable_mcp_tool_name
from verilex_runtime.storage.mcp_repository import MCPRepository
from verilex_runtime.excel_plugin.artifacts import with_excel_outputs


@pytest.mark.asyncio
async def test_conversation_reads_then_delivers_without_paid_model(env):
    root = env['ws_dir']
    workbook = Workbook(); workbook.active.title = '费用'
    workbook.active.append(['甲', 10]); workbook.active.append(['乙', 20]); workbook.save(root/'原件.xlsx')
    before = (root/'原件.xlsx').read_bytes()
    database = env['tool_calls']._database
    secrets = MCPSecretManager(MemoryKeyring(), MCPSecretJournal(database))
    repo = MCPRepository(database, secrets=secrets)
    pool = MCPConnectionPool()
    service = MCPService(database=database, repository=repo, threads=env['threads'], secrets=secrets, receipts=TestReceiptStore(), pool=pool)
    draft = MCPServerDraft(display_name='Excel 合成验收',transport='stdio',command='verilex-excel-plugin',cwd=str(root))
    try:
        tested = await service.test_draft(draft, MCPSecretPatch(), [])
        decisions = [ToolDecision(t.tool_name,True,repo._tool_material(t)[0],repo._tool_material(t)[2]) for t in tested.tools]
        server = await service.create_server(draft,MCPSecretPatch(),tested.test_receipt,decisions,accepted_risk_codes=[])
        repo.select_server(env['thread'].id, server.id)
        read_name = stable_mcp_tool_name(server.id,server.display_name,'excel_read_workbook')
        write_name = stable_mcp_tool_name(server.id,server.display_name,'excel_deliver_workbook')
        script = [[tool_call_delta('excel-read',read_name,json.dumps({'path':'原件.xlsx','sheet':'费用','cell_range':'A1:B2'}))],
            [tool_call_delta('excel-write',write_name,json.dumps({'output':'会话汇总.xlsx','writes':[{'sheet':'汇总','rows':[['合计'],[30]]}]}))],
            [text_delta('已生成会话汇总.xlsx，可在文件栏打开。')]]
        engine = make_engine(env,scripted_factory(script)); engine._mcp_service = service
        turn = engine.start(env['thread'].id,'请读取原件.xlsx，汇总金额并生成新文件，保留原件。',approval_mode='ask')
        events = []
        finished = await asyncio.wait_for(engine.run(turn.id,await emit_into(events)),30)
        assert finished.status.value == 'completed'
        assert load_workbook(root/'会话汇总.xlsx')['汇总']['A2'].value == 30
        assert (root/'原件.xlsx').read_bytes() == before
        files = with_excel_outputs([],env['thread'],env['tool_calls'],repo,env['workspaces'])
        assert [f['name'] for f in files] == ['会话汇总.xlsx']
        assert not [e for e in events if e.type == 'approval_requested']  # Explicit selected MCP owns authorization.
    finally:
        await service.aclose()
        await pool.aclose()
