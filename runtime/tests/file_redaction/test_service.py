import io
import pytest
from docx import Document
from verilex_runtime.storage.database import Database
from verilex_runtime.storage.migrations import run_migrations
from verilex_runtime.file_redaction.service import RedactionService
from verilex_runtime.file_redaction.errors import RedactionError
from .test_pipeline import sample,detector


def test_restart_rerun_delete(tmp_path):
    db=Database(tmp_path/'state.db'); run_migrations(db)
    s=RedactionService(db,detector)
    job=s.create('合同.docx',sample(),['phone'],['张三'])
    first=job['mapping'][0]
    job=s.update(job['id'],job['revision'],first['id'],True)
    s=RedactionService(db,detector)
    again=s.rerun(job['id'],job['revision'],['phone'],['张三','联系'])
    assert next(x for x in again['mapping'] if x['id']==first['id'])['keep']
    assert next(x for x in again['mapping'] if x['id']==first['id'])['placeholder']==first['placeholder']
    assert 'original' not in s.recent()[0] and 'mapping' not in s.recent()[0]
    s.delete_mapping(again['id'],again['revision'])
    assert not s.get(again['id'])['restorable']
    with pytest.raises(RedactionError): s.export(again['id'],again['revision']+1,'restore')
    payload=s.export(again['id'],again['revision']+1,'placeholder')
    assert '张三' not in Document(io.BytesIO(payload)).paragraphs[0].text
    with db.connect() as c: assert not c.execute('SELECT * FROM redaction_mappings').fetchall()


def test_stale_revision_does_not_change(tmp_path):
    db=Database(tmp_path/'state.db'); run_migrations(db); s=RedactionService(db,detector)
    job=s.create('合同.docx',sample(),[],['张三'])
    with pytest.raises(RedactionError): s.update(job['id'],0,job['mapping'][0]['id'],True)
    assert not s.get(job['id'])['mapping'][0]['keep']


def test_review_context_is_bounded_transient_and_hides_adjacent_hits(tmp_path):
    db=Database(tmp_path/'state.db'); run_migrations(db); service=RedactionService(db,detector)
    job=service.create('上下文.docx',sample('甲'*60+'张三联系13800138000，结束'+'乙'*60),['phone'],['张三'])
    hit=next(m for m in job['mapping'] if m['key'].startswith('word/document.xml') and m['old']=='张三')
    assert hit['context_before']=='…'+'甲'*36
    assert '【电话1】' in hit['context_after'] and '13800138000' not in hit['context_after']
    assert len(hit['context_after'])<=37
    with db.connect() as c:
        saved=c.execute('SELECT payload FROM redaction_mappings').fetchone()[0]
        assert 'context_before' not in saved and 'context_after' not in saved
    assert 'context' not in str(service.recent())
    service.delete_mapping(job['id'],job['revision'])
    assert service.get(job['id'])['mapping'] is None


def test_review_order_matches_document_and_preserves_old_mapping_ids(tmp_path):
    db=Database(tmp_path/'state.db'); run_migrations(db); service=RedactionService(db,detector)
    doc=Document()
    for i in range(12): doc.add_paragraph(f'第{i}项：张三')
    out=io.BytesIO();doc.save(out)
    job=service.create('顺序.docx',out.getvalue(),[],['张三'])
    assert [int(m['key'].split(':')[1]) for m in job['mapping']]==list(range(12))
    # Older persisted mappings used lexicographic part/paragraph ordering.
    import json,hashlib
    with db.connect() as c:
        row=c.execute('SELECT payload FROM redaction_mappings').fetchone()[0]
        value=json.loads(row);value['mapping'].sort(key=lambda m:m['key'])
        payload=json.dumps(value,ensure_ascii=False)
        c.execute('UPDATE redaction_mappings SET payload=?,digest=?',(payload,hashlib.sha256(payload.encode()).hexdigest()))
    loaded=service.get(job['id'])
    assert [m['id'] for m in loaded['mapping']]==[m['id'] for m in job['mapping']]


@pytest.mark.parametrize('mapping_deleted', [False, True])
def test_existing_jobs_remove_hidden_editor_text_before_reexport(tmp_path, mapping_deleted):
    import hashlib
    from verilex_runtime.documents.docx_package import parts, pack
    db=Database(tmp_path/'state.db');run_migrations(db);service=RedactionService(db,detector)
    job=service.create('旧脱敏文件.docx',sample(),['phone'],['张三'])
    with db.connect() as c:
        row=c.execute('SELECT canonical FROM redaction_jobs WHERE id=?',(job['id'],)).fetchone()
        items=parts(row[0]);items['verilex/editor.json']='{"old":"张三 13800138000"}'.encode()
        legacy=pack(items)
        c.execute('UPDATE redaction_jobs SET canonical=?,digest=? WHERE id=?',
            (legacy,hashlib.sha256(legacy).hexdigest(),job['id']))
        if mapping_deleted:c.execute('DELETE FROM redaction_mappings WHERE job_id=?',(job['id'],))
    output=service.export(job['id'],job['revision'],'placeholder')
    assert 'verilex/editor.json' not in parts(output)
    with db.connect() as c:
        stored=c.execute('SELECT canonical FROM redaction_jobs WHERE id=?',(job['id'],)).fetchone()[0]
    assert 'verilex/editor.json' not in parts(stored)
    assert service.get(job['id'])['revision'] == job['revision']
    if not mapping_deleted:
        restored=service.export(job['id'],job['revision'],'restore')
        assert Document(io.BytesIO(restored)).paragraphs[0].text=='张三联系张三，电话13800138000。'


def test_deleting_mapping_repairs_old_canonical_without_nested_write_lock(tmp_path):
    import hashlib
    from verilex_runtime.documents.docx_package import parts, pack
    db=Database(tmp_path/'state.db');run_migrations(db);service=RedactionService(db,detector)
    job=service.create('旧文件.docx',sample(),['phone'],['张三'])
    with db.connect() as c:
        raw=c.execute('SELECT canonical FROM redaction_jobs WHERE id=?',(job['id'],)).fetchone()[0]
        items=parts(raw);items['verilex/editor.json']=b'{"original":"private text"}'
        legacy=pack(items)
        c.execute('UPDATE redaction_jobs SET canonical=?,digest=? WHERE id=?',
                  (legacy,hashlib.sha256(legacy).hexdigest(),job['id']))
    result=service.delete_mapping(job['id'],job['revision'])
    assert result['revision']==job['revision']+1 and not result['restorable']
    output=service.export(job['id'],result['revision'],'placeholder')
    assert 'verilex/editor.json' not in parts(output)
