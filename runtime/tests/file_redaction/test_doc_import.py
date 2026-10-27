import io
from docx import Document
from verilex_runtime.attachments import legacy_doc
from verilex_runtime.file_redaction.service import RedactionService
from verilex_runtime.file_redaction.rules import detect_rules
from verilex_runtime.storage.database import Database
from verilex_runtime.storage.migrations import run_migrations


def test_doc_redaction_uses_conversion_then_existing_pipeline(tmp_path,monkeypatch):
    db=Database(tmp_path/'data.db');run_migrations(db)
    d=Document();d.add_paragraph('虚构合同：张三负责验收。');b=io.BytesIO();d.save(b)
    calls=[]
    def convert(raw): calls.append(raw);return b.getvalue()
    monkeypatch.setattr(legacy_doc,'convert_doc',convert)
    service=RedactionService(db,detector=detect_rules)
    raw=legacy_doc._MAGIC+b'synthetic'
    job=service.create('旧合同.doc',raw,[],['张三'])
    assert calls==[raw] and job['filename']=='旧合同.doc'
    output=service.export(job['id'],job['revision'],'placeholder')
    assert output.startswith(b'PK')
    assert '张三' not in '\n'.join(p.text for p in Document(io.BytesIO(output)).paragraphs)
