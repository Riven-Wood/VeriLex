"""Synthetic format variants found missing during the open-source comparison."""
import io
import pytest
from docx import Document
from verilex_runtime.file_redaction.rules import detect_rules, merge_hits
from verilex_runtime.file_redaction.pipeline import prepare, render, restore
from .test_pipeline import sample, detector

@pytest.mark.parametrize('value', ['138 0013 8000','138-0013-8000','＋８６ １３８ ００１３ ８０００','（０１０）１２３４５６７８','01012345678','010-12345678'])
def test_phone_format_keeps_exact_original_coordinates(value):
    text='电话：'+value+'；结束'
    hits=merge_hits(detect_rules(text,['phone'],[]))
    assert len(hits)==1
    assert hits[0]['old']==value
    assert text[hits[0]['start']:hits[0]['end']]==value

@pytest.mark.parametrize('value',['4111111111111111','4111 1111 1111 1111','４１１１－１１１１－１１１１－１１１１'])
def test_card_separators_and_width(value):
    hits=detect_rules('卡号：'+value+'。',['bank'],[])
    assert any(h['old']==value for h in hits)

@pytest.mark.parametrize('label',['银行账号','收款账号','付款账号','银行账户','账号'])
def test_labelled_account_does_not_require_card_checksum(label):
    text=label+'：123456789012；金额：5000元'
    hits=merge_hits(detect_rules(text,['bank'],[]))
    assert len(hits)==1 and hits[0]['old']=='123456789012'

@pytest.mark.parametrize('text',['订单：123456789012','金额：123456789012元','4111111111111112','0000000000000000','1138001380009','合同日期：2026-09-17','银行账号：1234567','银行账号：'+'1'*31])
def test_numeric_counterexamples(text):
    assert not detect_rules(text,['bank','phone'],[])

def test_formatted_values_roundtrip_and_per_occurrence_retention():
    text='收款账号：123456789012；电话：１３８ ００１３ ８０００；备用：１３８ ００１３ ８０００。'
    raw=sample(text); job=prepare(raw,['phone','bank'],[],detector)
    phones=[m for m in job['mapping'] if m['kind']=='phone']
    assert len(phones)==2 and phones[0]['placeholder']==phones[1]['placeholder']
    phones[0]['keep']=True
    out=Document(io.BytesIO(render(job,'permanent',detector))).paragraphs[0].text
    assert out=='收款账号：；电话：１３８ ００１３ ８０００；备用：。'
    assert Document(io.BytesIO(restore(job))).paragraphs[0].text==text

def test_custom_terms_still_match_exact_original_spelling():
    text='编号１２３；编号123'
    hits=detect_rules(text,[],['编号123'])
    assert len(hits)==1 and hits[0]['start']==6


def test_explicit_custom_name_corrects_real_model_boundary():
    from verilex_runtime.file_redaction.ner import detect
    text='张三确认本条为虚构验收数据。'
    job=prepare(sample(text),['person'],['张三'],detect)
    body=[m for m in job['mapping'] if m['key'].startswith('word/document.xml')]
    assert len(body)==1 and body[0]['old']=='张三'
    assert Document(io.BytesIO(render(job,'permanent',detect))).paragraphs[0].text=='确认本条为虚构验收数据。'
