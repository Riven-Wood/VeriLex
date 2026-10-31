import base64
import io
import pytest
from docx import Document
from docx.shared import Pt
from fastapi import FastAPI
from fastapi.testclient import TestClient
from verilex_runtime.documents.docx_package import W, parts, xml
from verilex_runtime.file_redaction.direct import export_direct, _resolve
from verilex_runtime.file_redaction.errors import RedactionError
from verilex_runtime.file_redaction.verify import verify_rejected_revisions
from verilex_runtime.documents.docx_package import pack, serialize
from verilex_runtime.documents.scrub import scrub_package
from verilex_runtime.file_redaction.service import RedactionService
from verilex_runtime.api.redaction import build_redaction_router
from verilex_runtime.storage.database import Database
from verilex_runtime.storage.migrations import run_migrations
from .test_pipeline import detector


def make_document():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('甲方张').bold = True
    p.add_run('三联系13800138000。张三负责。')
    doc.sections[0].header.paragraphs[0].text = '张三'
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '张三'
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


def test_word_tracks_replacements_across_runs_and_automatic_deletion():
    source=make_document()
    result=export_direct(source,['phone'],[{'old':'张三','new':'客户A'}],'word',detector)
    doc=xml(parts(result)['word/document.xml'])
    assert ''.join(n.text or '' for n in doc.iter(f'{{{W}}}delText')).count('张三')==3
    assert '13800138000' in ''.join(n.text or '' for n in doc.iter(f'{{{W}}}delText'))
    assert ''.join(''.join(n.itertext()) for n in doc.iter(f'{{{W}}}ins')).count('客户A')==3
    accepted=Document(io.BytesIO(_resolve(result,True)))
    assert accepted.paragraphs[0].text=='甲方客户A联系。客户A负责。'
    assert accepted.tables[0].cell(0,0).text=='客户A'
    assert accepted.sections[0].header.paragraphs[0].text=='客户A'
    rejected=Document(io.BytesIO(_resolve(result,False)))
    assert rejected.paragraphs[0].text=='甲方张三联系13800138000。张三负责。'
    assert rejected.sections[0].header.paragraphs[0].text=='张三'


def test_table_and_alternate_text_boxes_are_both_processed():
    from lxml import etree
    document=Document()
    table=document.add_table(rows=2, cols=2)
    table.style='Table Grid'
    table.cell(0,0).text='联系人张三'
    table.cell(0,1).text='电话13800138000'
    table.cell(1,0).merge(table.cell(1,1)).text='合并单元格张三'
    nested=table.cell(1,0).add_table(rows=1,cols=1)
    nested.cell(0,0).text='嵌套张三'
    stream=io.BytesIO();document.save(stream)
    items=parts(stream.getvalue())
    root=xml(items['word/document.xml'])
    paragraph=root.find(f'.//{{{W}}}p')
    run=etree.SubElement(paragraph,f'{{{W}}}r')
    mc='http://schemas.openxmlformats.org/markup-compatibility/2006'
    alt=etree.SubElement(run,f'{{{mc}}}AlternateContent')
    for branch in ('Choice','Fallback'):
        child=etree.SubElement(alt,f'{{{mc}}}{branch}')
        box=etree.SubElement(child,f'{{{W}}}txbxContent')
        p=etree.SubElement(box,f'{{{W}}}p')
        for text in ('页脚张','三'):
            r=etree.SubElement(p,f'{{{W}}}r')
            etree.SubElement(r,f'{{{W}}}t').text=text
    items['word/document.xml']=serialize(root)
    source=pack(items)
    output=export_direct(source,['phone'],[{'old':'张三','new':'某人'}],'word',detector)
    accepted=_resolve(output,True)
    accepted_doc=Document(io.BytesIO(accepted))
    assert accepted_doc.tables[0].cell(0,0).text=='联系人某人'
    assert accepted_doc.tables[0].cell(0,1).text=='电话'
    assert accepted_doc.tables[0].cell(1,0).text.startswith('合并单元格某人')
    assert accepted_doc.tables[0].cell(1,0).tables[0].cell(0,0).text=='嵌套某人'
    root=xml(parts(accepted)['word/document.xml'])
    assert [''.join(box.itertext()) for box in root.iter(f'{{{W}}}txbxContent')]==['页脚某人','页脚某人']
    rejected=_resolve(output,False)
    verify_rejected_revisions(scrub_package(source).copy_bytes,rejected)
    assert Document(io.BytesIO(rejected)).tables[0].cell(0,0).text=='联系人张三'


def test_revision_export_preserves_format_and_detects_format_damage():
    document = Document(io.BytesIO(make_document()))
    paragraph = document.paragraphs[0]
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.runs[0].font.size = Pt(16)
    document.tables[0].style = 'Table Grid'
    document.sections[0].footer.paragraphs[0].text = '页脚保留'
    image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')
    document.add_picture(io.BytesIO(image))
    stream = io.BytesIO(); document.save(stream)
    source = stream.getvalue()
    clean = scrub_package(source).copy_bytes
    output = export_direct(source, [], [{'old':'张三','new':'某人'}], 'word', detector)
    rejected = _resolve(output, False)
    verify_rejected_revisions(clean, rejected)
    reopened = Document(io.BytesIO(rejected))
    assert reopened.paragraphs[0].paragraph_format.left_indent == Pt(24)
    assert reopened.paragraphs[0].runs[0].bold and reopened.paragraphs[0].runs[0].font.size == Pt(16)
    assert reopened.tables[0].style.name == 'Table Grid'
    assert reopened.sections[0].footer.paragraphs[0].text == '页脚保留'
    assert len(reopened.inline_shapes) == 1
    damaged = parts(rejected)
    root = xml(damaged['word/document.xml'])
    root.find(f'.//{{{W}}}ind').set(f'{{{W}}}left', '0')
    damaged['word/document.xml'] = serialize(root)
    with pytest.raises(RedactionError):
        verify_rejected_revisions(clean, pack(damaged))


def test_markdown_removes_and_replaces_without_mapping():
    output=export_direct(make_document(),['phone'],[{'old':'张三','new':'某人'}],'markdown',detector).decode()
    assert '某人' in output and '张三' not in output and '13800138000' not in output
    removed=export_direct(make_document(),[],[{'old':'张三','new':''}],'markdown',detector).decode()
    assert '张三' not in removed and '客户A' not in removed


def test_direct_api_is_stateless_and_rejects_duplicate_terms(tmp_path):
    db=Database(tmp_path/'state.db');run_migrations(db)
    app=FastAPI();app.include_router(build_redaction_router(RedactionService(db,detector)))
    with TestClient(app) as client:
        body={'filename':'虚构.docx','content':base64.b64encode(make_document()).decode(),'kinds':[],
              'replacements':[{'old':'张三','new':'客户'}],'variant':'word'}
        res=client.post('/v1/redaction/process',json=body)
        assert res.status_code==200, res.text
        assert res.json()['filename']=='脱敏修订版.docx'
        assert client.get('/v1/redaction').json()==[]
        assert client.post('/v1/redaction/process',json={**body,'replacements':body['replacements']*2}).status_code==422


def test_markdown_input_preserves_unmatched_bytes_bom_and_line_endings(tmp_path):
    db=Database(tmp_path/'state.db');run_migrations(db)
    service=RedactionService(db,detector)
    original='\ufeff# 秘密合同\r\n\r\n**张三**：13800138000\r\n\r\n[材料](./附件.pdf)\r\n'.encode('utf-8')
    result=service.process('测试.MD',original,['phone'],[{'old':'张三','new':'某人'}],'markdown')
    assert result.decode('utf-8')=='\ufeff# 秘密合同\r\n\r\n**某人**：\r\n\r\n[材料](./附件.pdf)\r\n'
    assert service.recent()==[]
    assert original.decode('utf-8').find('张三')>=0


def test_markdown_replacement_cannot_change_structure(tmp_path):
    db=Database(tmp_path/'state.db');run_migrations(db)
    service=RedactionService(db,detector)
    original='# 合同\n\n**张三**：[材料](./证据.pdf)\n'.encode()
    from verilex_runtime.file_redaction.errors import RedactionError
    with pytest.raises(RedactionError) as error:
        service.process('测试.md',original,[],[{'old':'**张三**','new':'某人'}],'markdown')
    assert error.value.code=='redaction_markdown_structure'
    with pytest.raises(RedactionError) as error:
        service.process('测试.md',original,[],[{'old':'张三','new':'**某人**'}],'markdown')
    assert error.value.code=='redaction_markdown_structure'


def test_markdown_api_rejects_invalid_encoding_and_word_output(tmp_path):
    db=Database(tmp_path/'state.db');run_migrations(db)
    app=FastAPI();app.include_router(build_redaction_router(RedactionService(db,detector)))
    with TestClient(app) as client:
        body={'filename':'测试.md','content':base64.b64encode(b'\xff').decode(),'kinds':[], 'replacements':[], 'variant':'markdown'}
        res=client.post('/v1/redaction/process',json=body)
        assert res.status_code==422 and res.json()['error']['code']=='redaction_markdown_encoding'
        body['content']=base64.b64encode('正文'.encode()).decode();body['variant']='word'
        res=client.post('/v1/redaction/process',json=body)
        assert res.status_code==422 and res.json()['error']['code']=='redaction_markdown_format'
        body['variant']='markdown'
        res=client.post('/v1/redaction/process',json=body)
        assert res.status_code==200 and base64.b64decode(res.json()['content'])=='正文'.encode()
