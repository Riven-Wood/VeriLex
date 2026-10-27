import io
from docx import Document
import pytest
from verilex_runtime.file_redaction.pipeline import prepare, render, restore
from verilex_runtime.file_redaction.errors import RedactionError
from verilex_runtime.file_redaction.rules import detect_rules


def sample(text='张三联系张三，电话13800138000。'):
    d = Document(); p = d.add_paragraph(); p.add_run(text[:2]).bold=True; p.add_run(text[2:])
    d.sections[0].header.paragraphs[0].text='张三'
    d.core_properties.author='秘密作者'; out=io.BytesIO(); d.save(out); return out.getvalue()


def detector(text, kinds, custom):
    return detect_rules(text,kinds,custom)


def test_occurrence_keep_and_restore():
    raw=sample(); before=bytes(raw); job=prepare(raw, ['phone'], ['张三'], detector)
    names=[x for x in job['mapping'] if x['old']=='张三']
    assert len(names)==3 and len({x['placeholder'] for x in names})==1
    names[0]['keep']=True
    out=render(job, 'permanent', detector)
    d=Document(io.BytesIO(out)); assert d.paragraphs[0].text=='张三联系，电话。'
    assert not d.core_properties.author
    back=Document(io.BytesIO(restore(job))); assert back.paragraphs[0].text=='张三联系张三，电话13800138000。'
    assert raw==before


def test_existing_placeholder_unchanged_and_empty():
    job=prepare(sample('【自定义1】 张三'),[],['张三'],detector)
    assert job['mapping'][0]['placeholder']!='【自定义1】'
    assert '【自定义1】 张三'==Document(io.BytesIO(restore(job))).paragraphs[0].text
    job=prepare(sample('正文'),[],[],detector)
    assert not job['mapping']
    assert Document(io.BytesIO(render(job,'placeholder',detector))).paragraphs[0].text=='正文'


def test_corruption_and_overlap():
    job=prepare(sample(),['phone'],['张三联系','张三'],detector)
    job['mapping'][0]['start']+=1
    with pytest.raises(RedactionError): restore(job)


def test_rules():
    hits=detect_rules('13800138000 a@example.com 11010519491231002X', ['phone','email','identity'],[])
    assert {h['kind'] for h in hits}=={'phone','email','identity'}
    assert not detect_rules('110105194912310021', ['identity'],[])


def test_simple_field_is_rejected():
    from docx.oxml import OxmlElement
    d=Document(); p=d.add_paragraph(); field=OxmlElement('w:fldSimple');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='13800138000';r.append(t);field.append(r);p._p.append(field)
    out=io.BytesIO(); d.save(out)
    with pytest.raises(RedactionError,match='redaction_complex'): prepare(out.getvalue(),['phone'],[],detector)


def test_markdown_preserves_order():
    d=Document()
    for i in range(12): d.add_paragraph(f'正文段落{i}')
    out=io.BytesIO();d.save(out);job=prepare(out.getvalue(),[],[],detector)
    md=render(job,'markdown',detector).decode()
    assert md.split('\n\n')[:11]==[f'正文段落{i}' for i in range(11)]


def test_candidate_limit_before_merge():
    with pytest.raises(RedactionError,match='redaction_limit'): detect_rules('张'*5001,[],['张'])


def test_picture_alt_text_checked_by_real_ner():
    import base64
    from verilex_runtime.file_redaction.ner import detect
    d=Document();d.add_paragraph('正文')
    png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aI3sAAAAASUVORK5CYII=')
    shape=d.add_picture(io.BytesIO(png));shape._inline.docPr.set('descr','张三身份证照')
    out=io.BytesIO();d.save(out)
    with pytest.raises(RedactionError,match='redaction_residue'):prepare(out.getvalue(),['person'],[],detect)


@pytest.mark.parametrize('variant', ['placeholder', 'permanent'])
def test_redacting_editor_word_removes_hidden_body_and_keeps_original(variant):
    from verilex_runtime.litigation.products import Products, META, OWNER
    from verilex_runtime.documents.docx_package import parts
    from verilex_runtime.documents.scrub_verify import audit_processed_copy
    record = {'owner': OWNER, 'schema_version': 1, 'kind': 'documents',
              'id': 'a'*32, 'revision': 1, 'title': '虚构审查文书',
              'editor_receipt': {'before': {'type': 'doc', 'content': [{'type': 'paragraph',
                  'content': [{'type': 'text', 'text': '已撤销的秘密草稿'}]}]}},
              'document': {'type': 'doc', 'content': [{'type': 'paragraph',
                  'content': [{'type': 'text', 'text': '虚构联系人张三，电话13800138000。'}]}]}}
    original = Products(None)._encode(record)
    job = prepare(original, ['phone'], ['张三'], detector)
    output = render(job, variant, detector)
    assert META in parts(original)
    for processed in (job['canonical'], output):
        assert META not in parts(processed)
        assert all('张三'.encode() not in data and b'13800138000' not in data
                   and '已撤销的秘密草稿'.encode() not in data
                   for data in parts(processed).values())
        assert audit_processed_copy(processed, [], scrubbed=True).passed


def test_residue_audit_rejects_hidden_editor_part_even_without_xml_residue():
    from verilex_runtime.documents.docx_package import parts, pack
    from verilex_runtime.documents.scrub import scrub_package
    from verilex_runtime.documents.scrub_verify import audit_processed_copy
    items = parts(scrub_package(sample('已处理正文')).copy_bytes)
    items['verilex/editor.json'] = b'{"document":"hidden original"}'
    assert not audit_processed_copy(pack(items), [], scrubbed=True).passed


def test_entity_occurrence_does_not_replace_same_spelling_in_other_context():
    d=Document();d.add_paragraph('当事人高兴到场。');d.add_paragraph('我们十分高兴。')
    out=io.BytesIO();d.save(out)
    def contextual(text,kinds,custom):
        return [dict(kind='person',old='高兴',start=3,end=5,priority=2)] if text.startswith('当事人高兴') else []
    job=prepare(out.getvalue(),['person'],[],contextual)
    assert len(job['mapping'])==1
    result=Document(io.BytesIO(render(job,'permanent',contextual)))
    assert result.paragraphs[0].text=='当事人到场。'
    assert result.paragraphs[1].text=='我们十分高兴。'


def test_short_custom_term_is_not_matched_against_xml_ids_and_page_counts():
    job=prepare(sample('编号1'),[],['1'],detector)
    result=Document(io.BytesIO(render(job,'permanent',detector)))
    assert result.paragraphs[0].text=='编号'


def test_non_entity_categories_do_not_require_ner(monkeypatch):
    from verilex_runtime.file_redaction import ner
    def unavailable():raise AssertionError('unneeded model')
    monkeypatch.setattr(ner,'model',unavailable)
    assert ner.detect('电话13800138000',['phone'],[])[0]['kind']=='phone'
    assert ner.detect('自定义甲',[],['自定义甲'])[0]['old']=='自定义甲'
