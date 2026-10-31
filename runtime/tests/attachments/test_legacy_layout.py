import hashlib
import io
from zipfile import ZipFile

from docx import Document
from lxml import etree as E
import pytest

from verilex_runtime.attachments.legacy_layout import restore_layout
from verilex_runtime.attachments import legacy_layout
from verilex_runtime.attachments.errors import AttachmentDomainError

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def test_missing_packaged_reader_reports_the_component_instead_of_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy_layout, '_READER', tmp_path / 'missing-bundle')
    monkeypatch.setattr(legacy_layout, '_DEVELOPMENT_READER', tmp_path / 'missing-development')
    with pytest.raises(AttachmentDomainError) as error:
        legacy_layout.reader_root()
    assert error.value.code == 'attachment_doc_layout_reader_missing'


def fixture():
    d=Document();d.add_paragraph('虚构标题');p=d.add_paragraph()
    props=p._p.get_or_add_pPr();r=E.SubElement(props,W+'rPr')
    E.SubElement(r,W+'sz',{W+'val':'44'});E.SubElement(r,W+'szCs',{W+'val':'44'})
    d.add_paragraph('后续文字。');b=io.BytesIO();d.save(b)
    metadata={'version':1,'compression':1,'kernPunctuation':True,'paragraphs':[
        {'hash':hashlib.sha256(t.encode()).hexdigest(),'table':False} for t in ['虚构标题','','后续文字。']]}
    metadata['paragraphs'][1].update(emptySize=21,emptySizeCs=0)
    return b.getvalue(),metadata


def test_restore_source_empty_font_and_punctuation_without_rewriting_other_parts():
    data,metadata=fixture();out=restore_layout(data,metadata)
    with ZipFile(io.BytesIO(data)) as original, ZipFile(io.BytesIO(out)) as saved:
        assert original.namelist()==saved.namelist()
        for part in original.namelist():
            if part not in {'word/document.xml','word/settings.xml'}:assert original.read(part)==saved.read(part)
        root=E.fromstring(saved.read('word/document.xml'));p=root.find(W+'body').findall(W+'p')[1]
        assert p.find('./'+W+'pPr/'+W+'rPr/'+W+'sz').get(W+'val')=='21'
        assert p.find('./'+W+'pPr/'+W+'rPr/'+W+'szCs') is None
        settings=E.fromstring(saved.read('word/settings.xml'))
        assert settings.find(W+'characterSpacingControl').get(W+'val')=='compressPunctuation'
        assert settings.find(W+'noPunctuationKerning').get(W+'val')=='0'
    assert [p.text for p in Document(io.BytesIO(out)).paragraphs]==['虚构标题','','后续文字。']


@pytest.mark.parametrize('change',['hash','count','font','table'])
def test_ambiguous_or_invalid_source_mapping_fails_before_rewriting(change):
    data,metadata=fixture()
    if change=='hash':metadata['paragraphs'][0]['hash']='wrong'
    if change=='count':metadata['paragraphs'].pop()
    if change=='font':metadata['paragraphs'][1]['emptySize']=0
    if change=='table':metadata['paragraphs'][1]['table']=True
    with pytest.raises(AttachmentDomainError):restore_layout(data,metadata)


def test_intentionally_large_empty_paragraph_keeps_its_source_size():
    data,metadata=fixture();metadata['paragraphs'][1].update(emptySize=44,emptySizeCs=44)
    d=Document(io.BytesIO(restore_layout(data,metadata)))
    assert d.paragraphs[1]._p.pPr.find('./'+W+'rPr/'+W+'sz').get(W+'val')=='44'
