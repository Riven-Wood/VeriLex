from pathlib import Path
import io
import subprocess
import json
import hashlib
import pytest
from docx import Document
from verilex_runtime.attachments import legacy_doc
from verilex_runtime.attachments.errors import AttachmentDomainError
from verilex_runtime.attachments.inspection import inspect_file


@pytest.fixture(autouse=True)
def reader_resources(tmp_path, monkeypatch):
    reader=tmp_path/'reader';(reader/'bin').mkdir(parents=True);(reader/'reader').mkdir()
    (reader/'bin/java').write_text('unit-test executable');(reader/'bin/java').chmod(0o700)
    (reader/'reader/legacy-layout.jar').write_bytes(b'unit-test jar')
    monkeypatch.setattr(legacy_doc.legacy_layout,'_READER',reader)
    monkeypatch.setattr(legacy_doc.legacy_layout,'_DEVELOPMENT_READER',reader)


def _docx():
    d=Document();d.add_paragraph('虚构合同：验收后付款。');d.add_table(rows=1, cols=1).cell(0,0).text='付款阶段'
    out=io.BytesIO();d.save(out);return out.getvalue()


def test_legacy_signature_rejects_renamed_file():
    with pytest.raises(AttachmentDomainError) as error:
        legacy_doc.convert_doc(b'not Word')
    assert error.value.code == 'attachment_media_type_mismatch'


def test_missing_converter_has_actionable_error(monkeypatch):
    monkeypatch.setattr(legacy_doc.sys,'platform','unsupported')
    with pytest.raises(AttachmentDomainError) as error:
        legacy_doc.convert_doc(legacy_doc._MAGIC + b'fixture')
    assert error.value.code == 'attachment_doc_converter_missing'


def test_converter_uses_private_job_and_validates_output(monkeypatch):
    original=legacy_doc._MAGIC + b'fixture'; jobs=[]
    monkeypatch.setattr(legacy_doc.sys,'platform','darwin')
    real_is_file=Path.is_file
    monkeypatch.setattr(Path,'is_file',lambda p: True if str(p) in {'/Applications/LibreOffice.app/Contents/MacOS/soffice','/usr/bin/sandbox-exec'} else real_is_file(p))
    def run(command,**kwargs):
        root=kwargs['cwd'];jobs.append(root)
        assert (root/'source.doc').read_bytes()==original
        policy=(root/'import.sb').read_text()
        assert '(deny network*)' in policy and '(deny file-write*)' in policy
        assert '(subpath "/Users")' in policy
        assert 'MacroSecurityLevel' in (root/'profile/user/registrymodifications.xcu').read_text()
        if '-cp' in command:
            assert kwargs['timeout']==20
            metadata={'version':1,'compression':1,'kernPunctuation':True,'paragraphs':[
                {'hash':hashlib.sha256(t.encode()).hexdigest(),'table':table}
                for t,table in [('虚构合同：验收后付款。',False),('付款阶段',True)]
            ]}
            return subprocess.CompletedProcess(command,0,stdout=json.dumps(metadata).encode())
        assert kwargs['timeout']==45 and '--infilter=MS Word 97' in command
        (root/'source.docx').write_bytes(_docx())
        return subprocess.CompletedProcess(command,0)
    monkeypatch.setattr(legacy_doc.subprocess,'run',run)
    result=legacy_doc.convert_doc(original)
    assert Document(io.BytesIO(result)).tables[0].cell(0,0).text=='付款阶段'
    assert not jobs[0].exists()


def test_attachment_keeps_original_doc_bytes(tmp_path,monkeypatch):
    raw=legacy_doc._MAGIC+b'fixture';p=tmp_path/'upload.part';p.write_bytes(raw)
    monkeypatch.setattr(legacy_doc,'convert_doc',lambda data: _docx())
    result=inspect_file(p,'合同.doc','application/msword')
    assert result.media_type=='application/msword' and '验收后付款' in result.extracted_text
    assert p.read_bytes()==raw


@pytest.mark.parametrize('bad_result',[b'not zip',None])
def test_invalid_or_absent_converted_output_rejected(monkeypatch,bad_result):
    monkeypatch.setattr(legacy_doc.sys,'platform','darwin')
    real=Path.is_file
    monkeypatch.setattr(Path,'is_file',lambda p: True if str(p) in {'/Applications/LibreOffice.app/Contents/MacOS/soffice','/usr/bin/sandbox-exec'} else real(p))
    def run(command,**kw):
        if bad_result is not None: (kw['cwd']/'source.docx').write_bytes(bad_result)
        return subprocess.CompletedProcess(command,0)
    monkeypatch.setattr(legacy_doc.subprocess,'run',run)
    with pytest.raises(AttachmentDomainError): legacy_doc.convert_doc(legacy_doc._MAGIC+b'fixture')


def test_conversion_timeout_is_stable_and_cleans_job(monkeypatch):
    monkeypatch.setattr(legacy_doc.sys,'platform','darwin'); jobs=[]
    real=Path.is_file
    monkeypatch.setattr(Path,'is_file',lambda p: True if str(p) in {'/Applications/LibreOffice.app/Contents/MacOS/soffice','/usr/bin/sandbox-exec'} else real(p))
    def run(command,**kw):
        jobs.append(kw['cwd']);raise subprocess.TimeoutExpired(command,45)
    monkeypatch.setattr(legacy_doc.subprocess,'run',run)
    with pytest.raises(AttachmentDomainError) as error: legacy_doc.convert_doc(legacy_doc._MAGIC+b'fixture')
    assert error.value.code=='attachment_doc_conversion_failed'
    assert not jobs[0].exists()
