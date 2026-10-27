import hashlib
import asyncio
import os
import stat
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

from verilex_runtime.attachments.blob_store import AttachmentBlobStore
from verilex_runtime.attachments.errors import AttachmentDomainError
from verilex_runtime.attachments.models import InspectedAttachment
from verilex_runtime.attachments.models import StagedFile
from verilex_runtime.storage.models import AttachmentBlobRecord, ExtractionStatus


async def _chunks(*values: bytes):
    for value in values:
        yield value


@pytest.mark.asyncio
async def test_stage_streams_chunks_and_computes_sha_without_content_length(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"legal", b" text"), hard_limit=20)
    assert staged.path.read_bytes() == b"legal text"
    assert staged.size_bytes == 10
    assert staged.sha256 == hashlib.sha256(b"legal text").hexdigest()


@pytest.mark.asyncio
async def test_stage_rejects_before_writing_overflow_chunk_and_removes_part(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    with pytest.raises(AttachmentDomainError) as caught:
        await store.stage("upl-1", _chunks(b"a" * 8, b"b" * 8), hard_limit=10)
    assert caught.value.code == "attachment_too_large"
    assert list((tmp_path / "attachments" / "tmp").glob("*.part")) == []


@pytest.mark.asyncio
async def test_stage_preserves_limit_error_when_cleanup_unlink_fails(tmp_path, monkeypatch):
    store = AttachmentBlobStore(tmp_path / "attachments")
    original_unlink = Path.unlink

    def fail_part_unlink(path, *args, **kwargs):
        if path.name == "upl-1.part":
            raise OSError("cleanup failed")
        return original_unlink(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", fail_part_unlink)
    with pytest.raises(AttachmentDomainError) as caught:
        await store.stage("upl-1", _chunks(b"a" * 11), hard_limit=10)
    assert caught.value.code == "attachment_too_large"


@pytest.mark.asyncio
async def test_stage_accepts_exact_limit(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"a" * 10), hard_limit=10)
    assert staged.size_bytes == 10


@pytest.mark.asyncio
async def test_stage_does_not_delete_preexisting_part(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    old = tmp_path / "attachments" / "tmp" / "upl-1.part"
    old.write_bytes(b"do not delete")
    with pytest.raises(FileExistsError):
        await store.stage("upl-1", _chunks(b"new"))
    assert old.read_bytes() == b"do not delete"


@pytest.mark.asyncio
async def test_stage_cleans_part_when_iterator_is_cancelled(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")

    async def cancelled():
        yield b"partial"
        raise asyncio.CancelledError

    with pytest.raises(asyncio.CancelledError):
        await store.stage("upl-1", cancelled())
    assert list((tmp_path / "attachments" / "tmp").glob("*.part")) == []


@pytest.mark.asyncio
@pytest.mark.parametrize("upload_id", ["../escape", "nested/name", "", "."])
async def test_stage_rejects_untrusted_upload_id(tmp_path, upload_id):
    store = AttachmentBlobStore(tmp_path / "attachments")
    with pytest.raises(ValueError):
        await store.stage(upload_id, _chunks(b"x"))


def test_store_directories_are_owner_only(tmp_path):
    root = tmp_path / "attachments"
    AttachmentBlobStore(root)
    for path in (root, root / "blobs", root / "extracted", root / "tmp"):
        assert stat.S_IMODE(path.stat().st_mode) == 0o700


@pytest.mark.asyncio
async def test_staged_file_is_owner_only(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"secret"))
    assert stat.S_IMODE(staged.path.stat().st_mode) == 0o600


@pytest.mark.asyncio
async def test_promote_is_content_addressed_and_reuses_existing_files(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    first = await store.stage("upl-1", _chunks(b"same"))
    paths1 = store.promote(first, inspected)
    second = await store.stage("upl-2", _chunks(b"same"))
    paths2 = store.promote(second, inspected)
    assert paths1 == paths2
    assert store.resolve(paths1.stored_path).read_bytes() == b"same"
    assert store.resolve(paths1.extracted_text_path).read_text() == "same"
    assert not first.path.exists() and not second.path.exists()


@pytest.mark.asyncio
async def test_promote_rehashes_staged_bytes_before_publishing(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"good"))
    staged.path.write_bytes(b"evil")
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "good")
    with pytest.raises(ValueError, match="digest"):
        store.promote(staged, inspected)
    canonical = tmp_path / "attachments" / "blobs" / staged.sha256[:2] / staged.sha256
    assert not canonical.exists()


@pytest.mark.asyncio
async def test_promote_rejects_same_inode_rewrite_during_atomic_link(tmp_path, monkeypatch):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"good"))
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "good")
    original_link = os.link

    def rewrite_before_link(source, destination, *args, **kwargs):
        if Path(destination).name == staged.sha256:
            Path(source).write_bytes(b"evil")
        return original_link(source, destination, *args, **kwargs)

    monkeypatch.setattr(os, "link", rewrite_before_link)
    with pytest.raises(ValueError, match="digest"):
        store.promote(staged, inspected)
    canonical = tmp_path / "attachments" / "blobs" / staged.sha256[:2] / staged.sha256
    assert not canonical.exists()


@pytest.mark.asyncio
@pytest.mark.parametrize("_repeat", range(10))
async def test_two_store_instances_publish_same_sha_without_overwrite(
    tmp_path, monkeypatch, _repeat
):
    root = tmp_path / "attachments"
    first_store = AttachmentBlobStore(root)
    second_store = AttachmentBlobStore(root)
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    first = await first_store.stage("upl-1", _chunks(b"same"))
    second = await second_store.stage("upl-2", _chunks(b"same"))
    canonical = root / "blobs" / first.sha256[:2] / first.sha256
    original_replace = os.replace
    replace_barrier = threading.Barrier(2)
    replace_lock = threading.Lock()

    def reject_racing_overwrite(source, destination):
        if Path(destination) == canonical:
            replace_barrier.wait(timeout=2)
            with replace_lock:
                if canonical.exists():
                    raise AssertionError("canonical blob was overwritten")
                return original_replace(source, destination)
        return original_replace(source, destination)

    monkeypatch.setattr(os, "replace", reject_racing_overwrite)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda item: item[0].promote(item[1], inspected),
                ((first_store, first), (second_store, second)),
            )
        )
    assert results[0] == results[1]
    assert first_store.resolve(results[0].stored_path).read_bytes() == b"same"
    assert list((root / "tmp").glob("*.part")) == []


@pytest.mark.asyncio
async def test_reusing_concurrent_canonical_does_not_change_publishers_ctime(tmp_path, monkeypatch):
    root = tmp_path / "attachments"
    first_store = AttachmentBlobStore(root)
    second_store = AttachmentBlobStore(root)
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    first = await first_store.stage("upl-1", _chunks(b"same"))
    second = await second_store.stage("upl-2", _chunks(b"same"))
    publisher_hash_started = threading.Event()
    reuse_verified = threading.Event()
    digest_calls = 0
    original_digest = AttachmentBlobStore._stable_descriptor_digest

    def controlled_digest(descriptor):
        nonlocal digest_calls
        digest_calls += 1
        if digest_calls == 1:
            return original_digest(descriptor)
        before = os.fstat(descriptor)
        publisher_hash_started.set()
        assert reuse_verified.wait(timeout=2)
        os.lseek(descriptor, 0, os.SEEK_SET)
        digest = hashlib.sha256()
        size = 0
        while chunk := os.read(descriptor, 1024 * 1024):
            digest.update(chunk)
            size += len(chunk)
        after = os.fstat(descriptor)
        if (before.st_ctime_ns, before.st_mtime_ns) != (
            after.st_ctime_ns,
            after.st_mtime_ns,
        ):
            raise ValueError("Attachment content changed while being verified.")
        return size, digest.hexdigest()

    original_verify = second_store._verify_canonical_blob

    def signal_reuse(path, staged):
        assert publisher_hash_started.wait(timeout=2)
        try:
            return original_verify(path, staged)
        finally:
            reuse_verified.set()

    monkeypatch.setattr(first_store, "_stable_descriptor_digest", controlled_digest)
    monkeypatch.setattr(second_store, "_verify_canonical_blob", signal_reuse)
    with ThreadPoolExecutor(max_workers=2) as pool:
        publisher = pool.submit(first_store.promote, first, inspected)
        assert publisher_hash_started.wait(timeout=2)
        reuser = pool.submit(second_store.promote, second, inspected)
        assert publisher.result(timeout=3) == reuser.result(timeout=3)


def test_stable_digest_ignores_hardlink_unlink_ctime_only_change(
    tmp_path, monkeypatch
):
    source = tmp_path / "source"
    sibling = tmp_path / "sibling"
    source.write_bytes(b"same")
    os.link(source, sibling)
    descriptor = os.open(source, os.O_RDONLY)
    original_read = os.read
    removed = False

    def unlink_after_first_read(fd, amount):
        nonlocal removed
        chunk = original_read(fd, amount)
        if chunk and not removed:
            sibling.unlink()
            removed = True
        return chunk

    monkeypatch.setattr(os, "read", unlink_after_first_read)
    try:
        assert AttachmentBlobStore._stable_descriptor_digest(descriptor) == (
            4,
            hashlib.sha256(b"same").hexdigest(),
        )
    finally:
        os.close(descriptor)


def test_stable_digest_still_rejects_real_content_change_during_read(
    tmp_path, monkeypatch
):
    source = tmp_path / "source"
    source.write_bytes(b"a" * (1024 * 1024 + 1))
    descriptor = os.open(source, os.O_RDONLY)
    original_read = os.read
    changed = False

    def rewrite_after_first_read(fd, amount):
        nonlocal changed
        chunk = original_read(fd, amount)
        if chunk and not changed:
            with source.open("r+b") as target:
                target.seek(-1, os.SEEK_END)
                target.write(b"b")
                target.flush()
                os.fsync(target.fileno())
            changed = True
        return chunk

    monkeypatch.setattr(os, "read", rewrite_after_first_read)
    try:
        with pytest.raises(ValueError, match="changed"):
            AttachmentBlobStore._stable_descriptor_digest(descriptor)
    finally:
        os.close(descriptor)


@pytest.mark.asyncio
async def test_promote_tightens_reused_blob_and_cache_permissions(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    first = await store.stage("upl-1", _chunks(b"same"))
    paths = store.promote(first, inspected)
    blob = store.resolve(paths.stored_path)
    cache = store.resolve(paths.extracted_text_path)
    blob.chmod(0o644)
    cache.chmod(0o644)
    second = await store.stage("upl-2", _chunks(b"same"))
    store.promote(second, inspected)
    assert stat.S_IMODE(blob.stat().st_mode) == 0o600
    assert stat.S_IMODE(cache.stat().st_mode) == 0o600


@pytest.mark.asyncio
@pytest.mark.parametrize('corrupt', [False, True])
async def test_reupload_nested_word_upgrades_only_exact_legacy_cache(tmp_path, corrupt):
    from dataclasses import replace
    from docx import Document
    from verilex_runtime.attachments.inspection import inspect_file
    document = Document()
    document.add_paragraph('合同正文')
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = '付款约定'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = '嵌套违约责任'
    path = tmp_path / 'nested.docx'; document.save(path)
    inspected = inspect_file(path, 'nested.docx', None)
    assert inspected.legacy_extracted_text is not None
    assert '嵌套违约责任' not in inspected.legacy_extracted_text
    store = AttachmentBlobStore(tmp_path / 'attachments')
    data = path.read_bytes()
    first = await store.stage('legacy', _chunks(data))
    paths = store.promote(first, replace(inspected, extracted_text=inspected.legacy_extracted_text,
                                         legacy_extracted_text=None))
    cache = store.resolve(paths.extracted_text_path)
    if corrupt:
        cache.write_text('tampered cache')
    second = await store.stage('current', _chunks(data))
    if corrupt:
        with pytest.raises(ValueError, match='inconsistent'):
            store.promote(second, inspected)
        assert cache.read_text() == 'tampered cache'
    else:
        assert store.promote(second, inspected) == paths
        assert '嵌套违约责任' in cache.read_text()
        assert stat.S_IMODE(cache.stat().st_mode) == 0o600
    assert store.resolve(paths.stored_path).read_bytes() == data


@pytest.mark.asyncio
async def test_cache_publish_failure_leaves_correct_canonical_blob(tmp_path, monkeypatch):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"same"))
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    original_link = os.link

    def fail_cache_link(source, destination, *args, **kwargs):
        if Path(destination).suffix == ".txt":
            raise OSError("cache link failed")
        return original_link(source, destination, *args, **kwargs)

    monkeypatch.setattr(os, "link", fail_cache_link)
    with pytest.raises(OSError, match="cache link failed"):
        store.promote(staged, inspected)
    assert not staged.path.exists()
    assert list((tmp_path / "attachments").rglob("*.part")) == []
    canonical = tmp_path / "attachments" / "blobs" / staged.sha256[:2] / staged.sha256
    assert canonical.read_bytes() == b"same"


@pytest.mark.asyncio
async def test_failed_cache_promotion_never_removes_existing_blob(tmp_path, monkeypatch):
    store = AttachmentBlobStore(tmp_path / "attachments")
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    first = await store.stage("upl-1", _chunks(b"same"))
    paths = store.promote(first, inspected)
    store.resolve(paths.extracted_text_path).unlink()
    second = await store.stage("upl-2", _chunks(b"same"))
    original_link = os.link

    def fail_cache_link(source, destination, *args, **kwargs):
        if Path(destination).suffix == ".txt":
            raise OSError("cache link")
        return original_link(source, destination, *args, **kwargs)

    monkeypatch.setattr(os, "link", fail_cache_link)
    with pytest.raises(OSError, match="cache link"):
        store.promote(second, inspected)
    assert store.resolve(paths.stored_path).read_bytes() == b"same"


@pytest.mark.asyncio
async def test_promote_fsyncs_every_changed_directory_in_durable_order(tmp_path, monkeypatch):
    root = tmp_path / "attachments"
    store = AttachmentBlobStore(root)
    staged = await store.stage("upl-1", _chunks(b"same"))
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    calls = []
    monkeypatch.setattr(store, "_fsync_directory", lambda path: calls.append(path))
    store.promote(staged, inspected)
    blob_dir = root / "blobs" / staged.sha256[:2]
    assert calls == [
        root / "blobs",
        blob_dir,
        root / "tmp",
        root / "extracted",
        root / "extracted",
    ]


@pytest.mark.asyncio
async def test_reused_blob_unlink_is_fsynced_in_tmp(tmp_path, monkeypatch):
    root = tmp_path / "attachments"
    store = AttachmentBlobStore(root)
    inspected = InspectedAttachment("image.png", "image/png", 4, None)
    first = await store.stage("upl-1", _chunks(b"same"))
    store.promote(first, inspected)
    second = await store.stage("upl-2", _chunks(b"same"))
    calls = []
    monkeypatch.setattr(store, "_fsync_directory", lambda path: calls.append(path))
    store.promote(second, inspected)
    assert calls == [root / "tmp"]


@pytest.mark.asyncio
async def test_promote_preserves_publish_error_when_cleanup_fails(tmp_path, monkeypatch):
    store = AttachmentBlobStore(tmp_path / "attachments")
    staged = await store.stage("upl-1", _chunks(b"same"))
    inspected = InspectedAttachment("image.png", "image/png", 4, None)
    monkeypatch.setattr(os, "link", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("publish failed")))
    monkeypatch.setattr(Path, "unlink", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("cleanup failed")))
    with pytest.raises(OSError, match="publish failed"):
        store.promote(staged, inspected)


def test_promote_rejects_invalid_sha_and_staged_path(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    outside = tmp_path / "outside.part"
    outside.write_bytes(b"same")
    inspected = InspectedAttachment("case.txt", "text/plain", 4, "same")
    with pytest.raises(ValueError):
        store.promote(StagedFile(outside, "A" * 64, 4), inspected)
    assert outside.exists()


@pytest.mark.asyncio
async def test_stage_cleans_part_when_fsync_fails(tmp_path, monkeypatch):
    store = AttachmentBlobStore(tmp_path / "attachments")
    monkeypatch.setattr(os, "fsync", lambda _fd: (_ for _ in ()).throw(OSError("fsync")))
    with pytest.raises(OSError, match="fsync"):
        await store.stage("upl-1", _chunks(b"same"))
    assert list((tmp_path / "attachments" / "tmp").glob("*.part")) == []


def test_resolve_rejects_absolute_and_escaping_paths(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    with pytest.raises(ValueError):
        store.resolve(Path("/tmp/outside"))
    with pytest.raises(ValueError):
        store.resolve(Path("../outside"))
    with pytest.raises(ValueError):
        store.resolve(Path("blobs/../tmp/file.part"))
    with pytest.raises(ValueError):
        store.resolve("bad\x00path")


def test_resolve_rejects_symlinks_even_when_target_stays_inside_root(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    target = tmp_path / "attachments" / "blobs" / "target"
    target.write_bytes(b"x")
    link = tmp_path / "attachments" / "blobs" / "link"
    link.symlink_to(target)
    with pytest.raises(ValueError):
        store.resolve("blobs/link")


def test_open_stored_file_rejects_replaced_root_ancestor(tmp_path):
    base = tmp_path / "guarded"
    store = AttachmentBlobStore(base / "attachments")
    extracted = store.root / "extracted" / "probe.txt"
    extracted.parent.mkdir(parents=True, exist_ok=True)
    extracted.write_bytes(b"real-extracted")

    outside = tmp_path / "outside"
    outside_extracted = outside / "attachments" / "extracted" / "probe.txt"
    outside_extracted.parent.mkdir(parents=True)
    outside_extracted.write_bytes(b"attacker-extracted")

    base.rename(tmp_path / "moved-away")
    base.symlink_to(outside, target_is_directory=True)

    with pytest.raises((OSError, ValueError)):
        descriptor = store.open_stored_file("extracted/probe.txt")
        os.close(descriptor)


def test_remove_blob_files_removes_only_record_paths(tmp_path):
    store = AttachmentBlobStore(tmp_path / "attachments")
    raw = tmp_path / "attachments" / "blobs" / "aa" / ("a" * 64)
    extracted = tmp_path / "attachments" / "extracted" / (("a" * 64) + ".txt")
    raw.parent.mkdir(parents=True)
    raw.write_bytes(b"x")
    extracted.write_text("x")
    record = AttachmentBlobRecord(
        sha256="a" * 64,
        media_type="text/plain",
        size_bytes=1,
        stored_path=str(raw.relative_to(tmp_path / "attachments")),
        extracted_text_path=str(extracted.relative_to(tmp_path / "attachments")),
        extraction_status=ExtractionStatus.READY,
        extraction_error_code=None,
        created_at="2026-08-14T00:00:00Z",
    )
    store.remove_blob_files(record)
    assert not raw.exists() and not extracted.exists()
