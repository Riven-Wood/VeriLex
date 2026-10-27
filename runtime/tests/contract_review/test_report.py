"""Report integration: genuine DOCX, shared gates, content and source preservation."""
import base64
import io
from dataclasses import replace
from zipfile import ZipFile

import pytest
from docx import Document

from tests.contract_review.test_export import env, _review_with_contract, _opinion
from verilex_runtime.contract_review.errors import ContractReviewError
from verilex_runtime.contract_review.report import verify_report
from verilex_runtime.storage.models import TurnRecord, TurnStatus


def test_packaged_contract_smoke_checks_blank_contract_and_final_pagination(capsys):
    from verilex_runtime.contract_review.smoke import run
    assert run() == 0
    assert '"report": true' in capsys.readouterr().out


@pytest.mark.asyncio
async def test_report_uses_edited_adopted_content_and_template_geometry(env, tmp_path):
    service, repo, *_ = env
    handle = await _review_with_contract(env, tmp_path)
    rid = handle['review_id']
    before = repo.review_bytes(rid)
    _opinion(repo, rid, 'keep', body='旧的意见正文', clause_excerpt='不能冒充真实合同的摘录')
    _opinion(repo, rid, 'drop', body='不应交付的意见')
    service.update_opinion(rid, 'keep', edited_text='约定可计算的责任上限。')
    service.update_opinion(rid, 'drop', adopted=False)
    result = service.export_report(rid)
    assert result['filename'] == '技术服务合同（审查意见书）.docx'
    data = base64.b64decode(result['content'])
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == '合同审核意见书'
    assert '约定可计算的责任上限。' in texts
    assert '合同原文：千分之三向甲方支付违约金' in texts
    assert not any(word in '\n'.join(texts) for word in ['旧的意见正文', '不应交付的意见', '不能冒充真实合同的摘录'])
    assert len(doc.comments) == 0
    sec = doc.sections[0]
    assert (sec.page_width.twips, sec.page_height.twips) == (11906, 16838)
    assert (sec.left_margin.twips, sec.top_margin.twips) == (1800, 1440)
    assert doc.styles['Normal'].font.name == '宋体'
    assert doc.styles['Normal'].font.size.pt == 14
    assert doc.styles['Normal'].paragraph_format.line_spacing.pt == 25
    assert repo.review_bytes(rid) == before
    # Both independent outputs remain available after creating a report.
    comments = Document(io.BytesIO(base64.b64decode(service.export(rid)['content']))).comments
    assert len(comments) == 1 and '约定可计算的责任上限。' in list(comments)[0].text
    with ZipFile(io.BytesIO(data)) as z:
        assert not any(n.startswith('word/header') for n in z.namelist())


@pytest.mark.asyncio
@pytest.mark.parametrize('change,code', [
    ({'anchor_text': '合同不存在的内容'}, 'contract_annotation_write_failed'),
])
async def test_report_cannot_bypass_existing_export_gates(env, tmp_path, change, code):
    service, repo, *_ = env
    h = await _review_with_contract(env, tmp_path)
    _opinion(repo, h['review_id'], 'bad', **change)
    with pytest.raises(ContractReviewError) as error:
        service.export_report(h['review_id'])
    assert error.value.code == code


@pytest.mark.asyncio
async def test_report_preserves_edits_but_blocks_busy_and_no_adopted(env, tmp_path):
    service, repo, *_, turns = env
    h = await _review_with_contract(env, tmp_path)
    rid = h['review_id']
    _opinion(repo, rid, 'keep')
    service.update_opinion(rid, 'keep', edited_text='依照《虚构条例》处理。')
    result = service.export_report(rid)
    document = Document(io.BytesIO(base64.b64decode(result['content'])))
    assert any('依照《虚构条例》处理。' in p.text for p in document.paragraphs)
    service.update_opinion(rid, 'keep', adopted=False)
    with pytest.raises(ContractReviewError, match='contract_annotation_write_failed'):
        service.export_report(rid)
    turns.create(TurnRecord(id='busy', thread_id=h['thread_id'], status=TurnStatus.RUNNING,
        error_code=None, started_at='2026-10-03T00:00:00Z', finished_at=None, created_at='2026-10-03T00:00:00Z'))
    with pytest.raises(ContractReviewError, match='contract_thread_busy'):
        service.export_report(rid)


@pytest.mark.asyncio
async def test_report_independent_verifier_rejects_missing_or_changed_text(env, tmp_path):
    service, repo, *_ = env
    h = await _review_with_contract(env, tmp_path)
    _opinion(repo, h['review_id'], 'keep')
    data = base64.b64decode(service.export_report(h['review_id'])['content'])
    expected = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    damaged = Document(io.BytesIO(data))
    damaged.paragraphs[0].text = '篡改标题'
    out = io.BytesIO(); damaged.save(out)
    with pytest.raises(ContractReviewError, match='contract_report_failed'):
        verify_report(out.getvalue(), expected)
