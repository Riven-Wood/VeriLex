"""Synthetic contracts: exact-range edits, recoverable proposals and package fidelity."""
from __future__ import annotations

import asyncio
import base64
import io
from dataclasses import replace
from types import SimpleNamespace

import pytest
from docx import Document
from lxml import etree

from verilex_runtime.contract_review.editor import ContractEditor, _verify_preserved
from verilex_runtime.contract_review.errors import ContractReviewError
from verilex_runtime.documents.docx_package import W, parts, pack, xml, serialize
from verilex_runtime.litigation.editor import ParagraphEdit
from verilex_runtime.storage.models import TurnRecord, TurnStatus
from tests.contract_review.test_service import _env, NOW


def contract_bytes():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('甲方应当').bold = True
    p.add_run('在十日内支付价款。')
    doc.add_paragraph('甲方应当在十日内支付价款。')
    doc.add_paragraph('😀收件人信息待补充。')
    doc.add_paragraph('')
    doc.add_paragraph('第一页\t备注')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '交付期限'
    table.cell(0, 1).text = '三日'
    doc.sections[0].header.paragraphs[0].text = '虚构合同页眉'
    doc.sections[0].footer.paragraphs[0].text = '虚构合同页脚'
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.fixture
def env(tmp_path):
    service, repo, database, threads, calls, turns = _env(tmp_path)
    source = contract_bytes()
    handle = asyncio.run(service.start_review(
        source, '虚构合同.docx', provider_id='provider-1', model='model-1', skill='通用合同审查',
        party='neutral', strictness='normal', signature='VeriLex', retrieval_enabled=False,
    ))
    return service, repo, database, handle, source, turns


def bind_selection(env, *, index=0, text='十日', turn_id='edit-turn'):
    service, _, _, handle, _, turns = env
    state = service.editor.read(handle['review_id'])
    p = state['paragraphs'][index]
    start = p['text'].index(text)
    context = SimpleNamespace(thread_id=handle['thread_id'], turn_id=turn_id)
    turns.create(TurnRecord(id=turn_id, thread_id=context.thread_id, status=TurnStatus.RUNNING,
                           error_code=None, started_at=NOW, finished_at=None, created_at=NOW))
    payload = {'review_id': handle['review_id'], 'revision': state['revision'], 'paragraph_id': p['id'],
               'start': start, 'end': start + len(text), 'selected_text': text}
    snapshot = service.editor.prepare(context.thread_id, payload)
    service.editor.bind(context.thread_id, turn_id, snapshot)
    return context, payload


def test_current_word_edit_bound_to_version_and_committed_once(env):
    service, repo, _, h, source, turns = env
    doc=repo.create_word_document('word-test','虚构合同.docx',source,source,NOW)
    repo.link_word_document(h['review_id'],doc['id'])
    context=SimpleNamespace(thread_id=h['thread_id'],turn_id='word-turn')
    turns.create(TurnRecord(id=context.turn_id,thread_id=context.thread_id,status=TurnStatus.RUNNING,error_code=None,started_at=NOW,finished_at=None,created_at=NOW))
    selection={'texts':['😀甲方十日内支付'], 'start':0,'end':0,'from':4,'to':6}
    payload={'mode':'word','review_id':h['review_id'],'document_id':doc['id'],'revision':0,'selection':selection,'selected_text':'十日'}
    snapshot=service.editor.prepare(context.thread_id,payload)
    service.editor.bind(context.thread_id,context.turn_id,snapshot)
    result=service.editor.propose(context,{'before':'十日','after':'三十日\n另行通知','reason':'期限调整'})
    assert service.editor.propose(context,{'before':'十日','after':'三十日\n另行通知','reason':'期限调整'})['proposal_id']==result['proposal_id']
    pending=repo.pending_word_ai_edits(h['review_id']);assert len(pending)==1
    assert repo.save_word_document(doc['id'],0,source,NOW,None,result['proposal_id'])
    assert repo.pending_word_ai_edits(h['review_id'])==[]
    assert not repo.save_word_document(doc['id'],0,source,NOW,None,result['proposal_id'])
    with pytest.raises(ContractReviewError):service.editor.prepare(context.thread_id,payload)


def edit_one(editor, rid, state, index, text):
    p = state['paragraphs'][index]
    return editor.save(rid, state['revision'], [{'paragraph_id': p['id'], 'original': p['text'], 'text': text}])


def test_text_edit_export_preserves_original_parts_formatting_tables_and_restart(env):
    service, repo, _, h, source, _ = env
    editor = service.editor
    state = editor.read(h['review_id'])
    state = edit_one(editor, h['review_id'], state, 0, '甲方应当在三十日内支付价款。')
    state = edit_one(editor, h['review_id'], state, 3, '新增的空段文字')
    cell = next(i for i, p in enumerate(state['paragraphs']) if p['text'] == '三日')
    state = edit_one(editor, h['review_id'], state, cell, '五日')
    restored = ContractEditor(service).read(h['review_id'])
    assert restored == state
    assert repo.original_bytes(h['review_id']) == source
    assert repo.review_bytes(h['review_id']) == source
    exported = editor.export(h['review_id'], state['revision'])
    assert exported['filename'].endswith('（修改副本）.docx')
    output = base64.b64decode(exported['content'])
    before, after = parts(source), parts(output)
    assert before.keys() == after.keys()
    assert all(before[k] == after[k] for k in before if k != 'word/document.xml')
    doc = Document(io.BytesIO(output))
    assert doc.paragraphs[0].text == '甲方应当在三十日内支付价款。'
    assert doc.paragraphs[0].runs[0].bold
    assert doc.paragraphs[1].text == '甲方应当在十日内支付价款。'
    assert doc.paragraphs[3].text == '新增的空段文字'
    assert doc.tables[0].cell(0, 1).text == '五日'
    assert doc.sections[0].header.paragraphs[0].text == '虚构合同页眉'


def test_conflict_and_readonly_never_write_any_partial_changes(env):
    service, _, _, h, _, _ = env
    editor = service.editor
    state = editor.read(h['review_id'])
    p, readonly = state['paragraphs'][0], state['paragraphs'][4]
    assert not readonly['editable']
    for changes, code in [
        ([{'paragraph_id': p['id'], 'original': 'wrong', 'text': 'changed'}], 'contract_editor_conflict'),
        ([{'paragraph_id': p['id'], 'original': p['text'], 'text': 'a\nb'}], 'contract_editor_invalid'),
        ([{'paragraph_id': p['id'], 'original': p['text'], 'text': 'safe'},
          {'paragraph_id': readonly['id'], 'original': readonly['text'], 'text': 'unsafe'}], 'contract_editor_readonly'),
    ]:
        with pytest.raises(ContractReviewError) as error:
            editor.save(h['review_id'], 0, changes)
        assert error.value.code == code
        assert editor.read(h['review_id']) == state
    changed = edit_one(editor, h['review_id'], state, 0, 'changed')
    with pytest.raises(ContractReviewError):
        edit_one(editor, h['review_id'], state, 1, 'lost update')
    with pytest.raises(ContractReviewError):
        editor.export(h['review_id'], 0)
    assert editor.read(h['review_id']) == changed


def test_selection_proposal_accept_and_undo_are_scoped_persisted_and_idempotent(env):
    service, _, _, h, _, _ = env
    editor = service.editor
    context, _ = bind_selection(env)
    result = editor.propose(context, {'before': '十日', 'after': '十五日', 'reason': '延长付款时间'})
    assert editor.propose(context, {'before': '十日', 'after': '十五日'})['proposal_id'] == result['proposal_id']
    state = editor.read(h['review_id'])
    assert state['revision'] == 0 and state['proposals'][0]['can_accept']
    # An independent paragraph change does not invalidate this exact target.
    state = edit_one(editor, h['review_id'], state, 2, '收件人：虚构单位')
    assert '😀收件人' not in editor.excerpt(context)  # immutable selection paragraph only
    state = editor.transition(h['review_id'], result['proposal_id'], 'accept', state['revision'])
    assert state['paragraphs'][0]['text'] == '甲方应当在十五日内支付价款。'
    assert state['paragraphs'][1]['text'] == '甲方应当在十日内支付价款。'
    assert state['proposals'][0]['can_undo']
    assert editor.transition(h['review_id'], result['proposal_id'], 'accept', 0) == state
    state = editor.transition(h['review_id'], result['proposal_id'], 'undo', state['revision'])
    assert state['paragraphs'][0]['text'] == state['paragraphs'][0]['original']
    assert state['proposals'][0]['status'] == 'undone'
    assert editor.transition(h['review_id'], result['proposal_id'], 'undo', 0) == state
    with pytest.raises(ContractReviewError):
        editor.propose(context, {'before': '十日', 'after': '二十日'})


def test_rejected_and_touched_then_reverted_targets_cannot_be_reapplied(env):
    service, _, _, h, _, _ = env
    editor = service.editor
    context, _ = bind_selection(env)
    proposal = editor.propose(context, {'before': '十日', 'after': '十五日'})['proposal_id']
    state = editor.read(h['review_id'])
    state = edit_one(editor, h['review_id'], state, 0, '临时更改')
    state = edit_one(editor, h['review_id'], state, 0, state['paragraphs'][0]['original'])
    assert not state['proposals'][0]['can_accept']
    with pytest.raises(ContractReviewError):
        editor.transition(h['review_id'], proposal, 'accept', state['revision'])
    state = editor.transition(h['review_id'], proposal, 'reject', state['revision'])
    assert state['proposals'][0]['status'] == 'rejected'
    with pytest.raises(ContractReviewError):
        editor.propose(context, {'before': '十日', 'after': '二十日'})


def test_snapshot_ownership_unicode_and_original_opinion_boundary(env):
    service, _, _, h, _, turns = env
    editor = service.editor
    context, payload = bind_selection(env, index=2, text='收件人')
    assert payload['start'] == 1  # code points, not UTF-16 offset 2
    assert editor.has_snapshot(context)
    with pytest.raises(ContractReviewError):
        editor.prepare('another-thread', payload)
    with pytest.raises(ContractReviewError):
        editor.prepare(context.thread_id, {**payload, 'start': 2})
    with pytest.raises(ContractReviewError):
        editor.bind('another-thread', context.turn_id, editor.prepare(context.thread_id, payload))
    assert not editor.has_snapshot(SimpleNamespace(thread_id='another-thread', turn_id=context.turn_id))
    with pytest.raises(ContractReviewError):
        service.submit_opinion(context, {'body': '不能写入原件意见'})
    with pytest.raises(ContractReviewError):
        editor.propose(context, {'before': '另一处', 'after': '替换'})
    turns.update(replace(turns.get(context.turn_id), status=TurnStatus.CANCELLED))
    with pytest.raises(ContractReviewError):
        editor.propose(context, {'before': '收件人', 'after': '联系人'})


def test_snapshot_cannot_be_overwritten_and_cascades_with_review(env):
    service, repo, database, h, _, _ = env
    context, payload = bind_selection(env)
    snapshot = service.editor.prepare(context.thread_id, payload)
    with pytest.raises(ContractReviewError):
        service.editor.bind(context.thread_id, context.turn_id, {**snapshot, 'selected_text': 'changed'})
    assert repo.editor_turn(context.thread_id, context.turn_id) == snapshot
    with database.connect() as connection:
        connection.execute('DELETE FROM contract_reviews WHERE id = ?', (h['review_id'],))
        assert connection.execute('SELECT COUNT(*) FROM contract_editor_states').fetchone()[0] == 0
        assert connection.execute('SELECT COUNT(*) FROM contract_editor_turns').fetchone()[0] == 0


def test_independent_verifier_detects_lost_parts_and_unrelated_xml_changes(env):
    _, _, _, _, source, _ = env
    items = parts(source)
    changed = dict(items)
    changed.pop('word/header1.xml')
    with pytest.raises(ContractReviewError):
        _verify_preserved(source, pack(changed), [])
    root = xml(items['word/document.xml'])
    root.find(f'.//{{{W}}}sectPr').set('unintended', 'changed')
    changed = {**items, 'word/document.xml': serialize(root)}
    with pytest.raises(ContractReviewError):
        _verify_preserved(source, pack(changed), [])


def test_opinion_locations_use_original_overlap_and_tab_semantics(env):
    service, repo, _, h, source, _ = env
    def opinion(id, anchor, occurrence=None):
        repo.upsert_opinion(id=id, review_id=h['review_id'], clause_label='条款', clause_excerpt=anchor,
                            anchor_text=anchor, occurrence=occurrence, level='', category='', body='意见',
                            suggestion='', created_at=NOW, updated_at=NOW)
    opinion('ambiguous', '十日')
    opinion('exact', '十日', 1)
    opinion('tab', '页备注')  # must not manufacture a match across a omitted tab
    state = service.editor.read(h['review_id'])
    locations = {p['opinion_id']: p for p in state['opinion_locations']}
    assert locations['ambiguous']['status'] == 'unresolved'
    assert locations['tab']['status'] == 'unresolved'
    assert locations['exact']['paragraph_id'] == state['paragraphs'][1]['id']
    state = edit_one(service.editor, h['review_id'], state, 1, '删除全部文字')
    assert next(p for p in state['opinion_locations'] if p['opinion_id'] == 'exact')['status'] == 'changed'


def test_ai_reference_changes_remain_pending_until_user_adopts(env):
    service, _, _, h, source, _ = env
    before = service.editor.read(h['review_id'])
    context, _ = bind_selection(env)
    result = service.editor.propose(context, {'before': '十日', 'after': '依照《民法典》第五百八十五条约定的期限'})
    pending = service.editor.read(h['review_id'])
    assert pending['paragraphs'] == before['paragraphs']
    assert pending['proposals'][0]['can_accept']
    adopted = service.editor.transition(h['review_id'], result['proposal_id'], 'accept', pending['revision'])
    assert '第五百八十五条' in adopted['paragraphs'][0]['text']


def test_concurrent_autosaves_conflict_atomically_without_lost_changes(env, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    service, repo, _, h, _, _ = env
    state = service.editor.read(h['review_id'])
    original_save = repo.save_editor
    rendezvous = Barrier(2)
    def synchronized_save(*args):
        rendezvous.wait(timeout=5)
        return original_save(*args)
    monkeypatch.setattr(repo, 'save_editor', synchronized_save)
    def save(index):
        try:
            return edit_one(service.editor, h['review_id'], state, index, f'用户改稿{index}')
        except ContractReviewError as error:
            return error
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(save, [0, 1]))
    assert sum(isinstance(r, dict) for r in results) == 1
    assert sum(isinstance(r, ContractReviewError) and r.code == 'contract_editor_conflict' for r in results) == 1
    result = service.editor.read(h['review_id'])
    assert result['revision'] == 1
    assert sum(p['text'].startswith('用户改稿') for p in result['paragraphs']) == 1


def test_export_failure_from_corrupt_adapter_never_persists_state(env, monkeypatch):
    service, _, _, h, source, _ = env
    state = service.editor.read(h['review_id'])
    from verilex_runtime.contract_review import editor as module
    real_patch = module.patch_docx
    def corrupt(*args):
        data = parts(real_patch(*args))
        data['word/footer1.xml'] = data['word/footer1.xml'].replace('虚构'.encode(), '错误'.encode())
        return pack(data)
    monkeypatch.setattr(module, 'patch_docx', corrupt)
    with pytest.raises(ContractReviewError) as error:
        edit_one(service.editor, h['review_id'], state, 0, '更改正文')
    assert error.value.code == 'contract_editor_integrity'
    assert service.editor.read(h['review_id']) == state


def test_opinion_occurrence_overlap_does_not_shift_to_another_paragraph(env):
    service, repo, database, h, source, _ = env
    document = Document()
    document.add_paragraph('aaaa')
    document.add_paragraph('aaa')
    out = io.BytesIO(); document.save(out)
    with database.connect() as connection:
        connection.execute('UPDATE contract_reviews SET original = ? WHERE id = ?', (out.getvalue(), h['review_id']))
    repo.upsert_opinion(id='overlap', review_id=h['review_id'], clause_label='测试', clause_excerpt='aaa',
                        anchor_text='aaa', occurrence=1, level='', category='', body='意见', suggestion='',
                        created_at=NOW, updated_at=NOW)
    state = service.editor.read(h['review_id'])
    assert state['opinion_locations'][0]['paragraph_id'] == state['paragraphs'][0]['id']


def test_article_digit_only_selection_is_a_normal_pending_edit(env):
    service, _, _, h, _, _ = env
    state = service.editor.read(h['review_id'])
    state = edit_one(service.editor, h['review_id'], state, 0, '依照《民法典》第五百八十五条支付。')
    context, _ = bind_selection(env, text='五百八十五')
    service.editor.propose(context, {'before': '五百八十五', 'after': '五百八十六'})
    assert service.editor.read(h['review_id'])['paragraphs'] == state['paragraphs']


def test_complex_word_parts_comments_images_revisions_retained_and_readonly(env):
    service, repo, database, h, _, _ = env
    document = Document()
    document.add_paragraph('可以修改这一段。')
    commented = document.add_paragraph('原有批注文字。')
    document.add_comment(commented.runs, text='保留的原有意见', author='测试作者')
    # Deterministic 1×1 PNG; not a real user's document or image.
    png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl6CksAAAAASUVORK5CYII=')
    document.add_picture(io.BytesIO(png))
    revised = document.add_paragraph()
    insertion = etree.SubElement(revised._p, f'{{{W}}}ins')
    insertion.set(f'{{{W}}}id', '1')
    run = etree.SubElement(insertion, f'{{{W}}}r')
    etree.SubElement(run, f'{{{W}}}t').text = '保留的修订文字'
    out = io.BytesIO(); document.save(out)
    source = out.getvalue()
    with database.connect() as connection:
        connection.execute('UPDATE contract_reviews SET original = ? WHERE id = ?', (source, h['review_id']))
    state = service.editor.read(h['review_id'])
    assert set(state['retained_objects']) >= {'批注', '图片', '修订'}
    assert all(not p['editable'] for p in state['paragraphs'][1:])
    state = edit_one(service.editor, h['review_id'], state, 0, '仅修改正文这一段。')
    output = base64.b64decode(service.editor.export(h['review_id'], state['revision'])['content'])
    original_parts, edited_parts = parts(source), parts(output)
    assert original_parts.keys() == edited_parts.keys()
    assert all(v == edited_parts[k] for k, v in original_parts.items() if k != 'word/document.xml')
    assert repo.original_bytes(h['review_id']) == source
    assert len(Document(io.BytesIO(output)).comments) == 1
    assert len(Document(io.BytesIO(output)).inline_shapes) == 1


@pytest.mark.asyncio
async def test_registered_edit_tool_uses_bound_snapshot_and_redacted_errors(env):
    import json
    from verilex_runtime.contract_review.tools import register_contract_review_tools
    from verilex_runtime.tools.registry import ToolRegistry, ToolError
    service, _, _, h, _, _ = env
    context, _ = bind_selection(env)
    registry = ToolRegistry()
    register_contract_review_tools(registry, service)
    tool = registry.get('contract_edit_proposal')
    result = await tool.execute(context, {'before': '十日', 'after': '二十日'})
    assert result.ok and json.loads(result.output)['proposal_id']
    assert service.editor.read(h['review_id'])['paragraphs'][0]['text'].endswith('十日内支付价款。')
    with pytest.raises(ToolError) as error:
        await tool.execute(SimpleNamespace(thread_id='wrong', turn_id=context.turn_id), {'before': '秘密测试词', 'after': '错误目标'})
    assert error.value.code == 'contract_editor_selection'
    assert '秘密测试词' not in error.value.message
