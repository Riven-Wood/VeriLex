import base64
import io

from docx import Document

from verilex_runtime.contract_review.preview import build_preview


def sample():
    doc = Document()
    doc.add_paragraph('同一条款')
    p = doc.add_paragraph('同一条款')
    doc.add_comment(p.runs, '原有意见 <script>不执行</script>', '原作者')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格付款条款'
    doc.sections[0].header.paragraphs[0].text = '页眉'
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_preview_preserves_bytes_and_distinguishes_existing_comments():
    data = sample()
    result = build_preview(data)
    assert base64.b64decode(result['content']) == data
    assert result['paragraphs'] == ['同一条款', '同一条款', '表格付款条款']
    comment = result['comments'][0]
    assert comment['author'] == '原作者'
    assert comment['text'] == '原有意见 <script>不执行</script>'
    assert comment['paragraph_indices'] == [1]


def test_preview_empty_comments():
    doc = Document()
    doc.add_paragraph('正文')
    out = io.BytesIO()
    doc.save(out)
    assert build_preview(out.getvalue())['comments'] == []
