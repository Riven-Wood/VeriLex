"""虚构合同样本构造器：测试共用的脱敏结构（公司、人名、编号均为虚构）。"""
from __future__ import annotations

from pathlib import Path

from docx import Document


def build_contract(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("技术服务合同").style = doc.styles["Title"]
    doc.add_paragraph("合同编号：VL-2026-0901")
    doc.add_paragraph("甲方：示例科技有限公司")
    doc.add_paragraph("乙方：演示咨询工作室")
    doc.add_paragraph(
        "第一条 服务内容：乙方为甲方提供数据平台运维技术服务，"
        "服务范围以附件一《服务清单》为准。"
    )
    doc.add_paragraph("第二条 合同总价款为人民币叁拾万元整（¥300,000），包含全部税费。")
    doc.add_paragraph(
        "第三条 付款方式：合同生效后十个工作日内，甲方支付合同总价款的百分之四十作为预付款；"
    )
    doc.add_paragraph("验收合格后十个工作日内支付剩余百分之六十。")
    doc.add_paragraph("第四条 交付期限：乙方应于2026年11月30日前完成全部服务内容并通过甲方验收。")
    doc.add_paragraph(
        "第五条 保密义务：双方对在履行本合同过程中知悉的对方商业秘密承担保密责任，"
        "保密期限为合同终止后三年。"
    )
    doc.add_paragraph(
        "第六条 违约责任：任何一方违反本合同约定给对方造成损失的，应承担赔偿责任；"
        "乙方逾期交付的，每逾期一日按合同总价款的千分之三向甲方支付违约金。"
    )
    table = doc.add_table(rows=2, cols=3)
    for column, value in enumerate(("付款阶段", "比例", "金额（元）")):
        table.cell(0, column).text = value
    for column, value in enumerate(("预付款", "40%", "120,000")):
        table.cell(1, column).text = value
    doc.add_paragraph(
        "第七条 争议解决：因本合同引起的争议，双方应友好协商解决；"
        "协商不成的，提交北京仲裁委员会仲裁。"
    )
    doc.add_paragraph("第八条 本合同一式两份，双方各执一份，自双方签字盖章之日起生效。")
    doc.add_paragraph("甲方：示例科技有限公司（盖章）")
    doc.add_paragraph("乙方：演示咨询工作室（盖章）")
    doc.sections[0].header.paragraphs[0].text = "技术服务合同（演示样本）"
    doc.sections[0].footer.paragraphs[0].text = "第 1 页"
    doc.save(path)
    return path
