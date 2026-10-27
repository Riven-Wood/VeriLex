"""Opt-in real macOS converter checks (no mocked office process or sandbox).

VERILEX_RUN_DOC_CONVERTER_TESTS=1 runtime/.venv/bin/python -m pytest <this file>
Requires installed LibreOffice and permission to start a nested sandbox.
"""
from pathlib import Path
import base64
import io
import os
import socket
import subprocess
import tempfile

import pytest
from docx import Document

from verilex_runtime.attachments.legacy_doc import _sandbox_policy
from verilex_runtime.attachments.inspection import inspect_file
from verilex_runtime.contract_review.preview import build_preview
from verilex_runtime.documents.comment_verify import verify_comment_copy
from verilex_runtime.file_redaction.rules import detect_rules
from verilex_runtime.file_redaction.service import RedactionService
from verilex_runtime.storage.database import Database
from verilex_runtime.storage.migrations import run_migrations
from tests.contract_review.test_service import _env
from tests.contract_review.test_export import _opinion

pytestmark = pytest.mark.skipif(
    os.environ.get("VERILEX_RUN_DOC_CONVERTER_TESTS") != "1",
    reason="opt-in real macOS LibreOffice/sandbox test",
)
FIXTURE = Path(__file__).resolve().parents[1] / "fixtures/legacy-word.doc"


def test_sandbox_allows_job_but_blocks_other_files_writes_and_network():
    repo = Path(__file__).resolve().parents[3]
    with tempfile.TemporaryDirectory(prefix="vls-", dir="/private/tmp") as job, \
            tempfile.TemporaryDirectory(prefix="vls-out-", dir="/private/tmp") as outside, \
            tempfile.TemporaryDirectory(prefix=".vls-", dir=repo) as user_job:
        root = Path(job).resolve()
        own = root / "profile/user/registrymodifications.xcu"
        own.parent.mkdir(parents=True)
        own.write_text("job-private")
        blocked = [Path(outside) / "other-job.txt", Path(user_job) / "synthetic-user.txt"]
        for p in blocked:
            p.write_text("must-not-read")
        policy = root / "policy.sb"
        policy.write_text(_sandbox_policy(root))
        def run(*args):
            return subprocess.run(["/usr/bin/sandbox-exec", "-f", str(policy), *args],
                                  cwd=root, capture_output=True, timeout=5)
        allowed = run("/bin/cat", str(own))
        assert allowed.returncode == 0 and allowed.stdout == b"job-private"
        for p in blocked:
            denied = run("/bin/cat", str(p))
            assert denied.returncode != 0 and b"must-not-read" not in denied.stdout
            assert run("/usr/bin/touch", str(p)).returncode != 0
            assert p.read_text() == "must-not-read"
        with socket.socket() as listener:
            listener.bind(("127.0.0.1", 0))
            listener.listen()
            assert run("/usr/bin/nc", "-z", "-w", "1", "127.0.0.1", str(listener.getsockname()[1])).returncode != 0


@pytest.mark.asyncio
async def test_real_doc_contract_preview_comments_export_and_original(tmp_path):
    raw = FIXTURE.read_bytes()
    service, repository, *_ = _env(tmp_path)
    handle = await service.start_review(
        raw, "旧版虚构合同.doc", provider_id=None, model="model-1", skill="通用合同审查",
        party="neutral", strictness="normal", signature="审查测试", retrieval_enabled=False,
    )
    review_id = handle["review_id"]
    baseline = repository.review_bytes(review_id)
    preview = build_preview(baseline)
    assert base64.b64decode(preview["content"]) == baseline
    assert any("未约定验收期限" in p for p in preview["paragraphs"])
    assert preview["comments"][0]["text"] == "原有批注：请核对合同名称。"
    doc = Document(io.BytesIO(baseline))
    assert doc.tables[0].cell(1, 1).text == "10000元"
    assert doc.sections[0].header.paragraphs[0].text == "虚构合同 · 格式兼容检查"
    assert len(doc.inline_shapes) == 1
    assert any(run.bold for run in doc.paragraphs[0].runs)
    _opinion(repository, review_id, "doc-real-opinion", anchor_text="千分之三违约金")
    result = service.export(review_id)
    output = base64.b64decode(result["content"])
    assert result["filename"] == "旧版虚构合同（批注版）.docx"
    assert verify_comment_copy(baseline, output).passed
    assert len(Document(io.BytesIO(output)).comments) == 2
    assert repository.original_bytes(review_id) == raw == FIXTURE.read_bytes()


def test_real_doc_attachment_and_redaction(tmp_path):
    raw = FIXTURE.read_bytes()
    upload = tmp_path / "upload.part"
    upload.write_bytes(raw)
    inspected = inspect_file(upload, "虚构合同.doc", "application/msword")
    assert "未约定验收期限" in inspected.extracted_text
    assert inspected.media_type == "application/msword" and upload.read_bytes() == raw
    db = Database(tmp_path / "redaction.sqlite3")
    run_migrations(db)
    service = RedactionService(db, detector=detect_rules)
    job = service.create("虚构合同.doc", raw, [], ["张三"])
    output = service.export(job["id"], job["revision"], "placeholder")
    doc = Document(io.BytesIO(output))
    assert "张三" not in "\n".join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 1).text == "10000元"
    assert len(doc.inline_shapes) == 1
