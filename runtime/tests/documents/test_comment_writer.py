# runtime/tests/documents/test_comment_writer.py
"""Slice-1 comment writer tests. All fixtures are synthetic; no case data."""
from pathlib import Path

import pytest
from docx import Document

from verilex_runtime.documents.comments import (
    CommentOutcome,
    CommentSpec,
    write_comments,
)
from verilex_runtime.documents.errors import DocumentDomainError


def _contract(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("第一条 甲方应于收到发票后三十日内支付全部合同价款。")
    doc.add_paragraph("第二条 双方应对本合同内容严格保密。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "价款"
    table.cell(0, 1).text = "人民币壹拾万元整"
    p = tmp_path / "contract.docx"
    doc.save(p)
    return p


def test_write_comments_anchors_and_returns_copy(tmp_path):
    source = _contract(tmp_path)
    before = source.read_bytes()
    outcome = write_comments(
        source,
        [
            CommentSpec(
                anchor_text="收到发票后三十日内支付",
                comment_text="【高】必须修改：付款期限建议明确起算日。",
                author="VeriLex",
                initials="VL",
            )
        ],
    )
    assert isinstance(outcome, CommentOutcome)
    assert len(outcome.comment_ids) == 1
    assert outcome.copy_bytes[:2] == b"PK"
    assert source.read_bytes() == before  # original untouched

    out = tmp_path / "out.docx"
    out.write_bytes(outcome.copy_bytes)
    doc = Document(str(out))
    assert len(doc.comments) == 1
    comment = doc.comments.get(0)
    assert comment.author == "VeriLex"
    assert "付款期限" in comment.text
    assert "第一条" in doc.paragraphs[0].text  # body intact


def test_anchor_inside_table_cell(tmp_path):
    source = _contract(tmp_path)
    outcome = write_comments(
        source,
        [CommentSpec("壹拾万元整", "金额核对", "甲", "J")],
    )
    out = tmp_path / "t.docx"
    out.write_bytes(outcome.copy_bytes)
    assert len(Document(str(out)).comments) == 1


def test_partial_run_anchor_splits_runs(tmp_path):
    doc = Document()
    doc.add_paragraph("违约金为合同总价款的百分之二十。")
    source = tmp_path / "p.docx"
    doc.save(source)
    outcome = write_comments(
        source, [CommentSpec("总价款的百分之二十", "违约金过高", "乙", "Y")]
    )
    out = tmp_path / "p_out.docx"
    out.write_bytes(outcome.copy_bytes)
    opened = Document(str(out))
    assert len(opened.comments) == 1
    assert opened.paragraphs[0].text == "违约金为合同总价款的百分之二十。"


def test_multiple_comments_get_sequential_ids(tmp_path):
    source = _contract(tmp_path)
    outcome = write_comments(
        source,
        [
            CommentSpec("严格保密", "保密期建议", "甲", "J"),
            CommentSpec("三十日内支付", "期限起算", "甲", "J"),
        ],
    )
    assert outcome.comment_ids == tuple(sorted(outcome.comment_ids))


def test_anchor_not_found_rejects_whole_batch(tmp_path):
    source = _contract(tmp_path)
    with pytest.raises(DocumentDomainError) as err:
        write_comments(
            source,
            [
                CommentSpec("严格保密", "a", "甲", "J"),
                CommentSpec("不存在的锚定文字", "b", "甲", "J"),
            ],
        )
    assert err.value.code == "docx_comment_anchor_not_found"


def test_anchor_with_revisions_rejected(tmp_path):
    doc = Document()
    para = doc.add_paragraph("本合同经双方签署后生效。")
    run = para.runs[0]._r
    ins = run.makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ins",
        {"id": "1"},
    )
    run.addprevious(ins)
    ins.append(run)
    source = tmp_path / "rev.docx"
    doc.save(source)
    with pytest.raises(DocumentDomainError) as err:
        write_comments(source, [CommentSpec("双方签署", "x", "甲", "J")])
    assert err.value.code == "docx_comment_revisions_present"


def test_limitations_are_reported(tmp_path):
    outcome = write_comments(
        _contract(tmp_path), [CommentSpec("严格保密", "a", "甲", "J")]
    )
    assert any("回复" in line for line in outcome.limitations)
    assert any("页眉" in line for line in outcome.limitations)


@pytest.mark.parametrize('existing_comment', [False, True])
def test_comment_preserves_noncanonical_unrelated_xml_bytes(tmp_path, existing_comment):
    from verilex_runtime.documents.comments import write_verified_comments
    from verilex_runtime.documents.docx_package import parts, pack
    doc = Document()
    p = doc.add_paragraph('虚构合同名称')
    if existing_comment:
        doc.add_comment(p.runs, '原有批注', '原作者')
    doc.add_paragraph('验收合格后付款。')
    doc.sections[0].header.paragraphs[0].text = '虚构页眉'
    source = tmp_path / 'foreign-format.docx'
    doc.save(source)
    original = parts(source.read_bytes())
    # Valid XML with a serialization layout different from python-docx's save.
    for name in ('word/header1.xml', 'word/styles.xml', 'docProps/core.xml'):
        original[name] = original[name].replace(b'?>', b'?>\n\n', 1)
    source.write_bytes(pack(original))
    outcome = write_verified_comments(source, [CommentSpec('验收合格', '明确验收期限', '测试', 'CS')])
    result = parts(outcome.copy_bytes)
    for name, data in original.items():
        if name not in {'word/document.xml', 'word/comments.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}:
            assert result[name] == data, name
    assert len(Document(__import__('io').BytesIO(outcome.copy_bytes)).comments) == 1 + existing_comment
