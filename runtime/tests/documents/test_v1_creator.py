from __future__ import annotations

import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from verilex_runtime.documents.v1_creator import create_v1_docx
from verilex_runtime.documents.errors import DocumentDomainError
from verilex_runtime.documents.v1_markdown import MarkdownBlock


def test_create_docx_uses_v1_fixed_layout(tmp_path):
    output = tmp_path / "result.docx"

    report = create_v1_docx(
        output,
        title="合同审查意见",
        blocks=(
            MarkdownBlock("heading", "一、结论", 1),
            MarkdownBlock("heading", "（一）事实", 2),
            MarkdownBlock("heading", "1. 说明", 3),
            MarkdownBlock("paragraph", "这是正文。"),
            MarkdownBlock("bullet", "第一项"),
            MarkdownBlock("numbered", "第一步"),
            MarkdownBlock("pagebreak", ""),
        ),
    )

    document = Document(output)
    section = document.sections[0]
    assert report.block_count == 7
    assert report.text_char_count == 31
    one_twip = 635
    assert abs(section.page_width - Mm(210)) <= one_twip
    assert abs(section.page_height - Mm(297)) <= one_twip
    assert abs(section.top_margin - Mm(25.4)) <= one_twip
    assert abs(section.bottom_margin - Mm(25.4)) <= one_twip
    assert abs(section.left_margin - Mm(25.4)) <= one_twip
    assert abs(section.right_margin - Mm(25.4)) <= one_twip
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "合同审查意见",
        "一、结论",
        "（一）事实",
        "1. 说明",
        "这是正文。",
        "第一项",
        "第一步",
        "",
    ]


def test_create_docx_applies_v1_typography_and_page_break(tmp_path):
    output = tmp_path / "typography.docx"
    create_v1_docx(
        output,
        title="法律意见书",
        blocks=(
            MarkdownBlock("heading", "一级标题", 1),
            MarkdownBlock("heading", "二级标题", 2),
            MarkdownBlock("heading", "三级标题", 3),
            MarkdownBlock("paragraph", "中文正文 English text"),
            MarkdownBlock("bullet", "项目"),
            MarkdownBlock("numbered", "步骤"),
            MarkdownBlock("pagebreak", ""),
        ),
    )

    document = Document(output)
    title, heading_1, heading_2, heading_3, body, bullet, numbered, pagebreak = (
        document.paragraphs
    )
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].bold is True
    assert title.runs[0].font.size == Pt(18)
    assert heading_1.runs[0].font.size == Pt(14)
    assert heading_2.runs[0].font.size == Pt(13)
    assert heading_3.runs[0].font.size == Pt(12)
    assert body.runs[0].font.size == Pt(12)
    assert body.paragraph_format.line_spacing == 1.5
    assert body.paragraph_format.first_line_indent == Pt(24)
    assert bullet.style.name == "List Bullet"
    assert numbered.style.name == "List Number"
    for paragraph in document.paragraphs[:-1]:
        fonts = paragraph.runs[0]._element.get_or_add_rPr().rFonts
        assert fonts.get(qn("w:eastAsia")) == "宋体"
        assert fonts.get(qn("w:ascii")) == "Times New Roman"
    assert pagebreak._element.xpath('.//w:br[@w:type="page"]')


def test_validation_report_is_canonical_artifact_json(tmp_path):
    report = create_v1_docx(
        tmp_path / "report.docx",
        title="报告",
        blocks=(MarkdownBlock("paragraph", "正文"),),
    )

    assert json.loads(report.as_json()) == {
        "schema_version": 1,
        "structure": {"status": "passed", "issue_counts": {}},
        "render": {"status": "not_required"},
        "vision": {"status": "not_run"},
    }


def test_create_docx_reopen_guard_rejects_lost_list_structure(tmp_path, monkeypatch):
    from verilex_runtime.documents import v1_creator

    def document_with_damaged_reopen(path=None):
        document = Document(path)
        if path is not None:
            document.paragraphs[1].style = "Normal"
        return document

    monkeypatch.setattr(v1_creator, "Document", document_with_damaged_reopen)

    with pytest.raises(DocumentDomainError, match="document_validation_failed"):
        create_v1_docx(
            tmp_path / "damaged.docx",
            title="报告",
            blocks=(MarkdownBlock("bullet", "项目"),),
        )


def test_exact_initial_title_heading_is_rendered_only_once(tmp_path):
    path = tmp_path / 'single-title.docx'
    report = create_v1_docx(path, '申请书', (
        MarkdownBlock('heading', '申请书', 1),
        MarkdownBlock('paragraph', '申请人：待补'),
        MarkdownBlock('heading', '申请事项', 1),
        MarkdownBlock('paragraph', '申请书'),
    ))
    assert [p.text for p in Document(path).paragraphs] == ['申请书', '申请人：待补', '申请事项', '申请书']
    assert report.block_count == 3


def test_different_initial_heading_is_preserved(tmp_path):
    path = tmp_path / 'different-title.docx'
    create_v1_docx(path, '申请书', (MarkdownBlock('heading', '申请书使用说明', 1),))
    assert [p.text for p in Document(path).paragraphs] == ['申请书', '申请书使用说明']
