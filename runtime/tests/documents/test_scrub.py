# runtime/tests/documents/test_scrub.py
import io
import zipfile
from pathlib import Path

from docx import Document
import pytest

from verilex_runtime.documents.scrub import ScrubOutcome, scrub_package

_DC = "http://purl.org/dc/elements/1.1/"
_EP = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"


@pytest.mark.parametrize('external', [False, True])
def test_editor_relationships_are_removed_even_when_part_is_renamed(tmp_path, external):
    from lxml import etree
    from verilex_runtime.documents.docx_package import parts, pack, xml, serialize, REL, CT
    from verilex_runtime.documents.scrub import remove_editor_metadata, EDITOR_RELATIONSHIP
    from verilex_runtime.documents.scrub_verify import audit_processed_copy
    items = parts(scrub_package(_source(tmp_path)).copy_bytes)
    body = items['word/document.xml']
    rels = xml(items['_rels/.rels'])
    relation = etree.SubElement(rels, f'{{{REL}}}Relationship')
    relation.set('Id', 'privateEditor')
    relation.set('Type', EDITOR_RELATIONSHIP)
    relation.set('Target', 'https://example.invalid/editor.json' if external else 'private/editor.json')
    if external:
        relation.set('TargetMode', 'External')
    else:
        items['private/editor.json'] = b'{"before":"private discarded draft"}'
        ct = xml(items['[Content_Types].xml'])
        entry = etree.SubElement(ct, f'{{{CT}}}Override')
        entry.set('PartName', '/private/editor.json')
        entry.set('ContentType', 'application/json')
        items['[Content_Types].xml'] = serialize(ct)
    items['_rels/.rels'] = serialize(rels)
    source = pack(items)
    assert not audit_processed_copy(source, [], scrubbed=True).passed
    for clean in (remove_editor_metadata(source), scrub_package(source).copy_bytes):
        result = parts(clean)
        assert 'private/editor.json' not in result
        assert result['word/document.xml'] == body
        assert audit_processed_copy(clean, [], scrubbed=True).passed
        assert remove_editor_metadata(clean) == clean


def _source(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("正文段落。")
    doc.core_properties.author = "张律师"
    doc.core_properties.last_modified_by = "李助理"
    path = tmp_path / "s.docx"
    doc.save(path)
    return path


def _part(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name)


def _rewrite(data: bytes, *, replace: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, replace.get(item.filename, zin.read(item.filename)))
    return out.getvalue()


def test_scrubs_core_properties(tmp_path):
    outcome = scrub_package(_source(tmp_path))
    assert isinstance(outcome, ScrubOutcome)
    core = _part(outcome.copy_bytes, "docProps/core.xml")
    assert "张律师".encode() not in core and "李助理".encode() not in core
    assert any(item.target == "docProps/core.xml" and item.action == "cleaned" for item in outcome.report)


def test_removes_comment_parts_and_markers(tmp_path):
    from verilex_runtime.documents.comments import CommentSpec, write_comments
    commented = write_comments(
        _source(tmp_path), [CommentSpec("正文段落", "意见", "某人", "SR")]
    ).copy_bytes
    staged = tmp_path / "staged.docx"
    staged.write_bytes(commented)
    outcome = scrub_package(staged)
    with zipfile.ZipFile(io.BytesIO(outcome.copy_bytes)) as z:
        names = set(z.namelist())
    assert not any("comments" in n or "people.xml" in n for n in names)
    body = _part(outcome.copy_bytes, "word/document.xml")
    assert b"commentRangeStart" not in body and b"commentReference" not in body


def test_removes_rsids_and_track_changes_flag(tmp_path):
    source = _source(tmp_path)
    with zipfile.ZipFile(source) as z:
        settings = z.read("word/settings.xml").decode()
        document = z.read("word/document.xml").decode()
    if "rsids" not in settings:
        settings = settings.replace(
            "</w:settings>",
            '<w:rsids><w:rsid w:val="00112233"/></w:rsids><w:trackChanges/></w:settings>',
        )
    document = document.replace("<w:p>", '<w:p w:rsidR="00112233">')
    staged = tmp_path / "rsid.docx"
    staged.write_bytes(
        _rewrite(
            source.read_bytes(),
            replace={
                "word/settings.xml": settings.encode(),
                "word/document.xml": document.encode(),
            },
        )
    )
    outcome = scrub_package(staged)
    assert b"rsid" not in _part(outcome.copy_bytes, "word/settings.xml")
    assert b"trackChanges" not in _part(outcome.copy_bytes, "word/settings.xml")
    assert b"rsidR" not in _part(outcome.copy_bytes, "word/document.xml")


def test_original_untouched_and_absent_items_reported(tmp_path):
    source = _source(tmp_path)
    before = source.read_bytes()
    outcome = scrub_package(source)
    assert source.read_bytes() == before
    assert any(item.action == "absent" for item in outcome.report)  # e.g. custom.xml absent


def test_output_reopens(tmp_path):
    outcome = scrub_package(_source(tmp_path))
    out = tmp_path / "scrubbed.docx"
    out.write_bytes(outcome.copy_bytes)
    assert Document(str(out)).paragraphs[0].text == "正文段落。"
