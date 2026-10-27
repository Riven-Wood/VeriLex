# runtime/tests/documents/test_docx_slices_e2e.py
from pathlib import Path

import pytest
from docx import Document

from verilex_runtime.attachments.errors import AttachmentDomainError
from verilex_runtime.documents.comments import CommentSpec, write_verified_comments
from verilex_runtime.documents.errors import DocumentDomainError
from verilex_runtime.documents.replace import TextReplacement, process_copy


def _contract(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("甲方：张三。电话13800138000。")
    p = tmp_path / "c.docx"
    doc.save(p)
    return p


def test_verified_comments_happy_path(tmp_path):
    outcome = write_verified_comments(
        _contract(tmp_path), [CommentSpec("张三", "主体核验", "VeriLex", "VL")]
    )
    assert outcome.comment_ids and outcome.copy_bytes[:2] == b"PK"


def test_verified_comments_blocked_when_verification_fails(tmp_path, monkeypatch):
    import verilex_runtime.documents.comments as comments_mod
    from verilex_runtime.documents import comment_verify

    original = comment_verify.verify_comment_copy
    def sabotaged(orig, commented):
        from verilex_runtime.documents.comment_verify import VerificationResult
        return VerificationResult(False, ("semantic-diff:word/document.xml",))
    monkeypatch.setattr(comments_mod, "verify_comment_copy", sabotaged)
    with pytest.raises(DocumentDomainError) as err:
        write_verified_comments(
            _contract(tmp_path), [CommentSpec("张三", "x", "VeriLex", "VL")]
        )
    assert err.value.code == "docx_comment_verify_failed"


def test_process_copy_full_pipeline(tmp_path):
    outcome = process_copy(
        _contract(tmp_path),
        [TextReplacement("张三", "【人名1】"), TextReplacement("13800138000", "【电话1】")],
    )
    assert outcome.audit.passed
    assert outcome.counts == (1, 1)
    out = tmp_path / "out.docx"
    out.write_bytes(outcome.copy_bytes)
    assert Document(str(out)).paragraphs[0].text == "甲方：【人名1】。电话【电话1】。"


def test_guard_rejects_non_docx(tmp_path):
    fake = tmp_path / "fake.docx"
    fake.write_bytes(b"not a zip")
    with pytest.raises(AttachmentDomainError):
        write_verified_comments(fake, [CommentSpec("x", "y", "a", "A")])
    with pytest.raises(AttachmentDomainError):
        process_copy(fake, [TextReplacement("x", "y")])


def test_original_hash_unchanged_through_both_slices(tmp_path):
    import hashlib
    source = _contract(tmp_path)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    write_verified_comments(source, [CommentSpec("张三", "x", "VeriLex", "VL")])
    process_copy(source, [TextReplacement("张三", "【人名1】")])
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
