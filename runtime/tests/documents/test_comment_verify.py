# runtime/tests/documents/test_comment_verify.py
"""Independent verifier for slice 1. Adversarial matrix is the core."""
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from verilex_runtime.documents.comment_verify import verify_comment_copy
from verilex_runtime.documents.comments import CommentSpec, write_comments


@pytest.fixture()
def contract(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("第一条 甲方应于收到发票后三十日内支付全部合同价款。")
    doc.add_paragraph("第二条 双方应对本合同内容严格保密。")
    doc.add_paragraph("第三条 本合同一式两份。")
    p = tmp_path / "contract.docx"
    doc.save(p)
    return p


def _commented(contract: Path, anchor: str = "三十日内支付") -> bytes:
    return write_comments(
        contract, [CommentSpec(anchor, "意见", "VeriLex", "VL")]
    ).copy_bytes


def test_clean_comment_copy_passes(contract):
    result = verify_comment_copy(contract, _commented(contract))
    assert result.passed and result.differences == ()


def test_partial_run_anchor_copy_passes_verification(tmp_path):
    # mid-run anchor forces the writer to split runs (deepcopy keeps rPr);
    # the verifier must coalesce them back or a legitimate copy would fail
    doc = Document()
    doc.add_paragraph("违约金为合同总价款的百分之二十。")
    p = tmp_path / "mid.docx"
    doc.save(p)
    result = verify_comment_copy(p, _commented(p, anchor="总价款的百分之二十"))
    assert result.passed, result.differences


def test_tampered_body_text_detected(contract):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out, "w"
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace("三十日内".encode(), "二十日内".encode())
            zout.writestr(item, data)
    result = verify_comment_copy(contract, out.getvalue())
    assert not result.passed
    assert any("word/document.xml" in d for d in result.differences)


def test_extra_part_detected(contract):
    out_io = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out_io, "w"
    ) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/evil.txt", b"payload")
    result = verify_comment_copy(contract, out_io.getvalue())
    assert not result.passed and any("word/evil.txt" in d for d in result.differences)


def test_missing_part_detected(contract):
    out_io = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out_io, "w"
    ) as zout:
        for item in zin.infolist():
            if item.filename == "word/settings.xml":
                continue
            zout.writestr(item, zin.read(item.filename))
    result = verify_comment_copy(contract, out_io.getvalue())
    assert not result.passed and any("word/settings.xml" in d for d in result.differences)


def test_changed_style_part_detected(contract):
    with zipfile.ZipFile(contract) as z:
        styles = z.read("word/styles.xml")
    out_io = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out_io, "w"
    ) as zout:
        for item in zin.infolist():
            zout.writestr(
                item,
                styles + b" " if item.filename == "word/styles.xml" else zin.read(item.filename),
            )
    result = verify_comment_copy(contract, out_io.getvalue())
    assert not result.passed and any("word/styles.xml" in d for d in result.differences)


def test_leftover_comment_marker_without_part_detected(contract):
    """Delete comments.xml but keep markers: must fail, not pass by luck."""
    out_io = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out_io, "w"
    ) as zout:
        for item in zin.infolist():
            if item.filename == "word/comments.xml":
                continue
            zout.writestr(item, zin.read(item.filename))
    result = verify_comment_copy(contract, out_io.getvalue())
    assert not result.passed


def test_injected_revision_in_copy_detected(contract):
    import re
    out_io = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out_io, "w"
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(
                    "第三条".encode(),
                    "<w:ins w:id=\"99\"><w:r><w:t>第三条</w:t></w:r></w:ins>".encode(),
                )
            zout.writestr(item, data)
    result = verify_comment_copy(contract, out_io.getvalue())
    assert not result.passed


def test_paragraph_reorder_detected(contract):
    out_io = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_commented(contract))) as zin, zipfile.ZipFile(
        out_io, "w"
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                first = "第一条 甲方应于收到发票后三十日内支付全部合同价款。".encode()
                second = "第二条 双方应对本合同内容严格保密。".encode()
                from lxml import etree
                root = etree.fromstring(data)
                body = root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
                first_p, second_p = body[0], body[1]
                body.remove(second_p)
                body.insert(0, second_p)
                data = etree.tostring(root)
            zout.writestr(item, data)
    result = verify_comment_copy(contract, out_io.getvalue())
    assert not result.passed
