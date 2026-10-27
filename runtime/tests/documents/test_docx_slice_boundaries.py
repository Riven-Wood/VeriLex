"""Supported Word structures and adversarial outputs; fixtures contain no user data."""
import io
from pathlib import Path
import zipfile
import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from verilex_runtime.documents.comments import CommentSpec, write_comments, write_verified_comments
from verilex_runtime.documents.comment_verify import verify_comment_copy
from verilex_runtime.documents.replace import TextReplacement, replace_texts, process_copy
from verilex_runtime.documents.scrub import scrub_package
from verilex_runtime.documents.scrub_verify import audit_processed_copy
from verilex_runtime.documents.revisions import accept_revisions
from verilex_runtime.documents.errors import DocumentDomainError
from verilex_runtime.documents.docx_package import W, CT, REL, NS


def unpack(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z: return {n: z.read(n) for n in z.namelist()}


def pack(items):
    b = io.BytesIO()
    with zipfile.ZipFile(b, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for n, d in items.items(): z.writestr(n, d)
    return b.getvalue()


def source(tmp_path, text='甲方张三，乙方李四。'):
    d = Document(); d.add_paragraph(text); p = tmp_path / 'source.docx'; d.save(p); return p


def change(data, name, fn):
    items = unpack(data); root = etree.fromstring(items[name]); fn(root); items[name] = etree.tostring(root); return pack(items)


def test_multiple_disjoint_comments_in_same_paragraph(tmp_path):
    p = source(tmp_path)
    result = write_verified_comments(p, [CommentSpec('张三', '核对甲方', 'VL', 'VL'), CommentSpec('李四', '核对乙方', 'VL', 'VL')])
    assert len(result.comment_ids) == 2


def test_cross_adjacent_paragraphs_and_cell_boundaries(tmp_path):
    d = Document(); d.add_paragraph('付款期限'); d.add_paragraph('三十日。')
    t = d.add_table(rows=1, cols=2); t.cell(0,0).text='左'; t.cell(0,1).text='右'
    p=tmp_path/'cross.docx'; d.save(p)
    result=write_verified_comments(p,[CommentSpec('期限\n三十日','核对','VL','VL')])
    assert verify_comment_copy(p,result.copy_bytes).passed
    with pytest.raises(DocumentDomainError,match='docx_comment_anchor_invalid'):
        write_comments(p,[CommentSpec('左\n右','意见','VL','VL')])


def test_duplicate_anchor_needs_occurrence(tmp_path):
    p=source(tmp_path,'张三和张三')
    with pytest.raises(DocumentDomainError): write_comments(p,[CommentSpec('张三','核对','VL','VL')])
    result=write_verified_comments(p,[CommentSpec('张三','核对','VL','VL',occurrence=1)])
    assert result.comment_ids == (0,)


def test_existing_comment_preserved_and_id_allocated(tmp_path):
    p=source(tmp_path); d=Document(p); d.add_paragraph('第二段说明。'); d.save(p)
    existing=write_comments(p,[CommentSpec('张三','原意见','原作者','YZ')])
    p.write_bytes(existing.copy_bytes)
    result=write_verified_comments(p,[CommentSpec('第二段','新意见','VL','VL')])
    assert result.comment_ids == (1,)
    assert '原意见' in Document(io.BytesIO(result.copy_bytes)).comments.get(0).text


@pytest.mark.parametrize('kind',['ref_text','ref_style','header','existing_comment','body_bold','missing_type'])
def test_verifier_rejects_real_tampering(tmp_path,kind):
    p=source(tmp_path); d=Document(p); d.sections[0].header.paragraphs[0].text='页眉'; d.add_paragraph('第二条。'); d.save(p)
    if kind=='existing_comment':
        p.write_bytes(write_comments(p,[CommentSpec('第二条','原意见','VL','VL')]).copy_bytes)
    data=write_comments(p,[CommentSpec('张三','意见','VL','VL')]).copy_bytes
    assert verify_comment_copy(p,data).passed
    def body(root):
        if kind=='ref_text':
            ref=root.find('.//w:commentReference',NS); etree.SubElement(ref.getparent(),f'{{{W}}}t').text='偷偷增加'
        elif kind=='ref_style':
            ref=root.find('.//w:commentReference',NS); etree.SubElement(ref.getparent().find('w:rPr',NS),f'{{{W}}}b')
        else:
            r=root.find('.//w:r',NS); prop=etree.Element(f'{{{W}}}rPr'); etree.SubElement(prop,f'{{{W}}}b'); r.insert(0,prop)
    if kind in ('ref_text','ref_style','body_bold'): changed=change(data,'word/document.xml',body)
    elif kind=='header': changed=change(data,'word/header1.xml',lambda root: setattr(root.find('.//w:t',NS),'text','改动'))
    elif kind=='existing_comment': changed=change(data,'word/comments.xml',lambda root: setattr(root.find('.//w:t',NS),'text','抹除原意见'))
    else:
        def remove_type(root):
            for el in list(root):
                if el.get('PartName')=='/word/comments.xml':root.remove(el)
        changed=change(data,'[Content_Types].xml',remove_type)
    assert changed != data
    assert not verify_comment_copy(p,changed).passed


def test_overlap_and_cascade_do_not_deliver(tmp_path):
    p=source(tmp_path,'张三甲方')
    with pytest.raises(DocumentDomainError,match='docx_replace_overlap'):
        replace_texts(p,[TextReplacement('张三','A'),TextReplacement('三甲','B')])
    with pytest.raises(DocumentDomainError,match='docx_replace_not_found'):
        replace_texts(p,[TextReplacement('张三','新名'),TextReplacement('新名','最终')])


def test_first_character_format_and_tabs_preserved(tmp_path):
    d=Document(); p=d.add_paragraph(); r=p.add_run('电话\t张'); r.bold=True; p.add_run('三尾部')
    f=tmp_path/'tabs.docx'; d.save(f)
    out=replace_texts(f,[TextReplacement('张三','【姓名】')]); actual=Document(io.BytesIO(out.copy_bytes)).paragraphs[0]
    assert actual.text=='电话\t【姓名】尾部'
    assert any('【姓名】' in r.text and r.bold for r in actual.runs)


def test_repeated_matches_within_same_text_node(tmp_path):
    p=source(tmp_path,'张三/张三/李四/张三')
    out=replace_texts(p,[TextReplacement('张三','甲'),TextReplacement('李四','长一些的乙')])
    assert out.counts==(3,1)
    assert Document(io.BytesIO(out.copy_bytes)).paragraphs[0].text=='甲/甲/长一些的乙/甲'


def test_hyperlink_cannot_be_skipped_when_matching(tmp_path):
    d=Document(); p=d.add_paragraph(); p.add_run('张'); h=OxmlElement('w:hyperlink');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='中间';r.append(t);h.append(r);p._p.append(h);p.add_run('三')
    f=tmp_path/'hyperlink.docx';d.save(f)
    with pytest.raises(DocumentDomainError,match='docx_replace_not_found'):replace_texts(f,[TextReplacement('张三','姓名')])
    result=replace_texts(f,[TextReplacement('中间','链接文字')]); assert '链接文字' in ''.join(Document(io.BytesIO(result.copy_bytes)).element.itertext())


def test_headers_footers_notes_and_empty_body(tmp_path):
    d=Document(); d.sections[0].header.paragraphs[0].text='张三'; d.sections[0].footer.paragraphs[0].text='张三'
    f=tmp_path/'stories.docx'; d.save(f);items=unpack(f.read_bytes())
    rels=etree.fromstring(items['word/_rels/document.xml.rels']);cts=etree.fromstring(items['[Content_Types].xml'])
    for kind in ('footnote','endnote'):
        items[f'word/{kind}s.xml']=(f'<w:{kind}s xmlns:w="{W}"><w:{kind} w:id="1"><w:p><w:r><w:t>张三</w:t></w:r></w:p></w:{kind}></w:{kind}s>').encode()
        etree.SubElement(rels,f'{{{REL}}}Relationship',Id=kind,Type=f'http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s',Target=f'{kind}s.xml')
        etree.SubElement(cts,f'{{{CT}}}Override',PartName=f'/word/{kind}s.xml',ContentType=f'application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml')
    items['word/_rels/document.xml.rels']=etree.tostring(rels);items['[Content_Types].xml']=etree.tostring(cts);f.write_bytes(pack(items))
    out=process_copy(f,[TextReplacement('张三','【姓名】')]);assert out.counts==(4,);assert out.audit.passed


@pytest.mark.parametrize('place',['chart','smartart','attribute','field','split_runs'])
def test_residue_audit_covers_all_xml_text(tmp_path,place):
    p=source(tmp_path); out=process_copy(p,[TextReplacement('张三','【姓名】')]); items=unpack(out.copy_bytes)
    if place in ('chart','smartart'):
        path='word/charts/chart1.xml' if place=='chart' else 'word/diagrams/data1.xml';items[path]=b'<root><t>'+ '张三'.encode()+b'</t></root>'
    else:
        root=etree.fromstring(items['word/document.xml']);para=root.find('.//w:p',NS)
        if place=='attribute':para.set('test','张三')
        elif place=='field':etree.SubElement(etree.SubElement(para,f'{{{W}}}r'),f'{{{W}}}instrText').text='张三'
        else:
            for value in ('张','三'):etree.SubElement(etree.SubElement(para,f'{{{W}}}r'),f'{{{W}}}t').text=value
        items['word/document.xml']=etree.tostring(root)
    assert not audit_processed_copy(pack(items),[TextReplacement('张三','【姓名】')],scrubbed=True).passed


def test_audit_detects_missing_replacement_and_empty_content(tmp_path):
    out=process_copy(source(tmp_path),[TextReplacement('张三','【姓名】')])
    bad=change(out.copy_bytes,'word/document.xml',lambda root: root.find('w:body',NS).clear())
    assert not audit_processed_copy(bad,[TextReplacement('张三','【姓名】')],scrubbed=True).passed


def test_complete_metadata_matrix(tmp_path):
    p=source(tmp_path);items=unpack(p.read_bytes());rels=etree.fromstring(items['word/_rels/document.xml.rels']);cts=etree.fromstring(items['[Content_Types].xml'])
    for n in ('word/commentsExtended.xml','word/commentsIds.xml','word/people.xml','docProps/custom.xml','customXml/item1.xml'):
        items[n]=b'<root>personal-data</root>'
        etree.SubElement(rels,f'{{{REL}}}Relationship',Id=n,Type='test',Target='../'+n)
        etree.SubElement(cts,f'{{{CT}}}Override',PartName='/'+n,ContentType='application/xml')
    items['word/_rels/document.xml.rels']=etree.tostring(rels);items['[Content_Types].xml']=etree.tostring(cts)
    settings=etree.fromstring(items['word/settings.xml']);etree.SubElement(settings,f'{{{W}}}trackRevisions');items['word/settings.xml']=etree.tostring(settings)
    p.write_bytes(pack(items));out=process_copy(p,[TextReplacement('张三','【姓名】')]);assert out.audit.passed
    assert not any(n.startswith('customXml/') or n.endswith(('commentsExtended.xml','commentsIds.xml','people.xml','custom.xml')) for n in unpack(out.copy_bytes))


def test_move_and_property_only_revisions_are_accepted(tmp_path):
    p=source(tmp_path,'现有文字');items=unpack(p.read_bytes());root=etree.fromstring(items['word/document.xml']);para=root.find('.//w:p',NS)
    for tag,txt in (('moveFrom','旧名'),('moveTo','张三')):
        box=etree.SubElement(para,f'{{{W}}}{tag}');etree.SubElement(etree.SubElement(box,f'{{{W}}}r'),f'{{{W}}}t').text=txt
    for tag in ('moveFromRangeStart','moveFromRangeEnd','moveToRangeStart','moveToRangeEnd'):etree.SubElement(para,f'{{{W}}}{tag}')
    items['word/document.xml']=etree.tostring(root);p.write_bytes(pack(items));out=process_copy(p,[TextReplacement('张三','【姓名】')]);assert out.audit.passed
    assert '旧名' not in ''.join(Document(io.BytesIO(out.copy_bytes)).element.itertext())


def test_unsupported_cell_revision_rejects_without_mutating(tmp_path):
    p=source(tmp_path);data=change(p.read_bytes(),'word/document.xml',lambda root: etree.SubElement(root.find('.//w:p',NS),f'{{{W}}}cellMerge'))
    p.write_bytes(data)
    with pytest.raises(DocumentDomainError,match='docx_replace_revisions_present'):process_copy(p,[TextReplacement('张三','【姓名】')])
    assert p.read_bytes()==data


def test_malformed_outputs_fail_closed(tmp_path):
    p=source(tmp_path)
    assert not verify_comment_copy(p,b'not zip').passed
    assert not audit_processed_copy(b'not zip',[],scrubbed=True).passed


def test_empty_existing_comments_part(tmp_path):
    p=source(tmp_path);d=Document(p);assert len(d.comments)==0;d.save(p)
    result=write_verified_comments(p,[CommentSpec('张三','意见','VL','VL')]);assert result.comment_ids==(0,)


@pytest.mark.parametrize('fault',['duplicate_id','wrong_rel_tag','wrong_type_tag'])
def test_invalid_comment_relationship_shapes(tmp_path,fault):
    p=source(tmp_path);data=write_comments(p,[CommentSpec('张三','意见','VL','VL')]).copy_bytes
    def mutate(root):
        if fault=='wrong_type_tag':
            next(e for e in root if e.get('PartName')=='/word/comments.xml').tag='{urn:bad}Override'
        else:
            e=next(e for e in root if e.get('Type','').endswith('/comments'))
            if fault=='duplicate_id':e.set('Id',root[0].get('Id'))
            else:e.tag='{urn:bad}Relationship'
    bad=change(data,'[Content_Types].xml' if fault=='wrong_type_tag' else 'word/_rels/document.xml.rels',mutate)
    assert not verify_comment_copy(p,bad).passed


def test_existing_placeholder_cannot_mask_missing_replacement(tmp_path):
    p=source(tmp_path,'张三');d=Document(p);d.add_paragraph('【姓名】');d.save(p)
    replaced=replace_texts(p,[TextReplacement('张三','【姓名】')]);clean=scrub_package(replaced.copy_bytes)
    bad=change(clean.copy_bytes,'word/document.xml',lambda root: setattr(root.find('.//w:t',NS),'text','无关文字'))
    audit=audit_processed_copy(bad,[TextReplacement('张三','【姓名】')],scrubbed=True,expected_counts=replaced.counts,expected_paragraphs=replaced.expected_paragraphs)
    assert not audit.passed and any('replacement-position' in f for f in audit.failures)


def test_nonstandard_header_part_name(tmp_path):
    p=source(tmp_path,'正文');d=Document(p);d.sections[0].header.paragraphs[0].text='张三';d.save(p);items=unpack(p.read_bytes())
    items['word/h1.xml']=items.pop('word/header1.xml')
    for n in ('word/_rels/document.xml.rels','[Content_Types].xml'):items[n]=items[n].replace(b'header1.xml',b'h1.xml')
    p.write_bytes(pack(items));out=process_copy(p,[TextReplacement('张三','【姓名】')]);assert out.counts==(1,)
    assert Document(io.BytesIO(out.copy_bytes)).sections[0].header.paragraphs[0].text=='【姓名】'


def test_deleted_row_does_not_process_detached_descendants():
    root=etree.fromstring(f'<w:body xmlns:w="{W}"><w:tbl><w:tr><w:trPr><w:del/></w:trPr><w:tc><w:p><w:pPr><w:rPr><w:del/></w:rPr></w:pPr><w:r><w:t>删除的行</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body>'.encode())
    accept_revisions(root);assert not root.findall('.//w:tr',NS)


def test_unrelated_revision_does_not_block_comment(tmp_path):
    d=Document();p=d.add_paragraph();box=OxmlElement('w:ins');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='修订文字';r.append(t);box.append(r);p._p.append(box);p.add_run('独立条款')
    f=tmp_path/'unrelated-revision.docx';d.save(f)
    result=write_verified_comments(f,[CommentSpec('独立条款','意见','VL','VL')]);assert result.comment_ids==(0,)


def test_cross_paragraph_anchor_with_empty_middle_paragraph(tmp_path):
    d=Document();d.add_paragraph('第一段');d.add_paragraph();d.add_paragraph('第三段');f=tmp_path/'blank.docx';d.save(f)
    assert write_verified_comments(f,[CommentSpec('第一段\n\n第三段','意见','VL','VL')]).comment_ids==(0,)
