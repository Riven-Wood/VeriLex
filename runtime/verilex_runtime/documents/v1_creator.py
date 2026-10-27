from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
from typing import BinaryIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from verilex_runtime.documents.errors import DocumentDomainError
from verilex_runtime.documents.v1_markdown import MarkdownBlock


@dataclass(frozen=True, slots=True)
class V1ValidationReport:
    block_count: int
    text_char_count: int

    def as_json(self) -> str:
        return json.dumps(
            {
                "schema_version": 1,
                "structure": {"status": "passed", "issue_counts": {}},
                "render": {"status": "not_required"},
                "vision": {"status": "not_run"},
            },
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        )


def _apply_run_typography(run, size: Pt, *, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = size
    if bold is not None:
        run.bold = bold
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def _configure_style(style, size: Pt) -> None:
    style.font.name = "Times New Roman"
    style.font.size = size
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def _configure_document(document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(25.4)
    section.left_margin = Mm(25.4)
    section.right_margin = Mm(25.4)

    normal = document.styles["Normal"]
    _configure_style(normal, Pt(12))
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    for level, size in ((1, Pt(14)), (2, Pt(13)), (3, Pt(12))):
        _configure_style(document.styles[f"Heading {level}"], size)


def _reopened_structure_matches(document, blocks: tuple[MarkdownBlock, ...]) -> bool:
    paragraphs = document.paragraphs
    if len(paragraphs) != len(blocks) + 1:
        return False
    title = paragraphs[0]
    if (
        title.alignment != WD_ALIGN_PARAGRAPH.CENTER
        or len(title.runs) != 1
        or title.runs[0].bold is not True
        or title.runs[0].font.size != Pt(18)
    ):
        return False
    expected_styles = {
        "paragraph": "Normal",
        "bullet": "List Bullet",
        "numbered": "List Number",
        "pagebreak": "Normal",
    }
    for block, paragraph in zip(blocks, paragraphs[1:], strict=True):
        expected_style = (
            f"Heading {block.level}"
            if block.kind == "heading"
            else expected_styles.get(block.kind)
        )
        if expected_style is None or paragraph.style.name != expected_style:
            return False
        if block.kind == "pagebreak" and not paragraph._element.xpath(
            './/w:br[@w:type="page"]'
        ):
            return False
    return True


def create_v1_docx(
    path: Path | BinaryIO,
    title: str,
    blocks: tuple[MarkdownBlock, ...],
) -> V1ValidationReport:
    # document_create already supplies a standalone title. Some models repeat
    # exactly that title as the first Markdown heading; remove only that exact
    # redundant heading, never a different heading or a body paragraph.
    if blocks and blocks[0].kind == "heading" and blocks[0].text.strip() == title.strip():
        blocks = blocks[1:]
    document = Document()
    _configure_document(document)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.first_line_indent = Pt(0)
    title_run = title_paragraph.add_run(title)
    _apply_run_typography(title_run, Pt(18), bold=True)

    heading_sizes = {1: Pt(14), 2: Pt(13), 3: Pt(12)}
    for block in blocks:
        if block.kind == "heading":
            if block.level not in heading_sizes:
                raise DocumentDomainError("document_validation_failed")
            paragraph = document.add_heading(level=block.level)
            run = paragraph.add_run(block.text)
            _apply_run_typography(run, heading_sizes[block.level])
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.first_line_indent = Pt(24)
            run = paragraph.add_run(block.text)
            _apply_run_typography(run, Pt(12))
        elif block.kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(block.text)
            _apply_run_typography(run, Pt(12))
        elif block.kind == "numbered":
            paragraph = document.add_paragraph(style="List Number")
            run = paragraph.add_run(block.text)
            _apply_run_typography(run, Pt(12))
        elif block.kind == "pagebreak":
            paragraph = document.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        else:
            raise DocumentDomainError("document_validation_failed")

    document.save(path)
    if hasattr(path, "seek"):
        path.seek(0)
    reopened = Document(path)
    expected_text = [title, *(block.text for block in blocks)]
    actual_text = [paragraph.text for paragraph in reopened.paragraphs]
    if actual_text != expected_text or not _reopened_structure_matches(reopened, blocks):
        raise DocumentDomainError("document_validation_failed")
    if hasattr(path, "seek"):
        path.seek(0)

    return V1ValidationReport(
        block_count=len(blocks),
        text_char_count=len(title) + sum(len(block.text) for block in blocks),
    )
