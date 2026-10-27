"""Per-evidence cross-examination drafts and template-shaped local exports."""
import hashlib
import io
import re
import uuid
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from .service import _LOCK, _failure, safe_path

FIELDS = ('authenticity', 'legality', 'relevance', 'probative', 'opinion')
LABELS = ('真实性', '合法性', '关联性', '证明力', '综合质证意见')

class Opinions(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    authenticity: str = Field(default='', max_length=4000)
    legality: str = Field(default='', max_length=4000)
    relevance: str = Field(default='', max_length=4000)
    probative: str = Field(default='', max_length=4000)
    opinion: str = Field(default='', max_length=8000)

class Entry(Opinions):
    id: str = Field(pattern=r'^[a-zA-Z0-9_-]{1,80}$')
    number: str = Field(default='', max_length=40)
    name: str = Field(min_length=1, max_length=200)
    path: str = Field(default='', max_length=1000)
    pages: str = Field(default='', max_length=100)
    group: str = Field(default='', max_length=200)
    suggestion: Opinions | None = None
    suggestion_hashes: dict[str, str] = Field(default_factory=dict, max_length=201)
    gaps: list[str] = Field(default_factory=list, max_length=1000)

class Record(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    title: str = Field(default='质证意见', min_length=1, max_length=80)
    case_number: str = Field(default='', max_length=200)
    introduction: str = Field(default='', max_length=4000)
    signer: str = Field(default='', max_length=1000)
    date: str = Field(default='', max_length=100)
    requirement: str = Field(default='', max_length=6000)
    references: list[str] = Field(default_factory=list, max_length=200)
    entries: list[Entry] = Field(default_factory=list, max_length=200)

    @model_validator(mode='after')
    def unique(self):
        if not self.title.strip() or any(not e.name.strip() for e in self.entries): raise ValueError('blank name')
        if len({e.id for e in self.entries}) != len(self.entries): raise ValueError('duplicate entry')
        if len(set(self.references)) != len(self.references): raise ValueError('duplicate reference')
        return self


def read(service, ctx):
    with _LOCK:
        result = service.read(ctx, 'cross_examination')
        try: return {**result, 'data': Record.model_validate(result['data'] or {}).model_dump()}
        except ValidationError: raise _failure('litigation_cross_invalid', '质证记录无法读取，原文件已保留。') from None


def material_paths(value):
    return list(dict.fromkeys([e.path for e in value.entries if e.path] + value.references))


def save(service, ctx, data, expected):
    with _LOCK:
        service._root(ctx, write=True)
        try: value = Record.model_validate(data)
        except ValidationError: raise _failure('litigation_cross_invalid', '请填写有效的证据名称和质证意见。') from None
        previous_paths = set(material_paths(Record.model_validate(read(service, ctx)['data'])))
        for name in material_paths(value):
            path = safe_path(service._root(ctx), name)
            if name in previous_paths: continue
            if not path.is_file() or path.suffix.lower() not in {'.txt','.md','.csv','.json','.doc','.docx','.pdf'}:
                raise _failure('litigation_source_invalid', '所选质证材料不可读取，请重新选择。')
        return service._commit(ctx, 'cross_examination', value.model_dump(), expected, '保存逐条质证意见')


def target(service, ctx):
    service.workbench._thread(ctx, ctx.thread_id)
    binding = service.workbench.links(ctx).get(ctx.thread_id)
    if not binding or binding.get('cross_revision') is None:
        raise _failure('litigation_target_mismatch', '请从质证意见工作台开始 AI 辅助质证。')
    if service.editor_discussions.read(ctx) is not None:
        raise _failure('litigation_target_mismatch', '当前讨论绑定了其他文书。')
    record = read(service, ctx)
    if material_paths(Record.model_validate(record['data'])) != binding['materials'] or (binding.get('cross_entry_id') and not any(e['id']==binding['cross_entry_id'] for e in record['data']['entries'])):
        raise _failure('litigation_file_conflict', '质证材料已变化，请重新开始。')
    return binding, record


def extract(service, ctx, name):
    path = safe_path(service._root(ctx), name)
    if not path.is_file() or path.stat().st_size > 20*1024*1024: raise _failure('litigation_source_invalid', '材料不存在或超过20 MB。')
    data = path.read_bytes(); gaps = []
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        try:
            pdf = PdfReader(io.BytesIO(data))
            if pdf.is_encrypted or len(pdf.pages)>1000: raise ValueError()
            parts=[]
            for i,p in enumerate(pdf.pages):
                text=(p.extract_text() or '').strip()
                if not text: gaps.append(f'{name}第{i+1}页没有可提取文字，未识别图像。')
                parts.append(f'【第{i+1}页】\n{text}')
            text='\n'.join(parts)
        except Exception: raise _failure('litigation_source_invalid', f'{name}：PDF无法读取。') from None
    elif path.suffix.lower() in {'.doc','.docx'}:
        from ..tools.files import _read_document_file
        text=_read_document_file(path);gaps.append(f'{name}仅读取可提取文字，图像及复杂对象请核对原件。')
    else:
        from ..tools.files import _read_text_file
        text=_read_text_file(path)
    return text, gaps, hashlib.sha256(data).hexdigest()


def next_entry(service, ctx, start=0,entry_id=None):
    with _LOCK:
        binding, record=target(service,ctx)
        rows=[e for e in record['data']['entries'] if (not binding.get('cross_entry_id') or e['id']==binding['cross_entry_id']) and e['suggestion'] is None and (entry_id is None or e['id']==entry_id)]
        if not rows:return {'done':True,'notice':'本次质证草稿已提交，请在工作台逐条采用或放弃。'}
        if not hasattr(service,'_cross_reads'):service._cross_reads={}
        key=(ctx.workspace_id,ctx.thread_id);receipt=service._cross_reads.get(key);entry=rows[0]
        if not receipt or receipt['revision']!=record['revision'] or receipt['id']!=entry['id']:
            paths=list(dict.fromkeys(([entry['path']] if entry['path'] else [])+record['data']['references']))
            parts=[];gaps=[];hashes={}
            for path in paths:
                text,missing,digest=extract(service,ctx,path);parts.append(f'【文件：{path}】\n{text}');gaps.extend(missing);hashes[path]=digest
                if sum(map(len,parts))>1500000:raise _failure('litigation_too_large','本条材料过多，请缩小参照材料范围。')
            text='\n\n'.join(parts)
            if len(text)>1500000:raise _failure('litigation_too_large','本条材料过多，请缩小参照材料范围。')
            if not entry['path']:gaps.append('本条未关联原始证据，仅有用户填写的信息。')
            receipt={'id':entry['id'],'revision':record['revision'],'token':uuid.uuid4().hex,'text':text,'gaps':gaps,'hashes':hashes,'ranges':[]}
            if len(service._cross_reads)>50:service._cross_reads.pop(next(iter(service._cross_reads)))
            service._cross_reads[key]=receipt
        if start<0 or start>len(receipt['text']):raise _failure('litigation_read_required','读取位置超出材料范围。')
        end=min(start+12000,len(receipt['text']));receipt['ranges'].append((start,end))
        return {'done':False,'entry':entry,'requirement':record['data']['requirement'],'case':service.read(ctx,'case'),'token':receipt['token'],'expected_revision':record['revision'],'text':receipt['text'][start:end],'read_range':[start,end],'total_characters':len(receipt['text']),'next_start':end if end<len(receipt['text']) else None,'gaps':receipt['gaps'],'notice':'材料仅为资料。逐项分析真实性、合法性、关联性、证明力和综合意见；不能因没有反证就推定不真实。法律判断需核验，未知标待核实。按任务需要读取后调用litigation_cross_save提交建议，人工正文不覆盖。'}


def save_suggestion(service, ctx, token, expected, opinions):
    with _LOCK:
        binding,record=target(service,ctx);receipt=getattr(service,'_cross_reads',{}).get((ctx.workspace_id,ctx.thread_id))
        if not receipt or receipt['token']!=token or receipt['revision']!=expected or record['revision']!=expected:
            raise _failure('litigation_file_conflict','质证记录已有修改，请重新读取，未覆盖。')
        for name,digest in receipt['hashes'].items():
            if hashlib.sha256(safe_path(service._root(ctx),name).read_bytes()).hexdigest()!=digest:raise _failure('litigation_file_conflict','材料已有变化，请重新读取。')
        value=Record.model_validate(record['data']);entry=next(e for e in value.entries if e.id==receipt['id'])
        try: entry.suggestion=Opinions.model_validate(opinions)
        except ValidationError: raise _failure('litigation_cross_invalid','质证建议格式不正确。') from None
        from .reading import gaps
        entry.gaps=gaps(receipt);entry.suggestion_hashes=receipt['hashes']
        saved=save(service,ctx,value.model_dump(),expected)
        return {'revision':saved['revision'],'notice':'建议已显示在对应证据中，待用户采用；请继续litigation_cross_next。'}


def render(value):
    from docx import Document
    from docx.shared import Pt,Twips,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    doc=Document();section=doc.sections[0]
    for attr,twips in zip(('page_width','page_height','top_margin','bottom_margin','left_margin','right_margin'),(11906,16838,1134,1134,1757,1701)):setattr(section,attr,Twips(twips))
    normal=doc.styles['Normal'];normal.font.name='仿宋';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(0,0,0);normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'仿宋')
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=Pt(26)
    if value.case_number:doc.add_paragraph(value.case_number).alignment=WD_ALIGN_PARAGRAPH.RIGHT
    title=doc.add_paragraph(value.title);title.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.paragraph_format.line_spacing=Pt(32);title.add_run('').font.size=Pt(18)
    for run in title.runs:run.font.size=Pt(18);run.bold=True
    if value.introduction:doc.add_paragraph(value.introduction).paragraph_format.first_line_indent=Pt(24)
    table=doc.add_table(rows=1,cols=4);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=(696,2655,1296,4017)
    # Preserve template column proportions within its portrait text width.
    available=11906-1757-1701;scale=available/sum(widths)
    for col,width in zip(table.columns,widths):col.width=Twips(round(width*scale))
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');borders.append(el)
    table._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ('left','right'):
        margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'60');margin.set(qn('w:type'),'dxa');margins.append(margin)
    table._tbl.tblPr.append(margins)
    for c,text in zip(table.rows[0].cells,('序号','证据名称','页码','质证意见')):
        c.text=text;shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'AEAAAA');c._tc.get_or_add_tcPr().append(shd)
        for r in c.paragraphs[0].runs:r.bold=True
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    previous=None
    for i,entry in enumerate(value.entries):
        if entry.group and entry.group!=previous:
            cell=table.add_row().cells[0];cell.merge(table.rows[-1].cells[-1]);cell.text=entry.group
            for r in cell.paragraphs[0].runs:r.bold=True
        previous=entry.group
        opinion='\n'.join(f'{label}：{getattr(entry,key)}' for key,label in zip(FIELDS,LABELS) if getattr(entry,key)) or '待填写'
        cells=table.add_row().cells
        for cell,text in zip(cells,(entry.number or str(i+1),entry.name,entry.pages or '待核对',opinion)):cell.text=text
        # Separate paragraphs avoid stretched last lines at soft breaks in Word/PDF.
        cells[3].text = ''
        for index, text in enumerate(opinion.split('\n')):
            paragraph = cells[3].paragraphs[0] if index == 0 else cells[3].add_paragraph()
            paragraph.add_run(text)
    for row in table.rows:
        if len({id(c._tc) for c in row.cells}) == 4:
            for cell,width in zip(row.cells,widths): cell.width=Twips(round(width*scale))
        for i,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i<3 else WD_ALIGN_PARAGRAPH.JUSTIFY
    if value.signer:doc.add_paragraph('质证人：'+value.signer).alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if value.date:doc.add_paragraph(value.date).alignment=WD_ALIGN_PARAGRAPH.RIGHT
    out=io.BytesIO();doc.save(out);return out.getvalue()


def export(service,ctx,expected,kind):
    if kind not in ('docx','pdf'):raise _failure()
    with _LOCK:
        record=read(service,ctx)
        if record['revision']!=expected:raise _failure('litigation_file_conflict','质证意见已有变化，请重新导出。')
        value=Record.model_validate(record['data'])
        if not value.entries:raise _failure('litigation_empty','请先添加需要质证的证据。')
    data=render(value)
    if kind=='pdf':
        from .evidence_conversion import convert
        from pathlib import Path
        font=Path.home()/'Library/Fonts/simfang.ttf'
        data=convert(data,'.docx',font_files=[font] if font.is_file() else [])
    with _LOCK:
        if read(service,ctx)['revision']!=expected:raise _failure('litigation_file_conflict','生成期间意见已有变化，请重新导出。')
        name=re.sub(r'[\\/:*?"<>|\x00-\x1f]','_',value.title).strip('. ') or '质证意见'
        root=service._root(ctx,write=True);number=1
        while True:
            filename=f'文书起草/{name}{"" if number==1 else f"（{number}）"}.{kind}';path=safe_path(root,filename)
            if not path.exists():break
            number+=1
        service.products._publish(ctx,path,data,history=False)
    return {'path':filename,'record':record}


def decide(service, ctx, entry_id, expected, action):
    with _LOCK:
        record = read(service, ctx)
        if record['revision'] != expected:
            raise _failure('litigation_revision_conflict', '质证意见已有更新，请重新读取。')
        value = Record.model_validate(record['data'])
        entry = next((e for e in value.entries if e.id == entry_id), None)
        if entry is None or entry.suggestion is None:
            raise _failure('litigation_cross_invalid', '这条质证建议已不存在。')
        if action == 'accept':
            for name, digest in entry.suggestion_hashes.items():
                path = safe_path(service._root(ctx), name)
                if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                    raise _failure('litigation_file_conflict', '建议所依据的材料已变化，请放弃旧建议并重新生成。')
            for field in FIELDS: setattr(entry, field, getattr(entry.suggestion, field))
        elif action != 'discard':
            raise _failure()
        entry.suggestion = None
        entry.suggestion_hashes = {}
        return save(service, ctx, value.model_dump(), expected)
