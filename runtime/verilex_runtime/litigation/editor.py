"""Bounded editor/DOCX adapter. Original packages are never flattened to text.

The text-only patch path retains every untouched OPC part and paragraph. Rich
editing uses a new document; unsupported imported objects stay in the source.
No path I/O or workspace authorization lives here: callers own publication.
"""
from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from verilex_runtime.documents.docx_package import W, SPACE, checked_bytes, parts, pack, xml, serialize
from .service import _failure

MAIN = 'word/document.xml'
MAX_TEXT = 500_000


def fingerprint(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _text(p):
    return ''.join(n.text or '' for n in p.iter(f'{{{W}}}t')
                   if next(n.iterancestors(f'{{{W}}}p'), None) is p)


def _editable(p):
    # Only ordinary runs. A paragraph touching a field, picture, tracked change,
    # bookmark, comment or footnote stays read-only, even when it also has text.
    if any(a.tag in {f'{{{W}}}{name}' for name in ('ins','del','moveFrom','moveTo','txbxContent','sdt')} for a in p.iterancestors()):return False
    return all(c.tag in {f'{{{W}}}pPr', f'{{{W}}}r'} for c in p) and all(
        c.tag in {f'{{{W}}}rPr', f'{{{W}}}t'} for r in p.findall(f'{{{W}}}r') for c in r)


def inspect_docx(source: bytes) -> dict:
    data = checked_bytes(source)
    items = parts(data)
    root = xml(items[MAIN])
    tree = root.getroottree()
    paragraphs = [{'id': tree.getpath(p), 'text': _text(p), 'editable': _editable(p)}
                  for p in root.iter(f'{{{W}}}p')]
    if sum(len(p['text']) for p in paragraphs) > MAX_TEXT:
        raise _failure('litigation_too_large', '文档超过编辑上限，请拆分后再编辑。')
    complex_names = {'drawing': '图片', 'pict': '图形', 'footnoteReference': '脚注',
                     'endnoteReference': '尾注', 'commentRangeStart': '批注',
                     'ins': '修订', 'del': '修订', 'fldChar': '域', 'sdt': '内容控件',
                     'altChunk': '嵌入内容', 'object': '嵌入对象'}
    retained = sorted({label for node, label in complex_names.items() if root.find('.//' + f'{{{W}}}{node}') is not None})
    return {'file_hash': fingerprint(data), 'part': MAIN, 'paragraphs': paragraphs,
            'retained_objects': retained, 'mode': 'paragraph_text',
            'notice': '仅修改选中段落文字；其他内容与原包保留。'}


@dataclass(frozen=True)
class ParagraphEdit:
    paragraph_id: str
    original: str
    text: str


def patch_docx(source: bytes, expected_hash: str, edits: list[ParagraphEdit]) -> bytes:
    """Change exact structural paragraphs in this exact snapshot, never global replace."""
    data = checked_bytes(source)
    if fingerprint(data) != expected_hash:
        raise _failure('litigation_file_conflict', 'Word文件已有更新，请重新读取。')
    if not edits or len({e.paragraph_id for e in edits}) != len(edits):
        raise _failure()
    items = parts(data)
    root = xml(items[MAIN]); tree = root.getroottree()
    # Never evaluate a caller-supplied XPath.
    paragraphs = {tree.getpath(p): p for p in root.iter(f'{{{W}}}p')}
    if sum(len(e.text) for e in edits) > MAX_TEXT:
        raise _failure('litigation_too_large', '本次修改超过编辑上限。')
    for edit in edits:
        p = paragraphs.get(edit.paragraph_id)
        if p is None or not _editable(p):
            raise _failure('litigation_editor_readonly', '此段含有暂不支持编辑的对象，原内容已保留。')
        if _text(p) != edit.original:
            raise _failure('litigation_file_conflict', '选中段落已有更新，未应用修改。')
        if any(c in edit.text for c in '\n\r\t'):
            raise _failure('litigation_editor_readonly', '保格式编辑暂不支持拆分段落或插入制表符。')
        try:
            etree.Element('text').text = edit.text
        except (ValueError, TypeError):
            raise _failure() from None
        # Preserve unaffected run formatting and common prefix/suffix; inserted
        # characters inherit the first affected run, including cross-run Chinese.
        old = edit.original; new = edit.text; lo = 0
        while lo < min(len(old), len(new)) and old[lo] == new[lo]: lo += 1
        tail = 0
        while tail < min(len(old)-lo, len(new)-lo) and old[len(old)-1-tail] == new[len(new)-1-tail]: tail += 1
        hi = len(old)-tail
        inserted = new[lo:len(new)-tail if tail else len(new)]
        nodes = list(p.iter(f'{{{W}}}t')); offset = 0; injected = False
        for node in nodes:
            value = node.text or ''; end = offset + len(value)
            if (offset < hi and end > lo) or (lo == hi and offset <= lo <= end and not injected):
                prefix = value[:max(0, lo-offset)]
                suffix = value[max(0, hi-offset):] if hi < end else ''
                node.text = prefix + (inserted if not injected else '') + suffix
                node.set(SPACE, 'preserve'); injected = True
            offset = end
        if not injected and inserted:
            run = etree.SubElement(p, f'{{{W}}}r')
            t = etree.SubElement(run, f'{{{W}}}t'); t.text = inserted; t.set(SPACE, 'preserve')
    items[MAIN] = serialize(root)
    output = pack(items)
    checked_bytes(output)
    return output


def validate_document(value: dict) -> dict:
    """Strict shared subset of the editor schema; unknown objects cannot disappear."""
    count = 0
    allowed = {'doc', 'paragraph', 'heading', 'text', 'hardBreak', 'bulletList', 'orderedList',
               'listItem', 'table', 'tableRow', 'tableCell', 'tableHeader'}
    containers = {'doc': {'paragraph','heading','bulletList','orderedList','table'},
                  'paragraph': {'text','hardBreak'}, 'heading': {'text','hardBreak'},
                  'bulletList': {'listItem'}, 'orderedList': {'listItem'},
                  'listItem': {'paragraph','bulletList','orderedList'},
                  'table': {'tableRow'}, 'tableRow': {'tableCell','tableHeader'},
                  'tableCell': {'paragraph','heading','bulletList','orderedList'},
                  'tableHeader': {'paragraph','heading','bulletList','orderedList'}}
    def visit(node, depth=0):
        nonlocal count
        if depth > 20 or not isinstance(node, dict) or node.get('type') not in allowed:
            raise _failure('litigation_editor_unsupported', '文档包含暂不支持保存的内容。')
        kind = node['type']; count += 1
        if count > 50000: raise _failure('litigation_too_large', '文档超过编辑上限。')
        if set(node) - {'type','attrs','content','marks','text'}: raise _failure()
        attrs = node.get('attrs', {})
        if not isinstance(attrs, dict): raise _failure()
        accepted = {'paragraph': {'textAlign','indent','lineHeight'}, 'heading': {'level','textAlign','indent','lineHeight'},
                    'orderedList': {'start','type'}, 'tableCell': {'colspan','rowspan','colwidth','align'},
                    'tableHeader': {'colspan','rowspan','colwidth','align'}}.get(kind,set())
        if set(attrs) - accepted: raise _failure('litigation_editor_unsupported', '暂不支持保存此格式。')
        if attrs.get('textAlign') not in (None, 'left','center','right','justify'): raise _failure()
        if 'level' in attrs and (type(attrs['level']) is not int or attrs['level'] not in range(1,7)): raise _failure()
        if 'indent' in attrs and (type(attrs['indent']) not in (int,float) or not 0 <= attrs['indent'] <= 144): raise _failure()
        if 'lineHeight' in attrs and attrs['lineHeight'] is not None and attrs['lineHeight'] not in ('1','1.5','2'): raise _failure()
        if 'start' in attrs and (type(attrs['start']) is not int or not 1 <= attrs['start'] <= 10000): raise _failure()
        if attrs.get('type') not in (None, '1'): raise _failure('litigation_editor_unsupported', '暂不支持此编号格式。')
        if kind in ('tableCell','tableHeader') and (attrs.get('colspan',1) != 1 or attrs.get('rowspan',1) != 1):
            raise _failure('litigation_editor_unsupported', '合并单元格请保留在原文件中查看。')
        if kind in ('tableCell','tableHeader') and attrs.get('colwidth') is not None:
            widths=attrs['colwidth']
            if not isinstance(widths,list) or len(widths)!=1 or type(widths[0]) is not int or not 1<=widths[0]<=2000:raise _failure('litigation_editor_unsupported','表格列宽 colwidth 必须是单个正整数的数组，例如 [45]，不能写成数字 45。')
        marks = node.get('marks', [])
        if not isinstance(marks, list) or (marks and kind != 'text'): raise _failure()
        for mark in marks:
            if not isinstance(mark, dict) or set(mark)-{'type','attrs'} or mark.get('type') not in {'bold','italic','underline','strike','textStyle'}: raise _failure()
            ma = mark.get('attrs', {})
            if not isinstance(ma,dict) or set(ma)-({'fontFamily','fontSize'} if mark['type']=='textStyle' else set()): raise _failure()
            if ma.get('fontFamily') not in (None,'宋体','黑体','仿宋','楷体','SimSun','Arial','Times New Roman'): raise _failure()
            size = ma.get('fontSize')
            if size is not None and size not in [f'{n}pt' for n in (10,12,14,16,18,22,24,26,28,36)]: raise _failure()
        if kind == 'text':
            if not isinstance(node.get('text'),str) or not node['text'] or node.get('content'): raise _failure()
            try: etree.Element('t').text = node['text']
            except ValueError: raise _failure() from None
        children = node.get('content', [])
        if not isinstance(children,list): raise _failure()
        for child in children:
            if not isinstance(child,dict) or child.get('type') not in containers.get(kind,set()): raise _failure()
            visit(child,depth+1)
        if kind in {'doc','table','tableRow','tableCell','tableHeader','listItem','orderedList','bulletList'} and not children: raise _failure()
        if kind=='table' and len({len(r['content']) for r in children}) != 1: raise _failure()
    if not isinstance(value,dict) or value.get('type')!='doc': raise _failure()
    visit(value)
    if len(document_text(value)) > MAX_TEXT: raise _failure('litigation_too_large', '文档超过编辑上限。')
    return value


def document_text(node):
    if node.get('type')=='text': return node.get('text','')
    if node.get('type')=='hardBreak': return '\n'
    sep = '' if node.get('type') in {'paragraph','heading'} else '\n'
    return sep.join(document_text(c) for c in node.get('content',[]))


def render_docx(value: dict, *, compact: bool = False, pleading: bool = False) -> bytes:
    validate_document(value)
    doc = Document()
    normal = doc.styles['Normal']; normal.font.name = '宋体'; normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(f'{{{W}}}eastAsia','宋体')
    normal.paragraph_format.line_spacing = 1.5
    if compact:
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
    if pleading:
        from .drafting_format import configure
        configure(doc)
    def blocks(parent, nodes, list_id=None, level=0):
        for node in nodes:
            kind=node['type']; attrs=node.get('attrs',{})
            if kind in ('paragraph','heading'):
                p=parent.add_paragraph(style=f'Heading {attrs.get("level",1)}' if kind=='heading' else None)
                if list_id is not None:
                    np=etree.SubElement(p._p.get_or_add_pPr(),f'{{{W}}}numPr')
                    etree.SubElement(np,f'{{{W}}}ilvl').set(f'{{{W}}}val',str(level))
                    etree.SubElement(np,f'{{{W}}}numId').set(f'{{{W}}}val',str(list_id))
                    # Later paragraphs belong to the same item, not a new number.
                    list_id = None
                if attrs.get('textAlign'): p.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}[attrs['textAlign']]
                if pleading:
                    from .drafting_format import paragraph_defaults
                    paragraph_defaults(p, ''.join(c.get('text','') for c in node.get('content',[])), in_table=parent is not doc, numbered=list_id is not None or p._p.find(f'{{{W}}}pPr/{{{W}}}numPr') is not None)
                if attrs.get('indent'): p.paragraph_format.left_indent=Pt(attrs['indent'])
                if attrs.get('lineHeight'): p.paragraph_format.line_spacing=float(attrs['lineHeight'])
                for child in node.get('content',[]):
                    r=p.add_run('\n' if child['type']=='hardBreak' else child['text'])
                    for mark in child.get('marks',[]):
                        name=mark['type']
                        if name in ('bold','italic','underline'): setattr(r,name,True)
                        elif name=='strike': r.font.strike=True
                        elif name=='textStyle':
                            a=mark.get('attrs',{})
                            if a.get('fontFamily'):
                                r.font.name=a['fontFamily']; r._r.get_or_add_rPr().rFonts.set(f'{{{W}}}eastAsia',a['fontFamily'])
                            if a.get('fontSize'): r.font.size=Pt(int(a['fontSize'][:-2]))
            elif kind in ('bulletList','orderedList'):
                numbering=doc.part.numbering_part.element
                ids=[int(n.get(f'{{{W}}}abstractNumId')) for n in numbering.findall(f'{{{W}}}abstractNum')]
                aid=max(ids,default=-1)+1
                abstract=etree.SubElement(numbering,f'{{{W}}}abstractNum'); abstract.set(f'{{{W}}}abstractNumId',str(aid))
                lvl=etree.SubElement(abstract,f'{{{W}}}lvl'); lvl.set(f'{{{W}}}ilvl','0')
                for tag,val in [('start',str(attrs.get('start',1))),('numFmt','decimal' if kind=='orderedList' else 'bullet'),('lvlText','%1.' if kind=='orderedList' else '•')]:
                    etree.SubElement(lvl,f'{{{W}}}{tag}').set(f'{{{W}}}val',val)
                num=numbering.add_num(aid)
                for item in node['content']: blocks(parent,item['content'],num.numId,0)
            elif kind=='table':
                rows=node['content']; table=parent.add_table(rows=len(rows),cols=len(rows[0]['content'])); table.style='Table Grid'
                widths=[cell.get('attrs',{}).get('colwidth') for cell in rows[0]['content']]
                if all(widths):
                    available=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
                    total=sum(width[0] for width in widths)
                    table.autofit=False
                    from docx.shared import Emu
                    for j,width in enumerate(widths):
                        length=Emu(round(available*width[0]/total))
                        table.columns[j].width=length
                        for row in table.rows:row.cells[j].width=length
                for i,row in enumerate(rows):
                    for j,cell in enumerate(row['content']):
                        target=table.cell(i,j); empty=target.paragraphs[0]._p
                        blocks(target,cell['content']); target._tc.remove(empty)
    blocks(doc,value['content'])
    out=io.BytesIO(); doc.save(out); return checked_bytes(out.getvalue())


def import_docx(source: bytes, *, text_only=False) -> dict:
    """Import a supported simple document or an explicitly requested text copy.

Unknown structure is rejected rather than dropped. Imported files are never
written back; original-package text patching is a separate operation.
"""
    data=checked_bytes(source);inspection=inspect_docx(data)
    if text_only:
        return {'type':'doc','content':[{'type':'paragraph',**({'content':[{'type':'text','text':p['text']}]} if p['text'] else {})} for p in inspection['paragraphs']] or [{'type':'paragraph'}]}
    root=xml(parts(data)[MAIN])
    allowed={f'{{{W}}}{name}' for name in ('document','body','p','pPr','r','rPr','t','br','tab','tbl','tblPr','tblGrid','gridCol','tr','trPr','tc','tcPr','sectPr')}
    # Property descendants remain in the original, but unsupported body objects
    # require explicit text extraction. Conversion is a new editable copy.
    for node in root.iter():
        if any(etree.QName(a).localname in {'pPr','rPr','tblPr','tblGrid','trPr','tcPr','sectPr'} for a in node.iterancestors()):continue
        if node.tag not in allowed or node.tag==f'{{{W}}}tab':
            raise _failure('litigation_editor_readonly','此Word含有暂不支持转换的对象，请查看原件或明确提取正文副本。')
    document=Document(io.BytesIO(data))
    def paragraph(p):
        if p._p.find(f'.//{{{W}}}numPr') is not None:
            raise _failure('litigation_editor_readonly','此Word包含自动编号，请保留原件并提取正文副本。')
        content=[]
        for run in p.runs:
            marks=[]
            for name in ('bold','italic','underline'):
                if getattr(run,name):marks.append({'type':name})
            if run.font.strike:marks.append({'type':'strike'})
            attrs={}
            if run.font.name in ('宋体','黑体','仿宋','楷体','Arial','Times New Roman'):attrs['fontFamily']=run.font.name
            if run.font.size and run.font.size.pt in (10,12,14,16,18,22,24,26,28,36):attrs['fontSize']=f'{int(run.font.size.pt)}pt'
            if attrs:marks.append({'type':'textStyle','attrs':attrs})
            chunks=run.text.split('\n')
            for i,chunk in enumerate(chunks):
                if i:content.append({'type':'hardBreak'})
                if chunk:content.append({'type':'text','text':chunk,**({'marks':marks} if marks else {})})
        attrs={};kind='paragraph'
        if p.style and p.style.name.startswith('Heading ') and p.style.name[-1:] in '123456':
            kind='heading';attrs['level']=int(p.style.name[-1])
        align={WD_ALIGN_PARAGRAPH.LEFT:'left',WD_ALIGN_PARAGRAPH.CENTER:'center',WD_ALIGN_PARAGRAPH.RIGHT:'right',WD_ALIGN_PARAGRAPH.JUSTIFY:'justify'}.get(p.alignment)
        if align:attrs['textAlign']=align
        indent=p.paragraph_format.left_indent
        if indent is not None and 0<=indent.pt<=144:attrs['indent']=indent.pt
        spacing=p.paragraph_format.line_spacing
        if spacing in (1,1.5,2):attrs['lineHeight']=str(spacing).rstrip('0').rstrip('.')
        return {'type':kind,**({'attrs':attrs} if attrs else {}),**({'content':content} if content else {})}
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    def blocks(parent, element):
        result=[]
        for child in element:
            if child.tag==f'{{{W}}}p':result.append(paragraph(Paragraph(child,parent)))
            elif child.tag==f'{{{W}}}tbl':
                table=Table(child,parent);rows=[]
                for row in table.rows:
                    cells=[]
                    for cell in row.cells:
                        if cell._tc.find(f'.//{{{W}}}gridSpan') is not None or cell._tc.find(f'.//{{{W}}}vMerge') is not None or cell.tables:
                            raise _failure('litigation_editor_readonly','复杂表格保留在原件中，可提取正文副本。')
                        cells.append({'type':'tableCell','content':blocks(cell,cell._tc) or [{'type':'paragraph'}]})
                    rows.append({'type':'tableRow','content':cells})
                result.append({'type':'table','content':rows})
        return result
    value={'type':'doc','content':blocks(document,document.element.body) or [{'type':'paragraph'}]}
    return validate_document(value)
