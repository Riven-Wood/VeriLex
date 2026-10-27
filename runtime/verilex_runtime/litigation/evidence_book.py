"""Ordered material selection, independent evidence booklet and six-column directory."""
import io
import json
import hashlib
import uuid
from pathlib import Path
from typing import Literal
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from .service import safe_path, _LOCK, _failure
from .evidence_conversion import convert, pdf_pages, plugin_job, SUPPORTED, MAX_BYTES

class Entry(BaseModel):
    model_config=ConfigDict(extra='forbid',strict=True)
    name:str=Field(min_length=1,max_length=200)
    group:str=Field(default='',max_length=200)
    purpose:str=Field(default='',max_length=2000)
    content:str=Field(default='',max_length=2000)
    source:str=Field(default='',max_length=2000)
    path:str=Field(min_length=1,max_length=1000)
    first:int=Field(default=1,ge=1,le=10000)
    last:int|None=Field(default=None,ge=1,le=10000)
    file_hash:str=Field(default='',max_length=64)
    page_count:int|None=Field(default=None,ge=1,le=1000)
    manual:bool=False
    manual_fields:list[Literal["name","content","purpose","source"]]=Field(default_factory=list,max_length=4)
    status:str=Field(default='pending',pattern='^(pending|ready|unread)$')
    gaps:list[str]=Field(default_factory=list,max_length=1000)

class EvidenceGroup(BaseModel):
    model_config=ConfigDict(extra='forbid',strict=True)
    id:str=Field(min_length=1,max_length=200)
    name:str=Field(min_length=1,max_length=200)
    purpose:str=Field(default='',max_length=2000)
    numbering:Literal['continue','restart']='continue'
    start_number:int=Field(default=1,ge=1,le=10000)

class Inventory(BaseModel):
    model_config=ConfigDict(extra='forbid',strict=True)
    title:str=Field(default='证据目录',min_length=1,max_length=200)
    entries:list[Entry]=Field(default_factory=list,max_length=200)
    groups:list[EvidenceGroup]=Field(default_factory=list,max_length=200)

    @model_validator(mode='after')
    def arrange_groups(self):
        ids=[group.id for group in self.groups]
        if len(ids)!=len(set(ids)):raise ValueError('重复的证据分组')
        # Older inventories stored a group label directly on each entry.
        for entry in self.entries:
            if entry.group and entry.group not in ids:
                self.groups.append(EvidenceGroup(id=entry.group,name=entry.group));ids.append(entry.group)
        if len(self.groups)>200:raise ValueError('证据分组过多')
        order={value:index for index,value in enumerate(ids)}
        self.entries.sort(key=lambda entry:order.get(entry.group,len(ids)))
        return self


def numbered_entries(value):
    groups={group.id:group for group in value.groups} if len(value.groups)>1 else {};previous=None;number=1
    for entry in value.entries:
        group=groups.get(entry.group)
        if entry.group!=previous and group and group.numbering=='restart':number=group.start_number
        yield number,entry,group,entry.group!=previous
        previous=entry.group;number+=1


def location(service,ctx):return safe_path(service._root(ctx),'文书起草/编辑记录/evidence-directory.json')


def read(service,ctx):
    with _LOCK:
        path=location(service,ctx)
        if not path.exists():return {'revision':0,'data':Inventory().model_dump()}
        try:
            if path.stat().st_size>4*1024*1024:raise ValueError()
            record=json.loads(path.read_bytes())
            if type(record['revision']) is not int:raise ValueError()
            value=Inventory.model_validate(record['data']).model_dump()
            service._register_internal(ctx,path)
            return {**record,'data':value}
        except (ValueError,KeyError,TypeError):raise _failure('litigation_directory_invalid','证据清单无法读取，原文件已保留。') from None


def source_bytes(service,ctx,item):
    source=safe_path(service._root(ctx),item.path)
    if not source.is_file() or source.stat().st_size>20*1024*1024:
        raise _failure('litigation_source_invalid',f'{item.name}：文件不存在或超过 20 MB。')
    return source,source.read_bytes()


def persist(service,ctx,record,expected):
    path=location(service,ctx)
    if read(service,ctx)['revision']!=expected:raise _failure('litigation_file_conflict','证据清单已有更新，请重新读取，未覆盖你的修改。')
    service.products._publish(ctx,path,json.dumps(record,ensure_ascii=False).encode(),path.read_bytes() if path.exists() else None,internal=True)
    return record


def save(service,ctx,data,expected):
    with _LOCK:
        service._root(ctx,write=True);current=read(service,ctx)
        if current['revision']!=expected:raise _failure('litigation_file_conflict','证据清单已有更新，请重新读取，未覆盖你的修改。')
        try:value=Inventory.model_validate(data)
        except ValidationError:raise _failure('litigation_directory_invalid','请填写证据名称、文件和有效页范围。') from None
        previous={item['path']:item for item in current['data']['entries']}
        for item in value.entries:
            source,raw=source_bytes(service,ctx,item)
            if source.suffix.lower() not in SUPPORTED:raise _failure('litigation_source_invalid',f'{item.name}：暂不支持此格式。')
            if item.last is not None and item.first>item.last:raise _failure('litigation_source_invalid','结束页不能早于开始页。')
            digest=hashlib.sha256(raw).hexdigest();old=previous.get(item.path,{})
            item.page_count=old.get('page_count') if old.get('file_hash')==digest else None
            if old.get('file_hash') and old['file_hash']!=digest:
                item.status='pending';item.gaps=['原文件已有更新，请重新核对目录文字。']
            if old and (old.get('first',1)!=item.first or old.get('last')!=item.last):
                item.status='pending';item.gaps=['所选页范围已改变，请重新核对目录文字。']
            item.file_hash=digest
        return persist(service,ctx,{'revision':expected+1,'data':value.model_dump()},expected)


def cache_path(service,ctx,item):
    return safe_path(service._root(ctx),'文书起草/编辑记录/证据缓存/'+item.file_hash+'.pdf')


def prepared(service,ctx,item):
    source,raw=source_bytes(service,ctx,item)
    if hashlib.sha256(raw).hexdigest()!=item.file_hash:
        raise _failure('litigation_file_conflict',f'{item.name}已有更新，请重新保存清单后核对。')
    cached=cache_path(service,ctx,item)
    if cached.is_file() and cached.stat().st_size<=MAX_BYTES:
        data=cached.read_bytes();pdf_pages(data);return data
    try:data=convert(raw,source.suffix.lower())
    except Exception as error:
        from verilex_runtime.tools.registry import ToolError
        if isinstance(error,ToolError):raise _failure(error.code,f'{item.name}：{error.message}') from None
        raise
    with _LOCK:
        if not cached.exists():service.products._publish(ctx,cached,data,internal=True)
    return data


def prepare(service,ctx,expected):
    service._root(ctx,write=True)
    record=read(service,ctx)
    if record['revision']!=expected:raise _failure('litigation_file_conflict','证据清单已有更新，请重新读取。')
    value=Inventory.model_validate(record['data']);total=0
    for item in value.entries:
        data=prepared(service,ctx,item);total+=len(data)
        if total>MAX_BYTES:raise _failure('litigation_too_large','转换后证据总量超过 100 MB，请分册。')
        item.page_count=pdf_pages(data)
        if item.first>item.page_count or (item.last is not None and item.last>item.page_count):raise _failure('litigation_source_invalid',f'{item.name}：页范围超出文件实际页数。')
    mapping=page_mapping(value)
    with _LOCK:
        # Slow conversion must never overwrite a newer selection or source.
        for item in value.entries:
            if hashlib.sha256(source_bytes(service,ctx,item)[1]).hexdigest()!=item.file_hash:raise _failure('litigation_file_conflict',f'{item.name}在转换时已有更新，请重试。')
        if value.model_dump()==record['data']:return {**record,'mapping':mapping}
        result=persist(service,ctx,{'revision':expected+1,'data':value.model_dump()},expected)
        return {**result,'mapping':mapping}


def page_mapping(value):
    cursor=1;rows=[]
    for item in value.entries:
        if item.page_count is None:raise _failure('litigation_pages_missing','请先转换并核对页数。')
        last=item.last or item.page_count;count=last-item.first+1
        if count<1 or last>item.page_count:raise _failure('litigation_source_invalid',f'{item.name}：页范围无效。')
        rows.append({'name':item.name,'path':item.path,'first':cursor,'last':cursor+count-1,'source_first':item.first,'source_last':last})
        cursor+=count
    if cursor>1001:raise _failure('litigation_too_large','证据册超过 1000 页，请分册。')
    return rows


def directory_docx(value,mapping):
    # Styles only from the supplied grouped directory; no parties or case content.
    from docx import Document
    from docx.shared import Twips,Pt,RGBColor
    from docx.enum.text import WD_LINE_SPACING
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();section=doc.sections[0]
    section.page_width=Twips(15840);section.page_height=Twips(12240)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Twips(720)
    style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(11);style.font.color.rgb=RGBColor(0,0,0)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'楷体')
    style.paragraph_format.space_after=Pt(0);style.paragraph_format.space_before=Pt(0)
    style.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY;style.paragraph_format.line_spacing=Pt(16)
    p=doc.add_paragraph(value.title);p.alignment=1
    run=p.runs[0];run.bold=True;run.font.size=Pt(14);run.font.name='黑体';run._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
    doc.add_paragraph('AI 草稿，须律师审查。页码对应单独生成的 PDF 证据册。')
    table=doc.add_table(rows=1,cols=6);table.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        element=OxmlElement('w:'+edge)
        for key,val in {'val':'single','sz':'4','color':'000000'}.items():element.set(qn('w:'+key),val)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    widths=[941,2551,1972,2640,4932,933]
    for col,width in zip(table.columns,widths):col.width=Twips(width)
    for cell,text,width in zip(table.rows[0].cells,['编号','证据名称','证据内容','证明目的','证据来源','页码'],widths):
        cell.text=text;cell.width=Twips(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment=1;cell.paragraphs[0].runs[0].bold=True
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for (number,item,group,first),pages in zip(numbered_entries(value),mapping):
        if first and len(value.groups)>1:
            cells=table.add_row().cells;cell=cells[0].merge(cells[-1])
            cell.text=(group.name+('，证明：'+group.purpose if group.purpose else '') if group else '未分组证据')
            p=cell.paragraphs[0];p.runs[0].bold=True;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=True
        content=item.content or '待核对'
        if item.gaps:content+='\n待核对：'+'；'.join(item.gaps)
        texts=[str(number),item.name,content,item.purpose or '待核对',item.source or '待核对',str(pages['first']) if pages['first']==pages['last'] else f"{pages['first']}—{pages['last']}"]
        for cell,text,width in zip(table.add_row().cells,texts,widths):cell.text=text;cell.width=Twips(width)
    out=io.BytesIO();doc.save(out);return out.getvalue()


def export(service,ctx,expected,kind):
    if kind not in ('pdf','docx'):raise _failure()
    record=prepare(service,ctx,expected);value=Inventory.model_validate(record['data'])
    if not value.entries:raise _failure('litigation_empty','请先在左侧勾选证据文件。')
    mapping=record['mapping']
    if kind=='pdf':
        inputs={f'{i:04}.pdf':prepared(service,ctx,item) for i,item in enumerate(value.entries)}
        data=plugin_job({'operation':'merge','inputs':[{'name':f'{i:04}.pdf','first':row['source_first'],'last':row['source_last']} for i,row in enumerate(mapping)]},inputs)
        if pdf_pages(data)!=mapping[-1]['last']:raise _failure('litigation_conversion_failed','合并页数与清单不一致，未交付。')
    else:data=directory_docx(value,mapping)
    with _LOCK:
        if read(service,ctx)['revision']!=record['revision']:raise _failure('litigation_file_conflict','生成期间清单已有更新，请重新生成。')
        for item in value.entries:
            if hashlib.sha256(source_bytes(service,ctx,item)[1]).hexdigest()!=item.file_hash:raise _failure('litigation_file_conflict',f'{item.name}已有更新，未交付旧文件。')
        name=f'证据{"册" if kind=="pdf" else "目录"}-{uuid.uuid4().hex[:8]}.{kind}'
        service.products._publish(ctx,safe_path(service._root(ctx),'文书起草/'+name),data)
    return {'path':'文书起草/'+name,'directory_pages':0,'entries':mapping,'pages':mapping[-1]['last'],'record':record}
def directory_pdf(lines):
    """Embed the pinned OFL Chinese font; the result works without local fonts."""
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font='VeriLexNotoSansSC'
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font,str(Path(__file__).parent/'fonts/NotoSansSC.ttf')))
    output=io.BytesIO();canvas=Canvas(output,pagesize=(595,842),invariant=1,pageCompression=1)
    for offset in range(0,len(lines),35):
        canvas.setFont(font,11)
        for index,line in enumerate(lines[offset:offset+35]):canvas.drawString(45,795-index*20,line)
        canvas.showPage()
    canvas.save();return output.getvalue()
