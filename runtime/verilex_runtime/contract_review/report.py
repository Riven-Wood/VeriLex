"""Independent opinion letter; only numeric layout rules distilled from the user's reference.

No template document, sample wording, metadata, or model call is used at runtime.
The service's citation and verified-comment gates MUST run before this builder.
"""
from __future__ import annotations

import io
import re
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor

from verilex_runtime.contract_review.errors import ContractReviewError
from verilex_runtime.contract_review.opinions import Opinion
from verilex_runtime.documents.docx_package import checked_bytes


def _parties(original: bytes) -> str:
    """Only explicit, short party-name lines; do not infer entities from prose."""
    document = Document(io.BytesIO(original))
    candidates: dict[str, list[str]] = {}
    for paragraph in document.element.body.iter(qn("w:p")):
        text = "".join(
            node.text or "" for node in paragraph.iter(qn("w:t"))
        ).strip()
        match = re.fullmatch(r"([甲乙丙丁]方)\s*[：:]\s*([^\n；;。]{2,100})", text)
        if match:
            candidates.setdefault(match[1], []).append(match[2].strip())
    names = [f"{key}：{values[0]}" for key, values in candidates.items()
             if len(set(values)) == 1]
    return "\n".join(names) if names else "待补充"


def report_paragraphs(review: dict, opinions: Sequence[Opinion], original: bytes,
                      issued_at: str) -> list[tuple[str, str]]:
    day = datetime.fromisoformat(issued_at.replace("Z", "+00:00")).astimezone().date()
    parties = _parties(original).split("\n")
    rows = [
        ("title", "合同审核意见书"),
        ("blank", ""),
        ("metadata", f"合同名称：{Path(review['filename']).stem}"),
        ("metadata", f"合同签订方：{parties[0]}"),
        *[("party", name) for name in parties[1:]],
        ("metadata", "审查立场：" + (review.get("custom_party") or {"party_a": "甲方", "party_b": "乙方", "neutral": "中立"}[review["party"]])),
        ("heading", "审核意见："),
    ]
    for number, opinion in enumerate(opinions, 1):
        level = f"【{opinion.level}】" if opinion.level else ""
        rows.append(("opinion", f"{number}、{level}{opinion.category}"))
        # Use the exact anchor validated against this contract, never a model's
        # paraphrased clause_excerpt or an unverified clause heading.
        rows.append(("anchor", f"合同原文：{opinion.anchor_text}"))
        rows.append(("body", opinion.edited_text or opinion.body))
        if opinion.suggestion:
            rows.append(("body", f"修改建议：{opinion.suggestion}"))
        for citation in opinion.citations:
            rows.append(("body", f"参考来源：{citation.document}" + (f"（来源：{citation.source_label}，读取于 {citation.verified_at}）" if citation.source_call_id else "（未关联读取记录）")))
    rows.extend([
        ("blank", ""),
        ("signature", f"审查署名：{review['signature']}"),
        ("date", f"出具日期：{day.year}年{day.month}月{day.day}日"),
    ])
    return rows


def build_report(review: dict, opinions: Sequence[Opinion], original: bytes,
                 issued_at: str) -> bytes:
    try:
        rows = report_paragraphs(review, opinions, original, issued_at)
        document = Document()
        # python-docx's blank package includes an unused bibliography customXml.
        # This independent report has no hidden content or bibliography parts.
        for relationship in list(document.part.rels.values()):
            if relationship.reltype.endswith("/customXml"):
                document.part.drop_rel(relationship.rId)
        section = document.sections[0]
        section.page_width, section.page_height = Twips(11906), Twips(16838)
        section.top_margin = section.bottom_margin = Twips(1440)
        section.left_margin = section.right_margin = Twips(1800)
        section.header_distance = section.footer_distance = Twips(0)
        normal = document.styles["Normal"]
        normal.font.name = "宋体"
        normal.font.size = Pt(14)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        normal.paragraph_format.line_spacing = Pt(25)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.widow_control = True
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for index, (role, text) in enumerate(rows):
            paragraph = document.add_paragraph()
            if role == "title":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif role in {"signature", "date"}:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif role in {"body", "opinion", "anchor"}:
                paragraph.paragraph_format.first_line_indent = Pt(28)
            elif role == "party":
                paragraph.paragraph_format.left_indent = Pt(84)
            paragraph.paragraph_format.keep_with_next = (
                role in {"title", "heading", "opinion", "anchor", "signature", "blank"}
                or index == len(rows) - 4
            )
            if role == "metadata":
                label, value = text.split("：", 1)
                paragraph.add_run(label + "：").bold = True
                paragraph.add_run(value)
            else:
                run = paragraph.add_run(text)
                run.bold = role in {"title", "heading", "opinion"}
                if role == "title":
                    run.font.size = Pt(16)
        document.core_properties.author = review["signature"]
        document.core_properties.title = "合同审核意见书"
        document.core_properties.subject = ""
        document.core_properties.comments = ""
        document.core_properties.last_modified_by = ""
        document.core_properties.created = document.core_properties.modified = datetime.fromisoformat(issued_at.replace("Z", "+00:00"))
        buffer = io.BytesIO()
        document.save(buffer)
        data = checked_bytes(buffer.getvalue())
        verify_report(data, [text for _, text in rows])
        return data
    except ContractReviewError:
        raise
    except Exception:
        raise ContractReviewError("contract_report_failed") from None


def verify_report(data: bytes, expected: Sequence[str]) -> None:
    """Reopen the saved package; ensure complete text and no hidden deliveries."""
    try:
        document = Document(io.BytesIO(data))
        if [p.text for p in document.paragraphs] != list(expected):
            raise ValueError("report text mismatch")
        if document.tables or document.inline_shapes or len(document.sections) != 1:
            raise ValueError("unexpected report structure")
        with ZipFile(io.BytesIO(data)) as package:
            if any(name.startswith(("word/comments", "word/header", "word/footer", "word/embeddings", "customXml/")) for name in package.namelist()):
                raise ValueError("unexpected report parts")
    except Exception:
        raise ContractReviewError("contract_report_failed") from None
