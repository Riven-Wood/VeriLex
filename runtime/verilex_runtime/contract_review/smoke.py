"""Content-free packaged check using an internally generated fictional contract."""
from __future__ import annotations

import io
import json
import tempfile
from pathlib import Path


def run() -> int:
    stage = "create"
    try:
        from docx import Document
        from .annotator import CommentSpec, ContractAnnotator
        from .opinions import Opinion
        from .preview import build_preview
        from .report import build_report

        document = Document()
        anchor = "虚构合同：成果交付后由甲方验收。"
        document.add_paragraph(anchor)
        buffer = io.BytesIO()
        document.save(buffer)
        original = buffer.getvalue()
        stage = "preview_without_existing_comments"
        assert build_preview(original)["paragraphs"] == [anchor]
        stage = "first_comment"
        with tempfile.TemporaryDirectory() as folder:
            source, dest = Path(folder) / "source.docx", Path(folder) / "copy.docx"
            source.write_bytes(original)
            ContractAnnotator().write(source, [CommentSpec(anchor, "建议明确期限。", "验收")], dest)
            assert len(Document(dest).comments) == 1
            assert source.read_bytes() == original
        stage = "report"
        report = build_report(
            {"filename": "虚构合同.docx", "signature": "验收", "party": "neutral"},
            [Opinion("smoke", "验收", anchor, anchor, body="建议明确期限。")],
            original, "2026-10-03T00:00:00Z",
        )
        result = Document(io.BytesIO(report))
        assert result.paragraphs[0].text == "合同审核意见书"
        # Assert the final pagination rule is present in the frozen module too.
        assert result.paragraphs[-4].paragraph_format.keep_with_next is True
        assert result.paragraphs[-3].paragraph_format.keep_with_next is True
    except Exception as error:
        print(json.dumps({"ok": False, "stage": stage, "error_type": type(error).__name__}))
        return 2
    print(json.dumps({"ok": True, "preview": True, "first_comment": True, "report": True}))
    return 0
