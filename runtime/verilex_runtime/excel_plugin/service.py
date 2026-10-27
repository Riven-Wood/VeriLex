"""Narrow, local adapter around the vendored MIT Excel MCP cell operations.

Only .xlsx, bounded reads, and additive outputs. No formula calculation or
network access. The explicitly selected MCP folder is the plugin boundary.
"""
from __future__ import annotations

from copy import copy
import io
import json
import os
import stat
from pathlib import Path
import math
import threading
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pydantic import BaseModel, ConfigDict, Field

from .vendor.operations.cells import read_range, write_range
from .vendor.values import CellInput, TypedCell

MAX_FILE = 25 * 1024 * 1024
MAX_EXPANDED = 64 * 1024 * 1024
MAX_CELLS = 100_000
NOTICE = '公式仅保存和读取，不重算；缓存值可能为空或过期。图表不预览、不编辑。原件保留。'


class WriteBlock(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    sheet: str = Field(min_length=1, max_length=31)
    start_cell: str = 'A1'
    rows: list[list[CellInput]] = Field(min_length=1, max_length=2000)
    copy_style_from_row: int | None = Field(default=None, ge=1, le=1048576)


class ExcelPlugin:
    def __init__(self, root: Path):
        self.root = root.resolve(strict=True)
        if not self.root.is_dir() or self.root == Path(self.root.anchor):
            raise ValueError('请选择专门存放表格的文件夹。')
        self._lock = threading.Lock()

    def _path(self, name: str, *, new: bool = False) -> Path:
        path = Path(name)
        if not path.is_absolute():
            path = self.root / path
        # Deny symlinks in any descendant, even when their target is in root.
        try:
            relative = path.relative_to(self.root)
        except ValueError:
            raise ValueError('只能处理已选择的表格文件夹内的文件。') from None
        cursor = self.root
        for part in relative.parts:
            if part in ('.', '..'):
                raise ValueError('文件路径不允许返回上级目录。')
            cursor /= part
            if cursor.is_symlink():
                raise ValueError('不支持符号链接，请选择普通文件。')
        resolved = path.resolve(strict=not new)
        if not resolved.is_relative_to(self.root) or resolved.suffix.lower() != '.xlsx':
            raise ValueError('仅支持选定文件夹内的 .xlsx；请先将旧格式或宏文件另存为 .xlsx。')
        if new and (resolved.exists() or not resolved.parent.is_dir()):
            raise ValueError('成果必须使用尚不存在的文件名，并保存到已有文件夹。')
        if not new and (not resolved.is_file() or resolved.stat().st_size > MAX_FILE):
            raise ValueError('文件不存在或超过 25 MB。')
        return resolved

    def _load(self, path: Path, *, editing=False):
        directory = os.open(self.root, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
        try:
            relative = path.relative_to(self.root)
            for part in relative.parts[:-1]:
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=directory)
                os.close(directory)
                directory = child
            descriptor = os.open(relative.name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory)
            with os.fdopen(descriptor, 'rb') as handle:
                if not stat.S_ISREG(os.fstat(handle.fileno()).st_mode):
                    raise ValueError('仅支持普通工作簿文件。')
                data = handle.read(MAX_FILE + 1)
        finally:
            os.close(directory)
        if len(data) > MAX_FILE:
            raise ValueError('工作簿超过 25 MB。')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise ValueError('工作簿解压后过大，无法安全处理。')
            names = [i.filename.lower() for i in infos]
            if any('vbaproject' in n or n.startswith('xl/embeddings/') for n in names):
                raise ValueError('不处理包含宏或嵌入程序的工作簿。')
            complex_parts = any(n.startswith(('xl/drawings/', 'xl/charts/', 'xl/pivot', 'xl/externallinks/', 'xl/activex/', 'xl/slicer', 'xl/threadedcomments/', 'customxml/')) for n in names)
            cells = 0
            for item in infos:
                if item.filename.startswith('xl/') and item.filename.endswith('.xml'):
                    payload = archive.read(item)
                    if b'<!DOCTYPE' in payload.upper() or b'<!ENTITY' in payload.upper():
                        raise ValueError('工作簿包含不支持的 XML 定义。')
                    xml = ET.fromstring(payload)
                    cells += sum(1 for element in xml.iter() if element.tag.rpartition('}')[2] == 'c')
                    if cells > MAX_CELLS:
                        raise ValueError('工作簿超过十万个单元格，请拆分后处理。')
                    if editing and any(e.tag.rpartition('}')[2] == 'extLst' for e in xml.iter()):
                        complex_parts = True
            if editing and complex_parts:
                raise ValueError('此表含图表、透视表、外部链接或扩展部件；当前插件仅支持读取，请将需要处理的数据另存到简单工作簿。')
        workbook = load_workbook(io.BytesIO(data), keep_links=False)
        if len(workbook.worksheets) > 50:
            workbook.close()
            raise ValueError('最多处理 50 张工作表。')
        return workbook, data, complex_parts

    def list_files(self):
        paths = sorted(p for p in self.root.iterdir() if p.is_file() and not p.is_symlink() and p.suffix.lower() == '.xlsx')
        return {'folder': str(self.root), 'files': [p.name for p in paths[:100]], 'truncated': len(paths) > 100}

    def inspect(self, path: str, sheet: str | None = None, cell_range: str | None = None):
        workbook, data, complex_parts = self._load(self._path(path))
        cached = None
        try:
            result = {'path': path, 'sheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column, 'charts': len(s._charts), 'merged_ranges': [str(m) for m in list(s.merged_cells.ranges)[:100]]} for s in workbook.worksheets], 'complex_parts': complex_parts, 'notice': NOTICE}
            if sheet is not None:
                if sheet not in workbook.sheetnames:
                    raise ValueError('未找到这张工作表，请先查看 sheets。')
                values = read_range(workbook[sheet], cell_range, 2000).model_dump()
                cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
                values['cached_values'] = read_range(cached[sheet], values['range'], 2000).values
                result['data'] = values
            # Do not silently truncate individual cell strings into plausible data.
            if len(json.dumps(result, ensure_ascii=False).encode()) > 100_000:
                raise ValueError('读取内容过长，请缩小 cell_range。')
            return result
        finally:
            workbook.close()
            if cached is not None:
                cached.close()

    def deliver(self, output: str, writes: list[WriteBlock], source: str | None = None):
        if not writes or len(writes) > 50 or sum(len(row) for block in writes for row in block.rows) > 10000:
            raise ValueError('每次需提供 1–50 个写入区域，总计不超过一万个单元格。')
        for block in writes:
            for row in block.rows:
                for item in row:
                    value = item.value if isinstance(item, TypedCell) else item
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValueError('单元格文字超过 Excel 长度限制。')
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError('数字必须是有限数值。')
        with self._lock:
            target = self._path(output, new=True)
            workbook = None
            if source:
                workbook, _, _ = self._load(self._path(source), editing=True)
            else:
                workbook = Workbook()
                workbook.active.title = writes[0].sheet
            try:
                for block in writes:
                    if block.sheet not in workbook.sheetnames:
                        workbook.create_sheet(block.sheet)
                if len(workbook.worksheets) > 50:
                    raise ValueError('最多交付 50 张工作表。')
                for block in writes:
                    sheet = workbook[block.sheet]
                    written = write_range(sheet, block.start_cell, block.rows, 10000)
                    if block.copy_style_from_row is not None:
                        from .vendor.refs import parse_range
                        area = parse_range(written.range)
                        styles = {col: copy(sheet.cell(block.copy_style_from_row, col)._style) for col in range(area.min_col, area.max_col + 1)}
                        for row in sheet.iter_rows(min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col):
                            for cell in row:
                                cell._style = copy(styles[cell.column])
                    elif source is None and block.start_cell == 'A1':
                        for cell in sheet[1]:
                            cell.font = Font(bold=True)
                        sheet.freeze_panes = 'A2'
                buffer = io.BytesIO()
                workbook.save(buffer)
                payload = buffer.getvalue()
                if len(payload) > MAX_FILE:
                    raise ValueError('成果超过 25 MB。')
                # Link publishes atomically without overwriting an existing name.
                # A directory fd pins the authorized parent against path swaps.
                parent = self._path(str(target), new=True).parent
                dir_fd = os.open(self.root, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
                try:
                    for component in parent.relative_to(self.root).parts:
                        child_fd = os.open(component, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=dir_fd)
                        os.close(dir_fd)
                        dir_fd = child_fd
                    import uuid
                    temporary = '.excel-' + uuid.uuid4().hex + '.tmp'
                    fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600, dir_fd=dir_fd)
                    try:
                        with os.fdopen(fd, 'wb') as handle:
                            handle.write(payload)
                            handle.flush()
                            os.fsync(handle.fileno())
                        os.link(temporary, target.name, src_dir_fd=dir_fd, dst_dir_fd=dir_fd, follow_symlinks=False)
                    finally:
                        os.unlink(temporary, dir_fd=dir_fd)
                finally:
                    os.close(dir_fd)
                return {'created_file': str(target), 'bytes': len(payload), 'source_preserved': True, 'notice': NOTICE, 'delivery': '成果已保存在选定的表格文件夹。对话选择同一文件夹时，可从右侧文件列表用默认应用打开。'}
            finally:
                workbook.close()
