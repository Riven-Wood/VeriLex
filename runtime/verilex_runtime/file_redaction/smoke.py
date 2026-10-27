"""Synthetic packaged acceptance: real NER, Word replacement and no network."""
def run():
    import io
    import json
    import socket
    from docx import Document
    from .pipeline import prepare,render,restore
    from .ner import detect
    from .rules import DEFAULT_KINDS
    def forbidden(*args,**kwargs): raise AssertionError('redaction attempted network access')
    socket.create_connection=forbidden
    socket.socket.connect=forbidden
    d=Document();d.add_paragraph('张三与李四签署合同。联系地址：北京市海淀区中关村大街1号。电话13800138000。')
    d.add_paragraph('收款账号：123456789012；备用电话：１３８ ００１３ ８０００。');
    d.add_paragraph('张三再次签字。');d.core_properties.author='虚构作者'
    output=io.BytesIO();d.save(output);source=output.getvalue()
    job=prepare(source,list(DEFAULT_KINDS),[],detect)
    assert any(m['old']=='张三' and m['kind']=='person' for m in job['mapping'])
    assert any(m['old']=='13800138000' for m in job['mapping'])
    assert any(m['old']=='123456789012' and m['kind']=='bank' for m in job['mapping'])
    assert any(m['old']=='１３８ ００１３ ８０００' and m['kind']=='phone' for m in job['mapping'])
    result=render(job,'permanent',detect)
    text='\n'.join(p.text for p in Document(io.BytesIO(result)).paragraphs)
    assert '13800138000' not in text and '张三' not in text
    assert '123456789012' not in text and '１３８ ００１３ ８０００' not in text
    assert not Document(io.BytesIO(result)).core_properties.author
    assert Document(io.BytesIO(restore(job))).paragraphs[0].text==d.paragraphs[0].text
    print(json.dumps({'ok':True,'local_model':True,'network_disabled':True,'occurrences':len(job['mapping'])}))
