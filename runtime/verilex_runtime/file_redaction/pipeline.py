"""Occurrence-based processing built on the existing checked DOCX slice.

Canonical storage contains placeholders only. Original coordinates and canonical
coordinates are distinct, so restoring one duplicate never restores its peers.
"""
import copy
import hashlib
import io
from collections import defaultdict
from docx import Document
from lxml import etree
from verilex_runtime.documents.docx_package import W, checked_bytes, parts, pack, story_parts, xml, serialize
from verilex_runtime.documents.replace import _groups, _set_text
from verilex_runtime.documents.scrub import scrub_package
from verilex_runtime.documents.scrub_verify import audit_processed_copy
from .errors import RedactionError,import_error
from verilex_runtime.attachments.errors import AttachmentDomainError
from .rules import LABELS, merge_hits


def inventory(data):
    items=parts(data); roots={}; groups={}; covered=set()
    forbidden={f'{{{W}}}{n}' for n in ('txbxContent','object','altChunk')}
    for name in story_parts(items):
        root=xml(items[name]); roots[name]=root
        if any(e.tag in forbidden for e in root.iter()): raise RedactionError('redaction_complex')
        for pindex,p in enumerate(root.iter(f'{{{W}}}p')):
            for gindex,nodes in enumerate(_groups(p)):
                key=f'{name}:{pindex}:{gindex}'
                groups[key]=nodes
                covered.update(nodes)
    if any(n not in covered for root in roots.values() for n in root.iter(f'{{{W}}}t')):
        raise RedactionError('redaction_complex')
    for name,data in items.items():
        if name.startswith(('word/charts/','word/diagrams/','word/embeddings/')): raise RedactionError('redaction_complex')
        if name.endswith('.rels') and any(r.get('Type','').rsplit('/',1)[-1] in {'chart','diagramData','oleObject','package'} for r in xml(data)):
            raise RedactionError('redaction_complex')
    if sum(len(''.join(n.text or '' for n in ns)) for ns in groups.values())>200000: raise RedactionError('redaction_limit')
    return items,roots,groups


def location_order(key):
    part,paragraph,group=key.rsplit(':',2)
    return part,int(paragraph),int(group)


def review_mapping(job):
    """View-only context from canonical text; never add source snippets to storage."""
    _,_,groups=inventory(job['canonical'])
    texts={key:text_of(nodes) for key,nodes in groups.items()}
    result=[]
    for m in sorted(job['mapping'],key=lambda m:(location_order(m['key']),m['start'])):
        text=texts[m['key']]; lo=max(0,m['start']-36); hi=min(len(text),m['end']+36)
        result.append({**m,
            'context_before':('…' if lo else '')+text[lo:m['start']],
            'context_after':text[m['end']:hi]+('…' if hi<len(text) else ''),
        })
    return result


def text_of(nodes): return ''.join(n.text or '' for n in nodes)


def edit(data, changes):
    items,roots,groups=inventory(data)
    expected={key:text_of(nodes) for key,nodes in groups.items()}
    for key,spans in changes.items():
        if key not in groups: raise RedactionError('restore_unavailable')
        nodes=groups[key]; original=text_of(nodes)
        spans=sorted(spans,key=lambda s:s[0])
        previous=-1
        for lo,hi,old,new in spans:
            if not 0<=lo<hi<=len(original) or lo<previous or original[lo:hi]!=old:
                raise RedactionError('restore_unavailable')
            previous=hi
        offsets=[]; offset=0
        for n in nodes:
            offsets.append((offset,offset+len(n.text or ''),n)); offset=offsets[-1][1]
        for lo,hi,old,new in reversed(spans):
            expected[key]=expected[key][:lo]+new+expected[key][hi:]
            active=[(a,b,n) for a,b,n in offsets if a<hi and b>lo]
            for i,(a,b,n) in enumerate(active):
                value=n.text or ''
                _set_text(n,value[:max(0,lo-a)]+(new if i==0 else '')+(value[hi-a:] if hi<b else ''))
    for name,root in roots.items(): items[name]=serialize(root)
    output=pack(items); checked_bytes(output)
    # Independent paragraph audit detects wrong-position and count-preserving swaps.
    from .verify import verify_structure
    verify_structure(data,output,expected)
    if not audit_processed_copy(output,[],scrubbed=True).passed: raise RedactionError('redaction_residue')
    return output


def prepare(raw,kinds,custom,detector,previous=None):
    try: clean=scrub_package(checked_bytes(raw)).copy_bytes
    except AttachmentDomainError as e: raise import_error(e.code) from None
    except Exception: raise RedactionError('redaction_invalid') from None
    items,roots,groups=inventory(clean)
    texts={key:text_of(nodes) for key,nodes in groups.items()}
    all_text='\n'.join(texts.values())
    discovered={key:detector(t,kinds,custom) for key,t in texts.items()}
    # Entity recognition is occurrence-specific. Equal spellings in another
    # context do not inherit a decision; explicit custom terms are matched by
    # their detector at every requested occurrence.
    if sum(len(hits) for hits in discovered.values())>5000: raise RedactionError('redaction_limit')
    discovered={key:merge_hits(hits) for key,hits in discovered.items()}
    mapping=[]; numbers=defaultdict(int); placeholders={}; prior={}
    if previous:
        prior={(m['key'],m['source_start'],m['old'],m['kind']):m for m in previous}
        for m in previous: placeholders[(m['kind'],m['old'])]=m['placeholder']
    used=set(placeholders.values())
    for key in sorted(texts,key=location_order):
        delta=0
        for hit in discovered[key]:
            entity=(hit['kind'],hit['old'])
            if entity not in placeholders:
                while True:
                    numbers[hit['kind']]+=1
                    candidate=f'【{LABELS[hit["kind"]]}{numbers[hit["kind"]]}】'
                    if candidate not in all_text and candidate not in used: break
                placeholders[entity]=candidate; used.add(candidate)
            placeholder=placeholders[entity]
            old_entry=prior.get((key,hit['start'],hit['old'],hit['kind']))
            mapping.append(dict(id=hashlib.sha256(f'{key}:{hit["start"]}:{hit["kind"]}:{hit["old"]}'.encode()).hexdigest()[:24],
                key=key,source_start=hit['start'],source_end=hit['end'],start=hit['start']+delta,
                end=hit['start']+delta+len(placeholder),old=hit['old'],kind=hit['kind'],placeholder=placeholder,
                keep=bool(old_entry and old_entry['keep'])))
            delta+=len(placeholder)-len(hit['old'])
    if len(mapping)>5000: raise RedactionError('redaction_limit')
    changes=defaultdict(list)
    for m in mapping: changes[m['key']].append((m['source_start'],m['source_end'],m['old'],m['placeholder']))
    canonical=edit(clean,changes)
    result=dict(canonical=canonical,digest=hashlib.sha256(canonical).hexdigest(),mapping=mapping,kinds=kinds,custom=custom)
    from .verify import verify_residue
    verify_residue(canonical, mapping, None, detector, kinds, custom)
    return result


def validate(job):
    if job.get('mapping') is None or hashlib.sha256(job['canonical']).hexdigest()!=job['digest']:
        raise RedactionError('restore_unavailable')
    try:
        _,_,groups=inventory(job['canonical'])
        seen=set(); last={}
        for m in job['mapping']:
            t=text_of(groups[m['key']])
            if (m['id'] in seen or not 0<=m['start']<m['end']<=len(t)
                or m['start']<last.get(m['key'],0) or t[m['start']:m['end']]!=m['placeholder']):
                raise RedactionError('restore_unavailable')
            seen.add(m['id']); last[m['key']]=m['end']
    except (KeyError,TypeError,ValueError): raise RedactionError('restore_unavailable') from None


def restore(job):
    validate(job); changes=defaultdict(list)
    for m in job['mapping']: changes[m['key']].append((m['start'],m['end'],m['placeholder'],m['old']))
    return edit(job['canonical'],changes)


def render(job,variant,detector):
    validate(job)
    if variant=='restore': return restore(job)
    if variant not in {'placeholder','permanent','markdown'}: raise RedactionError()
    changes=defaultdict(list); keep_ranges=defaultdict(list); delta=defaultdict(int)
    # Both generated placeholders and explicitly retained positions are masked in
    # the second recognition pass; this is a position whitelist, never a word list.
    masks=defaultdict(list)
    for m in job['mapping']:
        value=m['old'] if m['keep'] else ('' if variant=='permanent' else m['placeholder'])
        start=m['start']+delta[m['key']]
        if value: masks[m['key']].append((start,start+len(value)))
        changes[m['key']].append((m['start'],m['end'],m['placeholder'],value))
        delta[m['key']]+=len(value)-len(m['placeholder'])
    output=edit(job['canonical'],changes)
    from .verify import verify_residue
    verify_residue(output,job['mapping'],masks,detector,job['kinds'],job['custom'])
    if variant=='markdown':
        _,_,groups=inventory(output)
        # Explicit text export: preserve ordering and separate each source paragraph.
        return ('\n\n'.join(text_of(nodes) for key,nodes in sorted(groups.items(),key=lambda pair:location_order(pair[0])))+'\n').encode()
    return output
