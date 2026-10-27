"""Only synthetic data; repeatable Word redaction demonstration."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from verilex_runtime.file_redaction.pipeline import prepare,render,restore
from verilex_runtime.file_redaction.ner import detect
from verilex_runtime.file_redaction.rules import DEFAULT_KINDS
import hashlib,json

root=Path('output/file-redaction');root.mkdir(parents=True,exist_ok=True)
d=Document();d.add_heading('虚构合同 · 脱敏验收样本',0)
p=d.add_paragraph();p.add_run('张').bold=True;p.add_run('三').bold=True;p.add_run('与李四签署本合同。')
d.add_paragraph('张三再次签字。')
table=d.add_table(rows=2,cols=2);table.style='Table Grid';table.cell(0,0).text='联系人';table.cell(0,1).text='张三';table.cell(1,0).text='联系地址';table.cell(1,1).text='联系地址：北京市海淀区中关村大街1号。'
d.add_paragraph('电话13800138000，邮箱demo@example.com。')
d.sections[0].header.paragraphs[0].text='联系电话13800138000'
d.sections[0].footer.paragraphs[0].text='仅用于功能检查，内容全部虚构'
d.core_properties.author='虚构作者';d.core_properties.last_modified_by='虚构编辑者'
source=root/'虚构合同.docx';d.save(source);raw=source.read_bytes()
job=prepare(raw,list(DEFAULT_KINDS),[],detect)
(root/'脱敏版.docx').write_bytes(render(job,'placeholder',detect))
first=next(m for m in job['mapping'] if m['old']=='张三');first['keep']=True
(root/'仅保留一处姓名.docx').write_bytes(render(job,'permanent',detect))
(root/'文字还原版.docx').write_bytes(restore(job))
(root/'脱敏版.md').write_bytes(render(job,'markdown',detect))
assert source.read_bytes()==raw
(root/'检查记录.json').write_text(json.dumps({'source_sha256':hashlib.sha256(raw).hexdigest(),'occurrences':len(job['mapping']),'original_unchanged':True},ensure_ascii=False,indent=2))
print('Synthetic sample and verified outputs:',root)
