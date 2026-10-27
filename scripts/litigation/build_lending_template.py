"""Adapt the official 2025 blank form. Only source text/OOXML, no upstream code.

Run after build_templates.py. Input is the reviewed, hash-pinned source extract;
no network or LibreOffice is needed when rebuilding application assets.
"""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parents[1] / 'runtime/verilex_runtime/litigation/templates'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
IDENT = 'private-lending-person-2025'


def build_variant(ident=IDENT, plaintiff_type='natural', defendant_type='natural', title='自然人', third_party_type=None):
    provenance = json.loads((ROOT/'sources/private-lending-2025.json').read_text())
    source = (ROOT/'sources/private-lending-2025.xml').read_bytes()
    assert hashlib.sha256(source).hexdigest() == provenance['extracted_sha256']
    original = etree.fromstring(source)
    document = Document()
    body = document.element.body
    for child in list(body): body.remove(child)
    for child in original: body.append(deepcopy(child))
    # Keep the official cell grids, borders, widths and labels. Flowing tables
    # replace the source's page fragments; rows may grow and cross pages.
    for tag in ('sectPr', 'bookmarkStart', 'bookmarkEnd', 'numPr', 'cantSplit'):
        for node in list(body.findall('.//w:'+tag, NS)): node.getparent().remove(node)
    # The converted source's paragraph style references make LibreOffice
    # flatten tables when placed in a fresh package. Direct formatting is
    # retained; inheriting the fresh package's default paragraph style works.
    for style in list(body.findall('.//w:pStyle', NS)):
        style.getparent().remove(style)
    for height in body.findall('.//w:trHeight', NS): height.set(qn('w:hRule'), 'atLeast')
    # Blank source paragraphs positioned labels on pre-paginated sheets.
    # Remove them inside non-empty cells so labels stay beside flowing values.
    for target in body.findall('.//w:tc', NS):
        paras = target.findall('w:p', NS)
        if any(''.join(p.itertext()).strip() for p in paras):
            for p in paras:
                if not any((t.text or '').strip() for t in p.findall('.//w:t', NS)):
                    target.remove(p)
    tables = list(body.findall('w:tbl', NS))
    fields = []

    def f(key, label, hint='', kind='text', choices=None):
        rule = dict(key=key,label=label,hint=hint,required=True,kind=kind)
        if choices: rule['choices'] = choices
        fields.append(rule)
        return '{{'+key+'}}'

    def cell(t, r, c, text):
        target=tables[t].findall('w:tr',NS)[r].findall('w:tc',NS)[c]
        paras=target.findall('w:p',NS)
        sample=next((p for p in paras if p.findall('.//w:t',NS)),paras[0])
        ppr=sample.find('w:pPr',NS)
        rpr=sample.find('.//w:rPr',NS)
        for p in paras:target.remove(p)
        p=etree.SubElement(target,qn('w:p'))
        if ppr is not None:
            props=deepcopy(ppr)
            # Source offsets position individual blank runs on fixed pages.
            # A filled answer must use the cell width, including long answers.
            for indent in list(props.findall('w:ind',NS)):props.remove(indent)
            p.append(props)
        run=etree.SubElement(p,qn('w:r'))
        if rpr is not None:run.append(deepcopy(rpr))
        etree.SubElement(run,qn('w:t')).text=text

    def natural(prefix, title):
        return '\n'.join(label+'：'+f(prefix+key,title+label,kind=kind) for key,label,kind in [
            ('name','姓名','text'),('sex','性别','sex'),('birth','出生日期','date'),
            ('ethnicity','民族','text'),('work','工作单位','text'),('job','职务','text'),
            ('phone','联系电话','text'),('address','住所地（户籍所在地）','text'),
            ('residence','经常居住地','text'),('id_type','证件类型','text'),('id_number','证件号码','text')])

    # Mark inapplicable blocks first; populate only the selected party rows.
    for t,rows in [(0,[2,3,4,5]),(1,[1,2,3,4,5]),(2,[0,1,2])]:
        for r in rows:
            for c in range(1,len(tables[t].findall('w:tr',NS)[r].findall('w:tc',NS))):cell(t,r,c,'无（不适用）')

    def party(prefix, title, kind, t, natural_row, organization_row, organization_table=None):
        if kind == 'natural':
            cell(t,natural_row,1,natural(prefix,title))
            return
        if organization_table is not None:
            t = organization_table
        def entry(key, label, hint=''):
            return label+'：'+f(prefix+key,title+label,hint)
        cell(t,organization_row,1,'\n'.join([
            entry('name','名称'),entry('address','住所地（主要办事机构所在地）')]))
        left=[entry('registration','注册地 / 登记地'),entry('representative','法定代表人 / 负责人')]
        if prefix == 'plaintiff_':
            left.append(entry('credit_code','统一社会信用代码'))
        left.append(entry('organization_type','类型','按实际填写：有限责任公司、股份有限公司（是否上市）、其他企业法人、事业单位、社会团体、基金会、社会服务机构、机关法人、农村集体经济组织法人、城镇农村的合作经济组织法人、基层群众性自治组织法人、个人独资企业、合伙企业、不具有法人资格的专业服务机构等'))
        cell(t,organization_row+1,1,'\n'.join(left))
        cell(t,organization_row+1,2,'\n'.join([entry('job','职务'),entry('phone','联系电话')]))
        cell(t,organization_row+2,1,entry('ownership','所有制性质','国有（注明控股或参股）、民营；其他请如实说明'))

    party('plaintiff_','原告',plaintiff_type,0,2,3)
    party('defendant_','被告',defendant_type,1,1,2)
    if third_party_type:
        party('third_party_','第三人',third_party_type,1,5,0,organization_table=2)
    cell(1,0,1,f('agent','委托诉讼代理人','未委托填“无”；有代理人时填写姓名、单位、职务、电话及一般/特别授权权限'))
    cell(2,4,0,'（可完整表述诉讼请求；为方便、准确梳理要点，相关内容请在下方要素式表格中填写）\n'+f('claims','诉讼请求完整表述'))
    for row,key,label,hint in [
        (5,'principal','本金请求','截至日期、尚欠本金金额及币种'),
        (6,'interest','利息请求','截至日期、尚欠利息金额、计算方式'),
        (7,'acceleration','提前还款或解除合同','是/否；如是，明确提前还款或解除合同及内容'),
        (8,'security_claim','担保权利请求','是/否；如是，填写具体内容'),
        (9,'recovery_cost','实现债权的费用','是/否；如是，填写费用明细'),
        (10,'litigation_cost','诉讼费用请求','明确是或否'),
        (11,'other_claims','其他请求','无其他请求填“无”'),
        (12,'total','标的总额','金额及币种'),
        (14,'jurisdiction','仲裁或法院管辖约定','有/无；如有，填写合同条款及内容'),
        (15,'preservation','诉前保全情况','是/否；如是，填写法院、时间、案号；诉讼保全须另交申请材料')]:
        cell(2,row,1,f(key,label,hint))
    cell(2,6,2,'是否请求支付至实际清偿之日止：'+f('interest_until_paid','利息是否请求至实际清偿日',kind='choice',choices=['是','否']))
    cell(3,1,0,'（可完整表述纠纷涉及的事实与理由；为方便、准确梳理要点，相关内容请在下方要素式表格中填写）\n'+f('reasons','事实与理由完整表述'))
    facts=[
        ('contract','合同签订情况','名称、编号、签订时间、地点等'),
        ('signatories','签订主体','出借人、借款人'),
        ('loan_amount','借款金额与提供方式','约定金额、实际提供金额；现金/转账/其他及说明'),
        ('loan_term','借款期限','是否到期；约定起止日期'),
        ('loan_rate','借款利率','利率、年/季/月及合同条款；无约定如实写明'),
        ('delivery','借款提供时间','逐笔日期和金额'),
        ('repayment_method','还款方式','一次还本付息/按月、季或年计息到期还本/其他及说明'),
        ('repayment','还款情况','已还本金、已还利息、还息截至日期'),
        ('overdue','逾期还款','是/否；如是，填写逾期起始时间及期间'),
        ('collateral_contract','抵押或质押合同','是/否；如是，填写签订时间'),
        ('collateral','担保人和担保物','无则填“无”'),
        ('maximum_security','最高额担保','是/否；如是，填写债权确定时间、担保额度'),
        ('registration','抵押或质押登记','是/否；正式登记/预告登记及情况'),
        ('guarantee','保证合同','是/否；签订时间、保证人、主要内容、一般保证/连带责任保证'),
        ('other_security','其他担保方式','是/否；形式、签订时间'),
        ('other_facts','其他需要说明的内容','无则填“无”'),
        ('basis','请求依据','合同约定、法律规定；编号法条须先在本案讨论核验'),
        ('evidence','证据清单','借款合同、借据、交付和还款凭证、担保资料等；按实际材料列明')]
    for i,(key,label,hint) in enumerate(facts):
        cell(3 if i<12 else 4,i+2 if i<12 else i-12,1,f(key,label,hint))
    cell(4,7,1,f('mediation_understood','是否了解调解的作用',kind='choice',choices=['了解','不了解']))
    benefits=tables[4].findall('w:tr',NS)[8].findall('w:tc',NS)[1]
    text=''.join(node.text or '' for node in benefits.findall('.//w:t',NS))
    labels=['先行调解程序与排期','调解诉讼费用','更换调解组织和调解员','调解保密性','调解协议法律效力']
    for i,label in enumerate(labels):
        text=text.replace('了解□    不了解□',f('mediation_'+chr(97+i),'是否了解'+label,kind='choice',choices=['了解','不了解']),1)
    assert '了解□' not in text
    cell(4,8,1,text)
    cell(5,0,1,f('mediation_choice','是否考虑先行调解','由当事人选择；模型不得代选',kind='choice',choices=['是','否','暂不确定，想要了解更多内容']))
    # Keep an unsigned signature line; the date is reviewed, never inferred.
    body.remove(body[-1])
    document.add_paragraph('具状人（签字、盖章）：\n日期：'+f('date','落款日期',kind='date'))
    document.add_paragraph('AI 草稿，须律师审查')
    heading=document.add_paragraph('核验状态汇总')
    heading.paragraph_format.page_break_before=True
    heading.paragraph_format.keep_with_next=True
    document.add_paragraph('{{verification_summary}}')
    section=deepcopy(original[-1].find('.//w:sectPr',NS))
    for tag in ('headerReference','footerReference','type'):
        for node in list(section.findall('w:'+tag,NS)):section.remove(node)
    body.append(section)
    # No external relationships or original document metadata enter the package.
    assert not body.xpath('.//@r:id | .//@r:embed')
    content=etree.tostring(document.element,encoding='UTF-8',xml_declaration=True,pretty_print=True)
    (OUT/(ident+'.xml')).write_bytes(content)
    manifest_path=OUT/'manifest.json'
    manifest=json.loads(manifest_path.read_text())
    manifest['templates']=[t for t in manifest['templates'] if t['id']!=ident]
    manifest['templates'].append(dict(id=ident,name=f'民间借贷起诉状（2025要素式·{title}）',version='2025-07-14/verilex-1',
        source_url=provenance['publication_url'],retrieved_at=provenance['retrieved_at'],rights=provenance['rights'],
        layout=f'2025官方空白表格独立适配；保留栏目、表格网格及边框，填充区扩容；{title}单原告单被告、'+('单第三人' if third_party_type else '无第三人')+'，非原件像素复刻',
        party_types=list(dict.fromkeys([plaintiff_type,defendant_type])),third_parties=bool(third_party_type),sha256=hashlib.sha256(content).hexdigest(),fields=fields))
    if ident != IDENT:
        manifest['templates'][-1]['party_scope'] = {'plaintiff_type':plaintiff_type,'defendant_type':defendant_type}
    if third_party_type:
        manifest['templates'][-1]['party_scope'].update(third_party_type=third_party_type,third_party_count=1)
    manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(f'{ident}: {len(fields)} fields')

def build():
    build_variant()
    build_variant('private-lending-person-organization-2025','natural','organization','个人起诉组织')
    build_variant('private-lending-organization-person-2025','organization','natural','组织起诉个人')
    build_variant('private-lending-organization-2025','organization','organization','组织起诉组织')
    for base, plaintiff, defendant, title in [
        ('person','natural','natural','个人双方'),
        ('person-organization','natural','organization','个人起诉组织'),
        ('organization-person','organization','natural','组织起诉个人'),
        ('organization','organization','organization','组织双方'),
    ]:
        for third, label in [('natural','个人第三人'),('organization','组织第三人')]:
            build_variant(f'private-lending-{base}-third-{third}-2025',plaintiff,defendant,title+'·'+label,third)



if __name__=='__main__':build()
