"""Rebuild independently typeset ordinary court forms; review XML/manifest diff."""
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm
from docx.oxml.ns import qn
from lxml import etree

OUT = Path(__file__).resolve().parents[2] / 'runtime/verilex_runtime/litigation/templates'
OUT.mkdir(exist_ok=True)


def field(key, label, hint='', required=True, kind='text'):
    return dict(key=key, label=label, hint=hint, required=required, kind=kind)

natural = [field('plaintiff_name','原告姓名'),field('sex','性别',kind='sex'),field('birth_date','出生日期',kind='date'),field('ethnicity','民族'),field('occupation','工作单位和职务或职业'),field('address','原告住所'),field('contact','原告联系方式')]
company = [field('plaintiff_name','原告名称'),field('address','原告住所'),field('representative','法定代表人或主要负责人','姓名、职务、联系方式')]
common = [field('legal_agent','法定或指定代理人','不适用时填“无”'),field('agent','委托诉讼代理人','姓名及基本信息；未委托填“无”'),field('defendant','被告基本信息','自然人姓名、性别、工作单位、住所等；法人名称、住所等'),field('claims','诉讼请求'),field('reasons','事实和理由'),field('evidence','证据和证据来源，证人姓名和住所','无证人时如实说明'),field('court','受诉法院','法院完整名称'),field('copies','起诉状副本份数',kind='count'),field('date','落款日期',kind='date')]
forms = [
 ('civil-person','普通民事起诉状（自然人）',181,natural+common, [
  '民事起诉状','原告：{{plaintiff_name}}，{{sex}}，{{birth_date}}生，{{ethnicity}}，{{occupation}}，住{{address}}。联系方式：{{contact}}。',
  '法定代理人/指定代理人：{{legal_agent}}。','委托诉讼代理人：{{agent}}。','被告：{{defendant}}。','诉讼请求：','{{claims}}','事实和理由：','{{reasons}}','证据和证据来源，证人姓名和住所：','{{evidence}}','此致','{{court}}','附：本起诉状副本{{copies}}份','起诉人(签名)','{{date}}']),
 ('civil-organization','普通民事起诉状（法人或其他组织）',187,company+[f for f in common if f['key']!='legal_agent'],[
  '民事起诉状','原告：{{plaintiff_name}}，住所{{address}}。','法定代表人/主要负责人：{{representative}}。','委托诉讼代理人：{{agent}}。','被告：{{defendant}}。','诉讼请求：','{{claims}}','事实和理由：','{{reasons}}','证据和证据来源，证人姓名和住所：','{{evidence}}','此致','{{court}}','附：本起诉状副本{{copies}}份','起诉人(公章和签名)','{{date}}']),
 ('property-preservation','诉讼财产保全申请书',126,[field('plaintiff_name','申请人姓名或名称'),field('applicant','申请人其他基本信息','自然人性别、出生日期、民族、工作单位/职业、住所、联系方式；法人住所及法定代表人姓名、职务、联系方式'),field('legal_agent','法定或指定代理人','不适用填“无”'),field('agent','委托诉讼代理人','未委托填“无”'),field('defendant','被申请人基本信息'),field('measure','保全措施',kind='measure'),field('property','保全财产','名称、性质、数量、数额、所在地等'),field('term','保全期限'),field('case_number','案号及当事人、案由'),field('reasons','申请保全的事实和理由'),field('guarantee','担保情况','担保财产名称、性质、数量、数额、所在地等；未提供担保如实写明'),field('court','受诉法院'),field('date','落款日期',kind='date')],[
  '申请书','申请人：{{plaintiff_name}}，{{applicant}}。','法定代理人/指定代理人：{{legal_agent}}。','委托诉讼代理人：{{agent}}。','被申请人：{{defendant}}。','请求事项：','{{measure}}被申请人的{{property}}，期限为{{term}}。','事实和理由：','{{case_number}}一案，{{reasons}}','担保情况：{{guarantee}}','此致','{{court}}','申请人(签名或者盖章)','{{date}}']),
]
manifest=[]
for ident,name,source,fields,paragraphs in forms:
    document=Document()
    section=document.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=Mm(25.4)
    section.left_margin=section.right_margin=Mm(28)
    for index,text in enumerate(paragraphs+['AI 草稿，须律师审查']):
        p=document.add_paragraph()
        p.paragraph_format.line_spacing=1.5
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.widow_control=True
        if index==0:
            p.alignment=1
            p.paragraph_format.keep_with_next=True
        r=p.add_run(text);r.font.name='宋体';r.font.size=Pt(20 if index==0 else 12 if text=='AI 草稿，须律师审查' else 14)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        r.bold=index==0
    summary_heading=document.add_paragraph('核验状态汇总')
    summary_heading.paragraph_format.page_break_before=True
    summary_heading.paragraph_format.keep_with_next=True
    document.add_paragraph('{{verification_summary}}')
    content=etree.tostring(document._element,xml_declaration=True,encoding='UTF-8',pretty_print=True)
    (OUT/(ident+'.xml')).write_bytes(content)
    manifest.append(dict(id=ident,name=name,version='2016-09-28/verilex-1',source_url=f'https://www.court.gov.cn/susongyangshi/xiangqing/{source}.html',retrieved_at='2026-09-09',rights='最高人民法院官方司法文书样式正文；著作权法第五条',layout='官方栏目及顺序；VeriLex独立印刷版式，非官方DOCX原件',sha256=hashlib.sha256(content).hexdigest(),fields=fields))
(OUT/'manifest.json').write_text(json.dumps({'schema_version':1,'templates':manifest},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

# The reviewed 2025 source form is part of the same deterministic rebuild.
from build_lending_template import build
build()
