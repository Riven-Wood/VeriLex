"""Generate synthetic DOCX slice acceptance samples. No model, network or user files."""
from __future__ import annotations
import hashlib
import io
import json
from pathlib import Path
import zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from lxml import etree
from verilex_runtime.documents.comments import CommentSpec, write_verified_comments
from verilex_runtime.documents.comment_verify import verify_comment_copy
from verilex_runtime.documents.replace import TextReplacement, replace_texts, process_copy
from verilex_runtime.documents.scrub_verify import audit_processed_copy


def generate(folder: Path):
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    normal = doc.styles['Normal']; normal.font.name = '宋体'; normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')
    doc.add_heading('文档能力验证样本', 0)
    doc.add_paragraph('本文件全部内容为虚构测试数据，仅用于检查 Word 文件处理效果。')
    doc.add_paragraph('甲方：张三；乙方：示例服务中心。')
    para = doc.add_paragraph('联系电话：'); para.add_run('1380').bold = True; para.add_run('0138000').italic = True
    doc.add_paragraph('测试身份编号：TEST-ID-0001（非真实证件号码）。')
    doc.add_heading('第一条 付款安排', 1)
    doc.add_paragraph('甲方应于收到付款通知后三十日内支付服务费用。')
    doc.add_heading('第二条 资料交接', 1)
    doc.add_paragraph('乙方应于服务完成后五日内交付全部资料。')
    table = doc.add_table(rows=2, cols=2); table.style = 'Table Grid'
    table.cell(0,0).text='项目'; table.cell(0,1).text='测试内容'
    table.cell(1,0).text='联系人'; table.cell(1,1).text='张三'
    doc.sections[0].header.paragraphs[0].text='测试合同｜联系电话：13800138000'
    doc.sections[0].footer.paragraphs[0].text='VeriLex 文档切片 · 虚构样本'
    doc.core_properties.author='测试作者'; doc.core_properties.last_modified_by='测试修改者'
    original=folder/'01-original.docx'; doc.save(original)
    digest=hashlib.sha256(original.read_bytes()).hexdigest()
    annotated=write_verified_comments(original,[CommentSpec('三十日内','测试批注一：核对付款期限。','VeriLex','VL'),CommentSpec('五日内','测试批注二：核对交付期限。','VeriLex','VL'),CommentSpec('TEST-ID-0001','测试批注三：此处为虚构编号。','VeriLex','VL')])
    (folder/'02-commented.docx').write_bytes(annotated.copy_bytes)
    # A second source contains tracked insertion/deletion and an existing comment.
    red_doc=Document(io.BytesIO(annotated.copy_bytes)); p=red_doc.add_paragraph('修订示例：')
    for tag, value in [('del','旧的交付说明'),('ins','新的交付说明')]:
        container=OxmlElement('w:'+tag);container.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author','测试修订者')
        run=OxmlElement('w:r');t=OxmlElement('w:delText' if tag=='del' else 'w:t');t.text=value;run.append(t);container.append(run);p._p.append(container)
    red_source=folder/'03-with-revisions.docx';red_doc.save(red_source)
    replacements=[TextReplacement('张三','【姓名1】'),TextReplacement('13800138000','【电话1】'),TextReplacement('TEST-ID-0001','【编号1】')]
    result=process_copy(red_source,replacements);(folder/'04-replaced-clean.docx').write_bytes(result.copy_bytes)
    def tamper(data, old, new):
        out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as z, zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as target:
            for n in z.namelist():
                content=z.read(n)
                if n=='word/document.xml': content=content.replace(old.encode(),new.encode())
                target.writestr(n,content)
        return out.getvalue()
    bad_comment=tamper(annotated.copy_bytes,'三十日内','九十日内');(folder/'05-tampered-comment.docx').write_bytes(bad_comment)
    bad_replaced=tamper(result.copy_bytes,'【姓名1】','张三');(folder/'06-tampered-replacement.docx').write_bytes(bad_replaced)
    report={'original_unchanged':hashlib.sha256(original.read_bytes()).hexdigest()==digest,
            'comments_verified':verify_comment_copy(original,annotated.copy_bytes).passed,
            'comment_tamper_rejected':not verify_comment_copy(original,bad_comment).passed,
            'replacement_verified':result.audit.passed,'replacement_counts':result.counts,
            'residue_tamper_rejected':not audit_processed_copy(bad_replaced,replacements,scrubbed=True).passed,
            'uncovered':result.uncovered_kinds}
    (folder/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':
    import sys
    generate(Path(sys.argv[1]))
