"""Regenerate the wholly fictional DOC compatibility fixture on macOS.

Run with PYTHONPATH=runtime runtime/.venv/bin/python <this file>.
Uses the installed LibreOffice inside the same filesystem/network sandbox as
the importer. The DOC binary is a fixture, never a user document.
"""
from pathlib import Path
import io
import struct
import subprocess
import tempfile
import zlib

from docx import Document
from docx.shared import Inches
from verilex_runtime.attachments.legacy_doc import _sandbox_policy


def main():
    destination = Path(__file__).resolve().parents[2] / "runtime/tests/fixtures/legacy-word.doc"
    with tempfile.TemporaryDirectory(prefix="vlf-", dir="/private/tmp") as temp:
        root = Path(temp).resolve()
        doc = Document()
        p = doc.add_paragraph("旧版 Word 合同兼容性样本（全部虚构）")
        p.runs[0].bold = True
        doc.add_comment(p.runs, "原有批注：请核对合同名称。", "原作者")
        doc.add_paragraph("甲方：虚构甲公司；乙方：虚构乙公司。")
        doc.add_paragraph("第一条：验收合格后付款，未约定验收期限。")
        doc.add_paragraph("第二条：乙方逾期交付，每日承担千分之三违约金。")
        doc.add_paragraph("联系人：张三（虚构）。")
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [("付款阶段", "金额"), ("验收后付款", "10000元")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.sections[0].header.paragraphs[0].text = "虚构合同 · 格式兼容检查"
        doc.sections[0].footer.paragraphs[0].text = "仅用于软件测试"
        def chunk(kind, data):
            return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
        png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
        png += chunk(b"IDAT", zlib.compress((b"\0" + b"\x20\x70\xc0" * 4) * 4)) + chunk(b"IEND", b"")
        doc.add_picture(io.BytesIO(png), width=Inches(.4))
        source = root / "legacy-word.docx"
        doc.save(source)
        profile = root / "profile"
        (profile / "user").mkdir(parents=True)
        (profile / "user/registrymodifications.xcu").write_text('''<?xml version="1.0"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>''')
        sandbox = root / "fixture.sb"
        sandbox.write_text(_sandbox_policy(root))
        subprocess.run([
            "/usr/bin/sandbox-exec", "-f", str(sandbox),
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
            f"-env:UserInstallation={profile.as_uri()}", f"-env:OSL_SOCKET_PATH={root}",
            "--headless", "--nologo", "--nodefault", "--norestore",
            "--convert-to", "doc:MS Word 97", "--outdir", str(root), str(source),
        ], cwd=root, env={"PATH": "/usr/bin:/bin", "HOME": str(root), "TMPDIR": str(root) + "/"},
            check=True, timeout=45, stdout=subprocess.DEVNULL)
        result = (root / "legacy-word.doc").read_bytes()
        assert result.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(result)
        print(f"Generated fictional fixture: {destination.name} ({len(result)} bytes)")


if __name__ == "__main__":
    main()
