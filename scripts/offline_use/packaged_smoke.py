"""macOS packaged Runtime smoke with only loopback networking allowed.
Uses disposable synthetic data, no model requests or installed user state.
Usage: python scripts/offline_use/packaged_smoke.py /path/to/verilex-runtime
"""
import json
import base64
import io
import os
from pathlib import Path
import signal
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
from docx import Document

binary = str(Path(sys.argv[1]).resolve())
profile = '''(version 1) (allow default) (deny network*)
(allow network-inbound (local ip "localhost:*"))
(allow network-outbound (remote ip "localhost:*"))
(allow network-bind (local ip "localhost:*"))'''
with tempfile.TemporaryDirectory(prefix='verilex-offline-packaged-') as temp:
    root = Path(temp)
    notes = root / '合成资料'; notes.mkdir()
    (notes / 'offline.md').write_text('# 离线合成资料\n断网检索标记，内容仅用于验收。', encoding='utf-8')
    with socket.socket() as reservation:
        reservation.bind(('127.0.0.1',0)); port=reservation.getsockname()[1]
    token='offline-use-isolated-check'
    env={**os.environ,'VERILEX_API_TOKEN':token}
    opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
    def request(path, body=None, method=None):
        req=urllib.request.Request(f'http://127.0.0.1:{port}'+path,
            data=body if isinstance(body, bytes) else None if body is None else json.dumps(body).encode(),
            method=method,
            headers={'X-VeriLex-Token':token,'Content-Type':'application/json'})
        with opener.open(req,timeout=4) as response: return json.load(response)
    def start():
        proc=subprocess.Popen(['/usr/bin/sandbox-exec','-p',profile,binary,'--port',str(port),'--state-dir',str(root/'state')],
            env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,start_new_session=True)
        for _ in range(150):
            if proc.poll() is not None: raise AssertionError('isolated runtime exited before readiness')
            try: request('/v1/health'); return proc
            except Exception: time.sleep(.2)
        proc.terminate();proc.wait(timeout=10)
        raise AssertionError('runtime readiness timed out')
    def stop(proc):
        os.killpg(proc.pid,signal.SIGTERM);proc.wait(timeout=15)
    proc=start()
    try:
        assert request('/v1/providers') == []
        ws=request('/v1/workspaces',{'root_path':str(notes),'display_name':'离线合成工作区'})['id']
        thread=request('/v1/threads',{'title':'离线合成历史','model':'unconfigured','workspace_id':ws})['id']
        library=request('/v1/knowledge',{'root':str(notes)})['id']
        query=urllib.parse.urlencode({'q':'断网检索标记'})
        doc=request(f'/v1/knowledge/{library}/documents?{query}')['items'][0]
        original=request(f'/v1/knowledge/{library}/documents/{doc["id"]}?version={doc["version"]}')
        assert '断网检索标记' in original['text']
        preview=request(f'/v1/workspaces/{ws}/file-preview?path=offline.md')
        assert preview['content']==original['text']
        # Only a synthetic, credential-free provider record; no inference is requested.
        with sqlite3.connect(root/'state'/'verilex.sqlite3') as db:
            db.execute("INSERT INTO providers (id,name,kind,base_url,default_model,auth_mode,secret_ref,enabled,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                       ('fixture','离线合成','openai_compatible','http://127.0.0.1:1/v1','fixture','none',None,1,'2026-10-14T00:00:00Z','2026-10-14T00:00:00Z'))
        source_doc=Document(); source_doc.add_paragraph('合成原文：十日内交付。')
        source_stream=io.BytesIO(); source_doc.save(source_stream); source=source_stream.getvalue()
        params=urllib.parse.urlencode({'filename':'离线合成合同.docx','provider_id':'fixture','model':'fixture','skill':'通用合同审查','party':'neutral','strictness':'normal','signature':'VeriLex','retrieval':'false'})
        rid=request('/v1/contract-reviews?'+params,source)['review_id']
        editor=f'/v1/contract-reviews/{rid}/editor'
        paragraph=request(editor)['paragraphs'][0]
        changed=request(editor,{'expected_revision':0,'edits':[{'paragraph_id':paragraph['id'],'original':paragraph['text'],'text':'合成修改：三十日内交付。'}]},'PATCH')
        assert changed['revision']==1
        exported=base64.b64decode(request(editor+'/export',{'expected_revision':1})['content'])
        assert Document(io.BytesIO(exported)).paragraphs[0].text=='合成修改：三十日内交付。'
        assert base64.b64decode(request(f'/v1/contract-reviews/{rid}/preview')['content'])==source
        stop(proc);proc=start()
        assert request(f'/v1/threads/{thread}')['title']=='离线合成历史'
        assert request(f'/v1/knowledge/{library}/documents?{query}')['items'][0]['version']==doc['version']
        assert request(editor)['paragraphs'][0]['text']=='合成修改：三十日内交付。'
        print('PASS: external network denied; packaged startup/restart, local thread persistence, Markdown search/read, file preview, deterministic DOCX edit/export/persistence and original unchanged; no model calls.')
    finally:
        if proc.poll() is None: stop(proc)
